"""Tạo báo cáo Word bằng số liệu đã lưu, không huấn luyện lại mô hình.

Chạy bằng Python của bộ công cụ tài liệu có python-docx. Hình và JSON được
tạo trước bởi build_progress_report_assets.py bằng môi trường của dự án.
"""
from pathlib import Path
import json
import csv
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'reports/professor_report'
D = json.loads((OUT / 'report_data.json').read_text(encoding='utf-8'))
TEST = {r['policy']: r for r in D['test_metrics']}
TEMP = {(r['period'], r['policy']): r for r in D['temporal_metrics']}
LABEL = {'persistence':'Persistence', 'standard':'Bi-LSTM chuẩn',
         'dual_global':'Hai đầu ngưỡng chung', 'dual_precision30':'Hai đầu ngưỡng riêng'}
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.7), Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.header_distance = sec.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0,0,0)
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = True
# Loại đường viền tiêu đề kế thừa từ template Word mặc định.
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
hp = sec.header.paragraphs[0]
hp.text = 'PBL6  |  Dự báo lượng mưa Đà Nẵng  |  Báo cáo tiến độ nghiên cứu'
hp.runs[0].font.size = Pt(8)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('23 tháng 9 năm 2026  •  Trang ').font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); fp._p.append(field)


def p(text, bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para


def h(text):
    return doc.add_heading(text, level=2)


def page(title):
    para = doc.add_heading(title, level=1)
    para.paragraph_format.page_break_before = True


def table(headers, rows, widths=None, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for c,w in zip(t.columns,widths): c.width = Inches(w)
    for i,head in enumerate(headers): t.rows[0].cells[i].text = str(head)
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row): cells[i].text=str(val)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        row.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        trpr=row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,cell in enumerate(row.cells):
            if widths: cell.width=Inches(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'244769' if ri==0 else ('EEF3F7' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1)
                para.paragraph_format.space_before=Pt(1)
                para.paragraph_format.line_spacing=1.03
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or len(str(cell.text))>25 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size=Pt(size)
                    run.font.bold=ri==0
                    run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t


def fig(name, caption, width=7.0):
    para=doc.add_paragraph()
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next=True
    para.paragraph_format.space_after=Pt(3)
    para.add_run().add_picture(str(OUT/'figures'/name),width=Inches(width))
    doc.add_paragraph(caption,style='Caption')


def equation(text):
    # OMML gốc của Word, không chèn LaTeX thô vào tài liệu.
    para=doc.add_paragraph()
    math=OxmlElement('m:oMath')
    run=OxmlElement('m:r'); word=OxmlElement('m:t');word.text=text
    run.append(word);math.append(run);para._p.append(math)
    return para


def f(value, digits=3):
    return '—' if value is None else f'{value:.{digits}f}'


doc.add_paragraph('Báo cáo tiến độ dự báo lượng mưa tại Đà Nẵng',style='Title')
p('Tiền xử lý dữ liệu và phát triển mô hình Bi LSTM cho dự báo từ 1 đến 6 giờ',bold=True)
p('Tài liệu tổng hợp để xây dựng slide báo cáo giảng viên hướng dẫn\nChốt kết quả ngày 23 tháng 9 năm 2026')
h('Kết luận chính')
p('Dự án đã có pipeline xử lý dữ liệu, đặc trưng thời gian, mô hình Bi-LSTM và quy trình đánh giá theo giờ lẫn theo đợt mưa. Nhánh dữ liệu bề mặt đầy đủ giúp tiếp tục nghiên cứu trong khi ERA5 PL850 còn thiếu. Mô hình hai đầu ra giảm sai số lượng mưa và tăng khả năng phát hiện giờ mưa trên test cũ, nhưng phát sinh nhiều cảnh báo sai và chưa ổn định qua các giai đoạn thời gian.')
table(['Kết quả tiêu biểu','Giá trị','Ý nghĩa'],[
    ['Hai đầu so với persistence','RMSE 0,996 so với 1,190 mm/h','Giảm khoảng 16,3% sai số RMSE trên cùng test'],
    ['Ngưỡng riêng theo horizon','FP 1.600 → 564','Giảm 64,8% cảnh báo sai, đồng thời giảm recall'],
    ['Kiểm tra thời gian mở rộng','3 đợt, huấn luyện lại 6 mô hình','Ngưỡng riêng không phát hiện đúng giờ mưa lớn trong cả 3 đợt'],
    ['Xác thực độc lập theo giờ','Chưa có dữ liệu trạm phù hợp','Chưa đủ bằng chứng cho vận hành thời gian thực']
], [2.0,1.8,3.3])
h('Hai giai đoạn nên trình bày')
p('Giai đoạn 1 xây dựng nền tảng: kiểm toán dữ liệu, tạo đặc trưng, chia theo thời gian và huấn luyện Bi-LSTM hồi quy với Weighted MSE. Thử lấy mẫu theo sự kiện được giữ như một kết quả âm có giá trị, vì không cải thiện tiêu chí validation.')
p('Giai đoạn 2 bổ sung đầu cảnh báo mưa lớn, thử ngưỡng theo từng horizon và huấn luyện lại trên ba giai đoạn lịch sử. Kết quả cho thấy cần tối ưu đồng thời sai số lượng mưa, khả năng phát hiện và số cảnh báo sai.')
h('Phạm vi và đóng góp')
p('Sinh viên xác định bài toán, dữ liệu và các hướng thử nghiệm; Codex hỗ trợ cùng triển khai mã nguồn, kiểm tra dữ liệu, chạy thí nghiệm trên GPU và tổng hợp bằng chứng. Báo cáo chỉ ghi nhận kết quả có trong workspace. Informer 24–72 giờ và xác thực bằng Vrain chưa được thực hiện.')
p('Các chỉ số chính là đánh giá hồi cứu so với nhãn Open-Meteo. Test cũ đã được xem trong quá trình phát triển, vì vậy các kết quả bổ sung trên tập này mang tính thăm dò.')

page('1 Bài toán và hiện trạng dữ liệu')
p('Mục tiêu dài hạn là dự báo lượng mưa cục bộ tại Đà Nẵng với đầu vào đa nguồn. Trong phạm vi hiện tại, mô hình nhận 24 giờ quá khứ và trả về đồng thời lượng mưa cho 6 giờ tiếp theo. Tọa độ đại diện là 16,0544°B và 108,2022°Đ; đây là bài toán tại một vị trí đại diện, chưa phải bản đồ mưa toàn thành phố.')
table(['Nguồn','Dữ liệu sử dụng','Hiện trạng'],[
    ['Open-Meteo','Mưa, nhiệt độ, độ ẩm, áp suất và gió 10 m theo giờ','102.264 giờ, 2015–08/2026'],
    ['ERA5 PL850','u, v và độ ẩm riêng tại 850 hPa','Thiếu 15 tháng, tương ứng 11.016 giờ'],
    ['ERA5 SST','Nhiệt độ mặt biển tại ô biển gần nhất','Đủ phạm vi dữ liệu hiện có'],
    ['NOAA ONI','Chỉ số ENSO theo mùa ba tháng, gắn tháng trung tâm','Đọc và hợp nhất; mặc định loại khỏi đầu vào mô hình']
], [1.3,3.0,2.8])
fig('fig01_data_coverage.png','Hình 1. Độ phủ dữ liệu theo tháng. Các tháng thiếu PL850 là 04–12/2024 và 03–08/2026.')
h('Lý do tiếp tục bằng nhánh dữ liệu bề mặt')
p('Chờ tải bù ERA5 sẽ trì hoãn các thí nghiệm mô hình. Nhánh surface-only sử dụng toàn bộ lịch giờ liên tục của Open-Meteo để huấn luyện và so sánh. Nhánh đa nguồn được giữ để đánh giá trên phần giao thời gian hợp lệ; không nội suy kéo dài qua các tháng thiếu.')
p('Open-Meteo Historical Weather API cung cấp dữ liệu mô hình/tái phân tích, không phải chuỗi đo mưa trạm tại đúng tọa độ. Cấu hình tải hiện tại chưa cố định một nguồn mô hình duy nhất, nên cần kiểm tra tính đồng nhất nguồn theo năm trước khi diễn giải biến đổi khí hậu hoặc triển khai [1].')

page('2 Tiền xử lý và hợp nhất dữ liệu')
fig('fig02_pipeline.png','Hình 2. Quy trình xử lý từ dữ liệu thô đến đánh giá. Các quyết định huấn luyện và chọn ngưỡng chỉ sử dụng train và validation.')
table(['Bước','Cách xử lý đang áp dụng'],[
    ['Chuẩn hóa thời gian','Chuyển về Asia/Ho_Chi_Minh, UTC+7; sắp xếp và ghép theo giờ. Phân biệt giờ hợp lệ của số đo với giờ dữ liệu thực sự có sẵn.'],
    ['Chuẩn hóa biến vật lý','Đồng nhất tên biến; chuyển SST từ kelvin sang °C trước khi trừ nhiệt độ không khí.'],
    ['Chọn vị trí ERA5','PL850 tại 16,00°B, 108,25°Đ; SST tại ô biển 16,25°B, 108,25°Đ.'],
    ['Hợp nhất đa nguồn','Phần giao có dữ liệu vật lý đầy đủ: 91.248 giờ. Xuất CSV, Parquet và chẩn đoán merge.'],
    ['Dữ liệu thiếu','Cho phép nội suy khoảng ngắn tối đa 2 giờ ở biến đầu vào theo cấu hình; số giá trị thực sự đã nội suy trong lần merge này là 0. Không nội suy lượng mưa làm nhãn.'],
    ['Khoảng trống dài','Tái lập lịch giờ trước tạo cửa sổ; loại mọi cửa sổ đi qua gap hoặc chứa giá trị thiếu.'],
    ['ONI','Forward-fill mốc tháng khi hợp nhất. Chưa có lịch công bố nên không dùng ONI hồi cứu như biến thời gian thực.']
], [1.55,5.55],size=10.5)
p('Nội suy tuyến tính hai phía, nếu dùng cho biến đầu vào trong dữ liệu hồi cứu, có thể cần số đo xuất hiện sau giờ dự báo. Khi chuyển sang chạy thực, phải kiểm tra available_at hoặc thay bằng phép xử lý chỉ dùng quá khứ. Việc chia train/test đúng thứ tự thời gian chưa tự chứng minh toàn bộ dữ liệu đã sẵn có tại thời điểm phát hành.')
p('CAPE, CIN và TCWV được tạm bỏ theo lựa chọn của đề tài. Đây là quyết định giảm phạm vi baseline, chưa phải bằng chứng các biến này không có ích. Dòng ẩm và ảnh hưởng địa hình hiện mới được biểu diễn gián tiếp qua biến khí tượng; chưa có bản đồ địa hình hay trường khí tượng không gian làm đầu vào.')

page('3 Kỹ thuật đặc trưng')
p('Nhánh bề mặt dùng 22 đặc trưng. Mỗi hàng biểu diễn một giờ; các đặc trưng chỉ sử dụng thông tin tại giờ đó và các giờ trước. Lượng mưa hiện tại được xem là đã biết ở cuối cửa sổ đầu vào trong thí nghiệm hồi cứu.')
table(['Nhóm','Số biến','Nội dung'],[
    ['Biến khí tượng gốc',6,'temperature_2m, relative_humidity_2m, surface_pressure, wind_speed_10m, wind_direction_10m, precipitation'],
    ['Chu kỳ thời gian',4,'sin_hour, cos_hour, sin_month, cos_month'],
    ['Mưa trễ',6,'Mưa trễ 1, 2, 3, 6, 12 và 24 giờ'],
    ['Độ ẩm bề mặt trễ',3,'Độ ẩm 2 m trễ 1, 3 và 6 giờ'],
    ['Tích lũy mưa quá khứ',3,'Tổng mưa 3, 6 và 24 giờ; shift(1) trước rolling']
], [1.65,.65,4.8],size=10.5)
h('Cách xây dựng và ý nghĩa')
equation('sin_hour = sin(2π × hour / 24)    cos_hour = cos(2π × hour / 24)')
equation('sin_month = sin(2π × month / 12)    cos_month = cos(2π × month / 12)')
p('Cặp sin/cos giữ quan hệ gần nhau giữa 23 giờ và 0 giờ hoặc giữa tháng 12 và tháng 1. Biến trễ mô tả tính duy trì của mưa và độ ẩm. Tổng mưa quá khứ mô tả trạng thái trước thời điểm dự báo; shift(1) bảo đảm không cộng nhãn tương lai vào đặc trưng.')
h('Đặc trưng bổ sung trong nhánh đa nguồn')
p('Nhánh đa nguồn có 31 đặc trưng sau khi loại ONI khỏi đầu vào. Các biến bổ sung gồm u_850, v_850, specific_humidity_850, SST, ba biến trễ độ ẩm riêng 1–3–6 giờ, tốc độ gió 850 hPa và chênh lệch SST với nhiệt độ 2 m.')
p('Tốc độ gió 850 hPa được tính bằng căn bậc hai của tổng bình phương u và v. Chênh lệch SST–không khí là SST trừ nhiệt độ 2 m sau khi đồng nhất °C. Đây là các chỉ báo vật lý tiềm năng; chưa thể kết luận quan hệ nhân quả chỉ từ tương quan hoặc một lần huấn luyện.')
h('Điểm cần cải thiện')
p('Hướng gió bề mặt vẫn được biểu diễn bằng góc. Một thử nghiệm tiếp theo có thể mã hóa sin/cos hướng gió hoặc thành phần u–v để tránh điểm gián đoạn 359°–0°. Cần so sánh trên cùng giao thức thời gian, không thay đặc trưng rồi chọn theo test cũ.')

page('4 Trực quan hóa trước huấn luyện')
fig('fig03_train_rainfall_eda.png','Hình 3. Phân bố lượng mưa và trung bình theo tháng, chỉ tính trên phần train. Trục số giờ ở hình trái dùng thang log.',width=7.0)
p(f'Trong {D["train_hours"]:,} giờ thuộc khoảng train, có {D["train_rain_counts"][0]:,} giờ mưa bằng 0 ({100*D["train_zero_fraction"]:.2f}%) và chỉ {D["train_rain_counts"][3]} giờ mưa >5 mm/h ({100*D["train_heavy_fraction"]:.3f}%). Giá trị lớn nhất là {D["train_max_rain"]:.1f} mm/h. Tỷ lệ hiếm giải thích vì sao sai số trung bình thấp có thể đi kèm bỏ sót nhiều giờ mưa lớn.')
fig('fig04_train_feature_correlations.png','Hình 4. Pearson r giữa một số đặc trưng và nhãn t+1, t+6 trong train; cả đầu vào và nhãn dùng tính tương quan đều nằm trong train.',width=5.7)
p('Tương quan hỗ trợ hiểu dữ liệu, không phải phép đánh giá nhân quả. Phụ thuộc phi tuyến và tương tác theo thời gian có thể không thể hiện đầy đủ trong heatmap. Không dùng tương quan trên test để lựa chọn đặc trưng.')

page('5 Chia dữ liệu và giao thức đánh giá')
table(['Phần dữ liệu','Mốc thời gian địa phương','Số cửa sổ'],[
    ['Train 70%','Trước 02/03/2023 23:00',f'{D["split"]["train_windows"]:,}'],
    ['Validation 15%','Từ 02/03/2023 23:00 đến trước 01/12/2024 03:00',f'{D["split"]["val_windows"]:,}'],
    ['Test 15%','Từ 01/12/2024 03:00 đến hết dữ liệu 01/09/2026 06:00',f'{D["split"]["test_windows"]:,}']
], [1.3,4.5,1.3])
p('Cửa sổ đầu vào gồm 24 giờ, nhãn gồm 6 giờ ngay sau đó. Cả 6 nhãn phải thuộc cùng một phần train, validation hoặc test. Đầu vào validation/test có thể nhìn lại các giờ quá khứ trước ranh giới; điều này phù hợp với dự báo cuốn chiếu. RobustScaler chỉ fit trên 71.560 hàng train hợp lệ. 24 hàng đầu thiếu đặc trưng trễ và 10 cửa sổ có nhãn cắt ranh giới được loại.')
h('Các tiêu chí đánh giá')
p('MAE đo sai số tuyệt đối trung bình; RMSE nhạy hơn với sai số lớn; R² đo phần phương sai giải thích được; CC là tương quan Pearson. Mưa >5 mm/h là ngưỡng nghiên cứu của dự án, không tự đồng nghĩa với phân loại mưa cực đoan chính thức.')
p('Đối với cảnh báo theo giờ: TP là báo đúng, FP là báo khi nhãn không vượt ngưỡng, FN là bỏ sót. Precision bằng TP/(TP+FP); recall bằng TP/(TP+FN); F1 là trung bình điều hòa của precision và recall. Không dùng accuracy tổng thể làm tiêu chí chính vì phần lớn giờ không mưa lớn.')
p('Đánh giá theo đợt nối các giờ mưa vượt ngưỡng nếu cách nhau không quá 6 giờ và lịch giờ không bị thiếu. Khoảng sự kiện kéo từ giờ vượt ngưỡng đầu đến cuối. Ghép tối đa một-một các khoảng dự báo và thực tế có giao nhau. Vì vậy một cảnh báo có thể trúng khoảng sự kiện dù không trúng đúng giờ vượt ngưỡng.')
h('Diễn giải số mẫu')
p('Test có 92.010 cặp horizon–giờ, gồm 702 cặp mưa lớn. Con số 264 đợt trong bảng gộp là tổng theo 6 horizon, tương ứng 44 đợt ở mỗi horizon; không phải 264 trận mưa độc lập. Các cửa sổ chồng lấn và nhãn lặp theo horizon không tạo ra các quan sát thống kê độc lập.')
p('Test này đã được sử dụng để thảo luận trong quá trình phát triển. Không gọi các cải tiến tiếp theo trên test cũ là xác nhận cuối cùng. Ba đợt kiểm tra bổ sung đều dùng dữ liệu trước test và fit lại mô hình theo thứ tự thời gian; chúng vẫn là thí nghiệm hồi cứu trên dữ liệu dự án đã có.')

page('6 Hai giai đoạn phát triển mô hình')
fig('fig05_architecture.png','Hình 5. Backbone dùng chung của Bi-LSTM. Giai đoạn 2 bổ sung nhánh dự báo xác suất vượt 5 mm/h. Đầu lượng mưa giai đoạn 1 được chặn về không âm khi suy luận; giai đoạn 2 dùng Softplus.')
table(['Thành phần','Giai đoạn 1','Giai đoạn 2'],[
    ['Đầu vào','24 giờ × 22 đặc trưng','24 giờ × 22 đặc trưng'],
    ['Backbone','Bi-LSTM 2 tầng, hidden 50 mỗi chiều','Cùng cấu hình backbone'],
    ['Regularization','Dropout 0,2','Dropout 0,2'],
    ['Đầu ra','6 lượng mưa, Linear → ReLU → Linear','6 lượng mưa Softplus và 6 logits mưa lớn'],
    ['Hàm mất mát','Weighted MSE, trọng số 4 khi nhãn >5 mm/h','Weighted Huber trọng số 4 khi nhãn >5 mm/h + 0,5 × BCE'],
    ['Trọng số BCE','Không áp dụng','pos_weight = 12'],
    ['Tối ưu','Adam, learning rate 0,001','Adam, learning rate 0,001'],
    ['Chọn mô hình','Validation loss; patience 7','Validation loss; patience 7']
], [1.35,2.85,2.9],size=10)
p('Bi-LSTM xử lý hai chiều bên trong cửa sổ quá khứ; không đưa dữ liệu của 6 giờ cần dự báo vào đầu vào. Trạng thái cuối của hai chiều được ghép thành vector 100 chiều trước lớp tuyến tính. Việc dùng bidirectional vì vậy không tự tạo rò rỉ tương lai trong cấu hình này [3].')
p('Giai đoạn 2 phân tách hai mục tiêu: ước lượng lượng mưa và nhận diện vượt ngưỡng. Xác suất sau sigmoid là điểm cảnh báo của mô hình; weighted BCE không bảo đảm các điểm này đã được hiệu chuẩn xác suất. Ngưỡng 0,40 là giá trị chọn trên validation cũ, không phải mặc định 0,50.')

page('7 Quá trình huấn luyện trên GPU')
fig('fig06_training_curves.png','Hình 6. Đường loss của hai giai đoạn trên nhánh bề mặt. Hai loss có công thức và thang đo khác nhau nên không so trực tiếp độ cao giữa hai biểu đồ.')
table(['Thử nghiệm','Epoch tốt nhất','Số epoch chạy','Quyết định'],[
    ['Bi-LSTM chuẩn',2,9,'Baseline được giữ'],
    ['Bi-LSTM lấy mẫu sự kiện',11,18,'Không chọn do validation F1 giảm'],
    ['Bi-LSTM hai đầu',2,9,'Giữ để phân tích lượng mưa và cảnh báo']
], [2.55,1.05,1.05,2.45])
p('Huấn luyện dùng CUDA, Adam với learning rate 0,001, batch size 256, seed 42 và early stopping patience 7. Đánh giá bằng checkpoint có validation loss thấp nhất. Ở hai mô hình chính, train loss tiếp tục giảm sau epoch 2 trong khi validation loss tăng hoặc dao động cao hơn. Khả năng tổng quát hóa chưa tốt; chạy thêm epoch không mặc nhiên cải thiện kết quả.')
h('Thử nghiệm lấy mẫu theo sự kiện')
p('Train có 1.212 cửa sổ chứa mưa lớn trong 71.531 cửa sổ, khoảng 1,69%, nhóm thành 135 đợt. Biến thể lấy mẫu tăng tỷ lệ cửa sổ mưa lớn được lấy lên khoảng 25%. Tuy nhiên, F1 theo giờ trên validation giảm từ khoảng 0,274 xuống 0,094; RMSE test tăng từ 0,999 lên 1,048 mm/h. Vì vậy biến thể này không được chọn.')
p('Thử nghiệm này gợi ý rằng chỉ tăng tần suất xuất hiện các cửa sổ mưa hiếm chưa giải quyết được bài toán. Các cửa sổ gần nhau trong cùng sự kiện có thể lặp nhiều thông tin; thay đổi cách lấy mẫu cũng làm thay đổi phân bố đầu vào mà mô hình học. Đây là cách diễn giải cần kiểm chứng thêm, không phải kết luận nhân quả từ một lần chạy.')
h('Khả năng tái lập')
p('Đã lưu checkpoint, cấu hình, scaler, lịch sử loss và dự đoán. Lần kiểm tra bổ sung huấn luyện lại 3 Bi-LSTM chuẩn và 3 mô hình hai đầu. Bộ 19 kiểm thử tự động đã chạy thành công, bao gồm xử lý gap, ranh giới chia tập, cảnh báo theo đợt, dữ liệu trạm và kiểm tra thời điểm dữ liệu có sẵn. Chưa chạy nhiều seed để ước lượng độ biến động huấn luyện.')

page('8 Kết quả trên cùng tập test hồi cứu')
table(['Mô hình','MAE','RMSE','R²','CC'],[
    [LABEL[k],f(TEST[k]['MAE']),f(TEST[k]['RMSE']),f(TEST[k]['R2']),f(TEST[k]['CC'])]
    for k in ['persistence','standard','dual_global']
], [2.5,1.15,1.15,1.15,1.15],size=10.5)
p('Đơn vị MAE và RMSE là mm/h. Persistence dự báo cả 6 giờ bằng lượng mưa cuối cửa sổ đầu vào. Mô hình hai đầu giảm RMSE khoảng 16,3% so với persistence, nhưng chỉ khoảng 0,3% so với Bi-LSTM chuẩn. Chưa có nhiều seed hoặc khoảng tin cậy để khẳng định chênh lệch nhỏ 0,003 mm/h là ổn định.')
table(['Chính sách','Precision','Recall','F1 giờ','F1 đợt'],[
    [LABEL[k],f(TEST[k]['precision']),f(TEST[k]['recall']),f(TEST[k]['hourly_F1']),f(TEST[k]['event_F1'])]
    for k in ['persistence','standard','dual_global','dual_precision30']
], [2.5,1.15,1.15,1.15,1.15],size=10.5)
p('Bi-LSTM chuẩn có tương quan cao nhất trong ba mô hình nhưng bỏ sót nhiều mưa lớn. Đầu cảnh báo với ngưỡng chung tăng recall giờ mưa lớn lên 47,2%, đổi lại precision chỉ 17,1%. Persistence vẫn có F1 theo giờ cao hơn cả hai chính sách cảnh báo của mô hình hai đầu trên test này.')
table(['Chính sách','TP','FP','FN','Đợt trúng / tổng'],[
    [LABEL[k],TEST[k]['TP'],TEST[k]['FP'],TEST[k]['FN'],f'{TEST[k]["event_hit"]} / {TEST[k]["event_observed"]}']
    for k in ['persistence','standard','dual_global','dual_precision30']
], [2.5,.85,.85,.85,2.05])
h('Nhận xét nên dùng khi báo cáo')
p('Kết quả lượng mưa trung bình có cải thiện, nhưng hệ thống cảnh báo còn đánh đổi mạnh giữa bỏ sót và báo sai. Không nên dùng một chỉ số RMSE để kết luận mô hình đã dự báo tốt mưa lớn. Ngưỡng riêng làm F1 theo đợt tăng trên test cũ nhưng không cải thiện đồng thời F1 theo giờ và recall.')
p('Các kết quả cùng bảng dùng đúng một tập test và một tập nhãn. Hai hàng ngưỡng của mô hình hai đầu dùng cùng dự đoán lượng mưa; thay ngưỡng chỉ thay cảnh báo, không thay MAE, RMSE, R² hoặc CC.')

page('9 Hiệu quả theo từng thời hạn dự báo')
fig('fig07_test_horizon_skill.png','Hình 7. So sánh RMSE, recall và precision từ t+1 đến t+6 trên test cũ. Đường lượng mưa của mô hình hai đầu không phụ thuộc ngưỡng cảnh báo.')
with (ROOT/'reports/threshold_temporal_study/test_metrics_by_horizon.csv').open(encoding='utf-8') as fh:
    byh=list(csv.DictReader(fh))
rows=[]
for hh in range(1,7):
    a=next(x for x in byh if x['policy']=='dual_global' and int(x['horizon_h'])==hh)
    b=next(x for x in byh if x['policy']=='dual_precision30' and int(x['horizon_h'])==hh)
    c=next(x for x in byh if x['policy']=='standard' and int(x['horizon_h'])==hh)
    rows.append([f't+{hh}',f(float(c['RMSE'])),f(float(a['RMSE'])),f(float(a['POD_recall'])),f(float(b['POD_recall']))])
table(['Horizon','RMSE chuẩn','RMSE hai đầu','Recall chung','Recall riêng'],rows,[1.0,1.5,1.5,1.55,1.55])
p('Chất lượng không đồng đều giữa các horizon. Cần trình bày cả đường cong thay vì chỉ trung bình 6 giờ. Ở t+1, persistence tận dụng mạnh tính liên tục của lượng mưa hiện tại; việc một mô hình có RMSE gộp tốt hơn không bảo đảm nó tốt hơn persistence ở từng horizon.')
p('Với ngưỡng chung của mô hình hai đầu, số đợt được phát hiện tại t+6 là 11/44, trong khi Bi-LSTM chuẩn chỉ đạt 2/44. Tuy nhiên, cảnh báo theo đợt có tiêu chí giao khoảng nên cần xem đồng thời số giờ cảnh báo sai và khả năng trúng đúng giờ.')
p('Đường recall sau chỉnh ngưỡng giảm ở các horizon dài phản ánh giới hạn thông tin của đầu vào một điểm với 24 giờ lịch sử. Chưa có đầu vào radar, trường mây hay dự báo số trị tương lai để mô tả các hệ thống mưa đang di chuyển tới vị trí dự báo.')

page('10 Thử giảm cảnh báo sai bằng ngưỡng riêng')
fig('fig08_threshold_tradeoff.png','Hình 8. Cảnh báo đúng và sai trên test cũ; các ngưỡng bên phải được chọn bằng validation trước khi đọc dự đoán test trong lần chạy bổ sung.')
p('Giao thức đặt mục tiêu precision tối thiểu 0,30 trong calibration, có ít nhất 10 cảnh báo và ít nhất một TP. Trong các ngưỡng đạt điều kiện, chọn recall cao nhất, sau đó ưu tiên precision và ngưỡng cao hơn khi hòa. Dải thử là 0,40–0,95 bước 0,05 và 0,99. Nếu không có ứng viên đạt yêu cầu, horizon bị tắt cảnh báo.')
table(['Horizon','t+1','t+2','t+3','t+4','t+5','t+6'],[
    ['Ngưỡng đã chọn',.55,.60,.70,.75,.75,.75],
    ['Precision validation']+[f(v['calibration_precision']) for v in D['thresholds']['horizons']]
], [1.8,.88,.88,.88,.88,.88,.9])
p('Trên test cũ, số FP giảm từ 1.600 xuống 564, tương đương 64,8%. Tuy vậy, TP giảm từ 331 xuống 178; recall giảm từ 47,2% xuống 25,4%. Precision chỉ đạt 24,0%, thấp hơn mục tiêu 30% trên calibration. Ràng buộc trong mẫu hiệu chỉnh không bảo đảm cùng precision khi phân bố dữ liệu thay đổi.')
p('F1 theo đợt tăng từ 0,258 lên 0,293 trên test cũ; F1 theo giờ giảm từ 0,251 xuống 0,247. Trên validation, F1 theo đợt giảm từ khoảng 0,293 xuống 0,217. Vì vậy bằng chứng hiện tại không đủ để chọn phương án này làm cấu hình vận hành chung.')
p('Ngưỡng và kết quả được lưu riêng dưới reports/threshold_temporal_study. Chính sách mới là thí nghiệm bổ sung; chưa thay ngưỡng của các công cụ ghi dự báo thực hoặc bộ đánh giá trạm đã đóng băng trước đó.')

page('11 Kiểm tra ổn định theo thời gian')
p('Mỗi fold huấn luyện từ 2015 đến mốc kết thúc train, fit scaler trên train, chọn checkpoint và ngưỡng trong calibration rồi đánh giá giai đoạn sau. Tất cả nhãn đều trước 01/12/2024; không dùng test cũ để fit lại các mô hình này.')
table(['Fold','Train kết thúc trước','Calibration','Assessment'],[
    ['F1','01/01/2023','01–06/2023','07–12/2023'],
    ['F2','01/07/2023','07–12/2023','01–06/2024'],
    ['F3','01/01/2024','01–06/2024','07–11/2024']
], [0.6,2.0,2.25,2.25])
fig('fig10_temporal_stability.png','Hình 9. Kết quả sau khi fit lại mô hình và ngưỡng ở mỗi fold. Ngưỡng chung được chọn lại, lần lượt 0,05; 0,10; 0,15, không cố định 0,40.',width=7.0)
table(['Fold','Giờ mưa lớn mỗi horizon','RMSE chuẩn','RMSE hai đầu','Recall ngưỡng riêng'],[
    [fold,int(TEMP[(fold,'standard')]['TP']+TEMP[(fold,'standard')]['FN'])//6,
     f(TEMP[(fold,'standard')]['RMSE']),f(TEMP[(fold,'dual_global')]['RMSE']),f(TEMP[(fold,'dual_precision30')]['recall'])]
    for fold in ['F1','F2','F3']
], [.6,2.0,1.5,1.5,1.5],size=10)
p('Mô hình hai đầu có RMSE thấp hơn Bi-LSTM chuẩn ở cả ba fold. Tuy nhiên, chính sách precision 30% tắt cả 6 horizon ở F1 và F3; F2 chỉ bật t+1, tạo 4 FP và 0 TP ở assessment. Recall theo giờ của chính sách này bằng 0 ở cả ba fold. F2 vẫn có một đợt trúng do tiêu chí giao khoảng, minh họa sự khác biệt giữa hai cách chấm.')
p('F2 chỉ có 5 giờ mưa lớn mỗi horizon, nên kết quả cảnh báo rất nhạy với ít sự kiện. Các fold có mùa và tần suất mưa khác nhau; train chồng lấn và chỉ một seed. Kết luận phù hợp là chưa có độ ổn định để vận hành, không phải ngưỡng riêng đã giải quyết được cảnh báo sai.')

page('12 Minh họa một đợt mưa lớn trên test')
fig('fig09_heavy_event_example.png','Hình 10. Khoảng ±36 giờ quanh đỉnh lớn nhất của nhãn test cũ. Đợt được chọn bằng nhãn để minh họa sai số tại cực đại; không dùng hình này như bằng chứng đại diện cho toàn bộ test.',width=7.0)
table(['Tại đỉnh 11/06/2025 21:00','Nhãn','Bi-LSTM chuẩn','Hai đầu','Persistence'],[
    [f't+{v["horizon"]}']+[f(v[k],2) for k in ['actual','standard','dual','persistence']]
    for v in D['example_peak_predictions']
], [2.1,1.0,1.5,1.25,1.25])
p('Đỉnh nhãn Open-Meteo đạt 41,9 mm/h. Ở t+1, Bi-LSTM chuẩn dự báo 8,55 mm/h và mô hình hai đầu 4,88 mm/h. Ở t+6, hai giá trị lần lượt chỉ khoảng 1,11 và 0,77 mm/h. Đây là biểu hiện đánh giá thấp đỉnh mưa dù RMSE trung bình toàn tập tương đối thấp.')
p('Một đầu cảnh báo có thể báo nguy cơ vượt ngưỡng trong khi đầu lượng mưa vẫn dự báo thấp. Khi làm slide cần phân biệt dự báo cường độ mưa với cảnh báo xác suất, và luôn ghi horizon cùng thời điểm hợp lệ trên hình.')

page('13 Sai số dự báo và giá trị của dữ liệu đa nguồn')
fig('fig11_actual_prediction_scatter.png','Hình 11. Nhãn và dự báo lượng mưa của mô hình hai đầu tại t+1 và t+6. Đường chéo là dự báo hoàn hảo; các điểm đỉnh nằm xa dưới đường cho thấy đánh giá thấp mưa mạnh.',width=6.8)
h('Nhận xét về phân bố sai số')
p('Phần lớn mẫu tập trung gần gốc vì mưa bằng 0 hoặc nhỏ. Sai số ở vài đỉnh lớn ảnh hưởng mạnh tới RMSE, nhưng vẫn có thể bị che bởi số lượng lớn giờ ít mưa khi đọc MAE. Cần theo dõi sai số có điều kiện trên giờ mưa >5 mm/h, đỉnh sự kiện và độ lệch thời gian xuất hiện đỉnh.')
h('Thử nghiệm đa nguồn đã có')
table(['Thí nghiệm cùng phần giao thời gian','RMSE mm/h','Cách hiểu'],[
    ['Chỉ dữ liệu bề mặt', '1,183', 'Đối chứng trên đúng phần giao dữ liệu'],
    ['Bề mặt và ERA5', '1,146', 'Giảm khoảng 3,1% RMSE so với đối chứng']
], [3.2,1.2,2.7])
p('Kết quả này cung cấp tín hiệu rằng ERA5 có thể hữu ích. Tuy nhiên, test của nhánh đa nguồn khác test đầy đủ của nhánh bề mặt; không so trực tiếp RMSE 1,146 với 0,996 để kết luận nhánh nào tốt hơn. Phép so sánh hợp lệ là 1,183 với 1,146 trong cùng giao thức của ablation đã lưu.')
p('Nhánh đa nguồn có 91.248 hàng và các đoạn thời gian bị thiếu. Trước khi mở rộng thí nghiệm cần bổ sung dữ liệu, kiểm tra đơn vị và khả năng sẵn có tại thời điểm phát hành. SST ở ô biển lân cận cũng không phải số đo nhiệt độ biển tại chính điểm dự báo trên đất liền.')
p('CAPE, CIN và TCWV vẫn là các biến ứng viên. Chúng nên được bổ sung theo nhóm và đánh giá bằng ablation khi có dữ liệu phù hợp; chưa cần đưa toàn bộ vào baseline để hoàn thành vòng nghiên cứu hiện tại.')

page('14 Kiểm chứng độc lập và khả năng chạy thực')
h('Kết quả kiểm tra NOAA năm 2025')
p('Đã kiểm tra trạm DA NANG VMW00041003, tọa độ 16,05°B và 108,20°Đ. Tệp GHCNh có 17.189 hàng nhưng không có giá trị lượng mưa tích lũy 1 giờ. Có 216 báo cáo 6 giờ, 271 báo cáo 12 giờ và 358 báo cáo 24 giờ. Kiểm tra ISD bổ sung cũng chưa tìm được chuỗi lượng mưa 1 giờ phù hợp [4, 5].')
table(['Đối chiếu trên 216 mẫu 6 giờ','RMSE mm/6h','CC'],[
    ['Tổng nhãn Open-Meteo', '12,729', '0,797'],
    ['Tổng dự báo Bi-LSTM chuẩn', '18,485', '0,392'],
    ['Persistence', '20,603', '0,296']
], [4.1,1.5,1.5])
p('Trong 216 mẫu có 163 báo cáo dương và 53 báo cáo mưa vết; mẫu không đại diện đầy đủ mọi giờ khô. Các số trên chỉ là chẩn đoán có điều kiện trên báo cáo hiện có, dùng giả định DATE là cuối kỳ tích lũy 6 giờ. Cần xác nhận quy ước thời gian; không tính recall hoặc tỷ lệ báo sai theo giờ từ tập thưa này. Mô hình hai đầu chưa được xác nhận độc lập trong phép so sánh này.')
h('Vrain và quy trình tiếp nhận dữ liệu trạm')
p('Hiện chưa có API key hoặc tệp xuất theo giờ của Vrain, nên tạm hoãn nhánh xác thực này. Đã chuẩn bị nhập CSV/XLSX có tên trạm, tọa độ, múi giờ, thời lượng tích lũy và quy ước đầu/cuối kỳ; kiểm tra trùng giờ và khoảng cách; không biến dữ liệu thiếu thành mưa bằng 0. Nếu có tệp xuất hợp lệ trong tương lai, có thể dùng quy trình này mà không phụ thuộc tải qua API.')
h('Kiểm tra thời điểm dữ liệu thực sự có sẵn')
p('Audit hiện tại thiếu issued_at ở dự đoán và available_at ở dữ liệu đầu vào; số cửa sổ được xác minh là 0. Vì vậy chưa thể gọi các kết quả hồi cứu là kết quả vận hành thời gian thực. Công cụ ghi dự báo đã được chuẩn bị để lưu giờ phát hành, giờ ghi nhận, snapshot đầu vào và mã băm, đồng thời từ chối dữ liệu có sẵn quá muộn. Chưa có kết quả đánh giá trên luồng dự báo thực mới.')
p('Tên đề tài hướng đến thời gian thực là mục tiêu nghiên cứu. Slide kết quả hiện tại cần ghi “thí nghiệm hồi cứu trên dữ liệu lịch sử”, tránh nhầm mục tiêu với năng lực đã được chứng minh.')

page('15 Kết luận và hướng nghiên cứu tiếp theo')
h('Những phần đã hoàn thành')
p('Đã xây dựng pipeline đa nguồn và nhánh bề mặt liên tục; tạo đặc trưng và cửa sổ có kiểm tra gap; fit scaler trên train; huấn luyện Bi-LSTM chuẩn và hai đầu trên GPU; so sánh với persistence; đánh giá theo giờ, theo horizon và theo đợt; thử ngưỡng riêng và kiểm tra ba giai đoạn thời gian; chuẩn bị công cụ nhập trạm và kiểm tra khả năng sẵn có của dữ liệu.')
h('Những kết luận dữ liệu hiện tại hỗ trợ')
p('Mô hình hai đầu cải thiện RMSE so với persistence trên test cũ và thấp hơn Bi-LSTM chuẩn trong cả ba fold bổ sung. Đầu cảnh báo tăng khả năng phát hiện nhưng có nhiều FP. Giảm FP bằng mục tiêu precision 30% làm mất khả năng phát hiện trên các fold mới; phương án này chưa phù hợp làm ngưỡng vận hành mặc định.')
table(['Ưu tiên','Hành động tiếp theo','Tiêu chí đánh giá'],[
    ['1','Thiết lập dữ liệu và dự báo được ghi thực sự theo thời gian, hoặc một giai đoạn xác nhận chưa dùng để lựa chọn','Mỗi dự báo có issued_at và input available_at; khóa cấu hình trước đánh giá'],
    ['2','Xác định chi phí bỏ sót và báo sai; thử hiệu chuẩn điểm xác suất và kiểm tra theo mùa trên dữ liệu phát triển','Báo đồng thời precision, recall, số cảnh báo và độ ổn định; không chọn theo test cũ'],
    ['3','Phân tích từng sự kiện thất bại; thử biến hướng gió và baseline bổ sung trên cùng split','Ablation có đối chứng, nhiều seed cho các cấu hình vào vòng cuối'],
    ['4','Bổ sung ERA5 theo từng nhóm thời gian và tìm tệp trạm hợp lệ khi có điều kiện','Cùng nhãn, cùng giờ test và nguồn dữ liệu được kiểm chứng'],
    ['5','Sau khi ổn định dữ liệu và giao thức, triển khai Informer cho 24–72 giờ','Đầu vào và baseline phù hợp dự báo dài; báo cáo riêng với bài toán 1–6 giờ']
], [.65,3.3,3.15],size=10.5)
h('Các giới hạn cần trình bày với giảng viên')
p('Nhãn Open-Meteo là dữ liệu mô hình/tái phân tích; chưa có xác thực trạm theo giờ. Thí nghiệm một điểm chưa mô tả đầy đủ địa hình và chuyển động hệ thống mưa. Test cũ đã được xem; các fold hồi cứu không thay thế đánh giá tương lai độc lập. Chỉ một seed, số sự kiện hiếm và các cửa sổ chồng lấn hạn chế độ chắc chắn của kết luận. Informer chưa được triển khai [6].')

page('16 Gợi ý kịch bản 12 slide')
p('Mạch trình bày đề xuất đi từ vấn đề dữ liệu đến các quyết định mô hình, sau đó chứng minh cả cải thiện và giới hạn. Có thể dùng nguyên các hình PNG 300 DPI trong thư mục figures; tên hình được ghi dưới đây.')
table(['Slide','Nội dung chính và thông điệp','Hình hoặc bảng'],[
    [1,'Bài toán Đà Nẵng, đầu vào 24 giờ, đầu ra 1–6 giờ; mục tiêu dài hạn 24–72 giờ','Tóm tắt trang đầu'],
    [2,'Dữ liệu đa nguồn và 15 tháng thiếu ERA5; lý do mở nhánh bề mặt','fig01_data_coverage'],
    [3,'Chuẩn hóa thời gian, đơn vị, gap và ngăn rò rỉ dữ liệu','fig02_pipeline'],
    [4,'Mưa lớn hiếm; 22 đặc trưng và tương quan trên train','fig03 và fig04'],
    [5,'Chia 70/15/15 theo thời gian; persistence và cách chấm sự kiện','Bảng ở mục 5'],
    [6,'Giai đoạn 1 Bi-LSTM chuẩn và thử lấy mẫu sự kiện','fig05_architecture'],
    [7,'Giai đoạn 2 thêm đầu cảnh báo; loss và early stopping','fig06_training_curves'],
    [8,'Kết quả cùng test: RMSE cải thiện nhưng cảnh báo còn yếu','Bảng ở mục 8 và fig07'],
    [9,'Ngưỡng riêng giảm FP 64,8% nhưng recall giảm','fig08_threshold_tradeoff'],
    [10,'Ba fold phát hiện chính sách ngưỡng chưa ổn định','fig10_temporal_stability'],
    [11,'Ví dụ đánh giá thấp đỉnh mưa; giới hạn của nhãn và NOAA','fig09 và fig11'],
    [12,'Kết luận, phần đã làm và kế hoạch dữ liệu cùng đánh giá thực','Mục 14 và 15']
], [.55,4.55,2.0],size=10.5)
h('Câu trả lời ngắn cho các câu hỏi có thể gặp')
p('Vì sao dùng Bi-LSTM vẫn không nhìn trước tương lai? Hai chiều chỉ chạy trong 24 giờ quá khứ. Vì sao RMSE tốt mà cảnh báo kém? Phần lớn mẫu ít mưa, còn sự kiện lớn rất hiếm. Vì sao chưa chuyển ngay sang Informer? Bài toán nhãn và giao thức đánh giá cần ổn định trước khi tăng độ phức tạp. Vì sao chưa đạt mục tiêu thời gian thực? Chưa chứng minh toàn bộ dữ liệu đầu vào sẵn có tại giờ phát hành.')

page('17 Tái lập kết quả và nguồn tham khảo')
h('Các lệnh đã dùng cho vòng bổ sung')
p('Chạy trong thư mục gốc dự án với môi trường .venv đã cài PyTorch CUDA. Lệnh đầu huấn luyện lại 6 mô hình, nên không cần chạy chỉ để đọc báo cáo.')
for command in [
    '.venv\\Scripts\\python.exe run_threshold_temporal_study.py --device cuda',
    '.venv\\Scripts\\python.exe build_progress_report_assets.py',
    '.venv\\Scripts\\python.exe -m unittest discover -s tests -v']:
    para=p(command)
    for run in para.runs: run.font.name='Consolas'; run.font.size=Pt(9)
p('build_professor_report.py tạo tài liệu Word từ JSON và PNG đã lưu; dùng môi trường Python có python-docx. Các script dựng báo cáo không tự huấn luyện mô hình.')
table(['Bằng chứng','Đường dẫn tương đối trong dự án'],[
    ['Chia tập và baseline','models/surface_only/comparison.json'],
    ['Dự đoán giai đoạn 2','reports/rain_alert_experiment/test_predictions.parquet'],
    ['Giao thức và ngưỡng mới','reports/threshold_temporal_study/protocol.json\nreports/threshold_temporal_study/selected_thresholds.json'],
    ['Ba fold và metric','reports/threshold_temporal_study/fold_summary.json\nreports/threshold_temporal_study/temporal_metrics.csv'],
    ['Số liệu dựng báo cáo','reports/professor_report/report_data.json'],
    ['NOAA và audit thời gian thực','reports/independent_observations/noaa_audit_2025.json\nreports/realtime_readiness/availability_audit.json']
], [1.7,5.4],size=9.5)
h('Tài liệu kỹ thuật và dữ liệu')
refs=[
    '[1] Open-Meteo. Historical Weather API. https://open-meteo.com/en/docs/historical-weather-api',
    '[2] NOAA Climate Prediction Center. Oceanic Niño Index. https://www.cpc.ncep.noaa.gov/products/analysis_monitoring/ensostuff/ONI_v5.php',
    '[3] PyTorch. LSTM documentation. https://docs.pytorch.org/docs/stable/generated/torch.nn.LSTM.html',
    '[4] NOAA NCEI. Global Historical Climatology Network hourly. https://www.ncei.noaa.gov/products/global-historical-climatology-network-hourly',
    '[5] NOAA NCEI. Integrated Surface Data format. https://www.ncei.noaa.gov/pub/data/noaa/isd-format-document.pdf',
    '[6] Zhou và cộng sự. Informer Beyond Efficient Transformer for Long Sequence Time-Series Forecasting. https://arxiv.org/abs/2012.07436'
]
for reference in refs:
    para=p(reference)
    for run in para.runs:run.font.size=Pt(9)
p('Số liệu trong báo cáo lấy từ các artifact cục bộ đã lưu. Tài liệu bên ngoài dùng để mô tả nguồn và phương pháp, không thay thế số liệu thí nghiệm của dự án.')

doc.core_properties.title='Báo cáo tiến độ dự báo lượng mưa tại Đà Nẵng'
doc.core_properties.subject='Tổng hợp dữ liệu, Bi-LSTM và đánh giá theo thời gian cho báo cáo giảng viên'
doc.core_properties.author='Sinh viên thực hiện đề tài với hỗ trợ kỹ thuật từ Codex'
doc.core_properties.keywords='Đà Nẵng, lượng mưa, Bi-LSTM, báo cáo tiến độ'
path=OUT/'Bao_cao_tong_hop_PBL6_Da_Nang.docx'
doc.save(path)
print(path)

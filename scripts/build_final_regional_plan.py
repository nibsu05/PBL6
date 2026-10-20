"""Xuất Word/CSV cho phương án vùng trực tiếp; không huấn luyện mô hình."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
DEFAULT = ROOT / 'reports/regional_research_plan_2026_10_20'
CRITERIA = {
    'C1': 'Kiểm chứng nhãn với số đo thật',
    'C2': 'Đầu vào và độ trễ tại thời điểm dự báo',
    'C3': 'Cùng mẫu, baseline và khoảng tin cậy',
    'C4': 'Chỉ số phù hợp sự kiện mưa hiếm',
    'C5': 'Khả thi với RAM, GPU và thời gian',
    'C6': 'Tái sử dụng công việc đã làm',
}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def decision_appendix(decisions):
    blocks = []
    for row in decisions:
        blocks.extend([
            f"### {row['id']}. {row['topic']}", '',
            f"**Phương án nhóm:** {row['group']}", '',
            f"**Phương án nghiên cứu đã trao đổi:** {row['research']}", '',
            f"**Lựa chọn cuối:** {row['choice']}", '',
            '| Tiêu chí | Vì sao chọn |', '| --- | --- |',
            *[f'| {key} — {label} | {row[key]} |' for key, label in CRITERIA.items()], '',
        ])
    return '\n'.join(blocks)


def inline(paragraph, text):
    pattern = r'(\[[^\]]+\]\([^)]+\)|\*\*.*?\*\*|`[^`]+`)'
    for part in re.split(pattern, text):
        if not part:
            continue
        link = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
        if link:
            relation = paragraph.part.relate_to(link[2], RT.HYPERLINK, is_external=True)
            node = OxmlElement('w:hyperlink'); node.set(qn('r:id'), relation)
            run = OxmlElement('w:r'); prop = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '175A7A'); prop.append(color)
            underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); prop.append(underline)
            run.append(prop); value = OxmlElement('w:t'); value.text = link[1]
            run.append(value); node.append(run); paragraph._p.append(node)
        elif part.startswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            run = paragraph.add_run(part[1:-1]); run.font.name = 'Consolas'; run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def widths(headers):
    n = len(headers)
    if n == 2:
        return [4.5, 12.5]
    if n == 3 and headers[-1].startswith('Số cột'):
        return [4.1, 10.8, 2.1]
    if n == 3 and headers[0] == 'Mục gốc':
        return [2.3, 7., 7.7]
    if n == 4:
        return [6.2, 3.6, 3.6, 3.6]
    return [17. / n] * n


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Shading Accent 1'
    table.autofit = False
    for i, width in enumerate(widths(rows[0])):
        table.columns[i].width = Cm(width)
    for i, texts in enumerate(rows):
        cells = table.rows[0].cells if i == 0 else table.add_row().cells
        for j, value in enumerate(texts):
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[j].width = table.columns[j].width
            paragraph = cells[j].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.space_before = Pt(3)
            inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if i == 0:
                    run.bold = True
        table.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def word_copy(source):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.)
    section.header_distance = section.footer_distance = Cm(.8)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet']:
        style = doc.styles[name]
        style.font.name = 'Arial'; style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.1
    for name, size in [('Title', 24), ('Heading 1', 16), ('Heading 2', 12), ('Heading 3', 11)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string('16465C')
    section.header.paragraphs[0].text = 'PBL6  |  PHƯƠNG ÁN NGHIÊN CỨU  |  DỰ ĐOÁN VÙNG TRỰC TIẾP'
    section.header.paragraphs[0].runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('20/10/2026  •  Trang ').font.size = Pt(8)
    page = OxmlElement('w:fldSimple'); page.set(qn('w:instr'), 'PAGE'); footer._p.append(page)
    lines = source.read_text(encoding='utf-8').splitlines(); i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip() or line.startswith('<!--'):
            i += 1; continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                row = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+', c) for c in row):
                    rows.append(row)
                i += 1
            add_table(doc, rows); continue
        if line.startswith('# '):
            paragraph = doc.add_paragraph(style='Title'); value = line[2:]
        elif line.startswith('## '):
            if 'Phụ lục A.' in line:
                doc.add_page_break()
            paragraph = doc.add_paragraph(style='Heading 1'); value = line[3:]
        elif line.startswith('### '):
            paragraph = doc.add_paragraph(style='Heading 2'); value = line[4:]
        elif line.startswith('- '):
            paragraph = doc.add_paragraph(style='List Bullet'); value = line[2:]
        else:
            paragraph = doc.add_paragraph(); value = line
        inline(paragraph, value)
        i += 1
    doc.core_properties.title = 'Phương án nghiên cứu dự đoán mưa vùng trực tiếp tại Đà Nẵng'
    doc.core_properties.author = 'Nhóm PBL6'
    doc.core_properties.subject = 'Hợp nhất phương án và quyết định theo sáu tiêu chí'
    target = source.with_suffix('.docx'); doc.save(target)
    return target


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--report', type=Path, default=DEFAULT)
    parser.add_argument('--source', type=Path, default=Path.home()/'Downloads/TONG_HOP_DE_XUAT.md')
    args = parser.parse_args()
    out = args.report
    decisions = json.loads((out/'decisions.json').read_text(encoding='utf-8'))
    assert len(decisions) == 18 and len({d['id'] for d in decisions}) == 18
    for d in decisions:
        assert all(isinstance(d[c], str) and len(d[c]) > 30 for c in CRITERIA)
    groups = json.loads((out/'feature_groups.json').read_text(encoding='utf-8'))
    all_features = [f for g in ['B', 'A', 'M', 'C'] for f in groups[g]]
    assert len(all_features) == len(set(all_features)) == 44
    catalog_path = out/'feature_catalog_reference.csv'
    catalog = list(csv.DictReader(catalog_path.read_text(encoding='utf-8').splitlines()))
    v1_sources = {r['source'] for r in catalog if r['version'] == 'V1'}
    assert set(all_features).issubset(v1_sources)
    expected = {'E0': 11, 'E1': 23, 'E2': 31, 'E3': 12, 'E4': 43, 'E5': 24, 'E6': 32, 'E7': 44}
    for eid, selected_groups in groups['experiments'].items():
        assert sum(len(groups[g]) for g in selected_groups) == expected[eid]
    plan = out/'PHUONG_AN_CUOI.md'
    content = plan.read_text(encoding='utf-8')
    for protected in ['D1', 'D2', 'D9', 'D10', 'D14', 'D15']:
        assert protected in content
    assert 'CSI/BSS là chỉ số chính' in content
    start, stop = '<!-- DECISIONS_START -->', '<!-- DECISIONS_END -->'
    assert content.count(start) == content.count(stop) == 1
    before, rest = content.split(start); _, after = rest.split(stop)
    content = before + start + '\n\n' + decision_appendix(decisions) + '\n' + stop + after
    plan.write_text(content, encoding='utf-8', newline='\n')
    decision_csv = out/'QUYET_DINH_6_TIEU_CHI.csv'
    with decision_csv.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=['id', 'topic', 'group', 'research', 'choice', *CRITERIA],
                                lineterminator='\n')
        writer.writeheader(); writer.writerows(decisions)
    word = word_copy(plan)
    checked = Document(word)
    assert len(checked.tables) >= len(decisions)
    for d in decisions:
        assert d['id'] in content
    manifest = {
        'document_date': '2026-10-20', 'status': 'agreed_objective_and_planned_protocol_not_training_results',
        'objective': 'Direct regional mean rainfall prediction; research first',
        'decisions': len(decisions), 'criteria_per_decision': list(CRITERIA),
        'protected_decisions_retained': ['D1', 'D2', 'D9', 'D10', 'D14', 'D15'],
        'same_protocol_superior_evidence_used_to_replace_protected': False,
        'primary_feature_counts': expected,
        'reference_catalog_commit': 'd3f39360b8ab28e88be4b56310fbae10db8b9db7',
        'word_tables': len(checked.tables), 'word_paragraphs': len(checked.paragraphs),
        'models_trained_by_this_build': False, 'test_2025_scored_by_this_build': False,
        'files': {p.name: {'sha256': sha256(p), 'bytes': p.stat().st_size}
                  for p in [plan, word, decision_csv, out/'decisions.json', out/'feature_groups.json', catalog_path]},
    }
    if args.source.exists():
        manifest['group_source'] = {'filename': args.source.name, 'sha256': sha256(args.source),
                                    'lines': len(args.source.read_text(encoding='utf-8').splitlines()),
                                    'referenced_commit': 'e531982'}
    (out/'build_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2),
                                        encoding='utf-8', newline='\n')
    print(json.dumps({'word': str(word), 'decisions': len(decisions), 'criteria_each': len(CRITERIA),
                      'tables': len(checked.tables), 'feature_counts': expected}, ensure_ascii=False))


if __name__ == '__main__':
    main()

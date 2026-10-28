"""Prepare a two-page brief and five-slide talk from existing audited results."""
from pathlib import Path
import hashlib
import json
import shutil
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import pandas as pd
from docx import Document
from docx.shared import Cm,Pt,Inches as WordInches
from pptx import Presentation
from pptx.util import Inches,Pt as SlidePt
from pptx.dml.color import RGBColor
from build_final_regional_plan import inline,add_table

SOURCE=ROOT/'reports/short_range_ablation_2026_10_21'
OUT=ROOT/'reports/short_range_brief_2026_10_21'


def write(name,text):
    (OUT/name).write_text(text,encoding='utf-8',newline='\n')


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    run=json.loads((SOURCE/'run.json').read_text(encoding='utf-8'))
    audit=json.loads((SOURCE/'checkpoint_audit.json').read_text(encoding='utf-8'))
    assert run['status']=='complete' and audit['status']=='passed' and audit['full_runs_replayed']==96
    scores=pd.read_csv(SOURCE/'scores_ensemble_pooled.csv')
    scores=scores[(scores.season=='all')&(scores.threshold_mm==16)]
    ci=pd.read_csv(SOURCE/'paired_bootstrap.csv')
    def score(model,exp,lead):
        return scores[(scores.model==model)&(scores.experiment==exp)&(scores.lead_day==lead)].iloc[0]
    def effect(exp):
        return ci[(ci.model=='tft')&(ci.contrast==exp+'-E0')&(ci.lead_day==1)&(ci.metric=='BSS16')&(ci.block_days==7)].iloc[0]
    table_rows=[]
    for lead in [1,2,3]:
        table_rows.append([str(lead),*[f'{score(m,e,lead).BSS:.3f}' for m,e in [('tft','E0'),('tft','E7'),('lgbm','E0'),('lgbm','E7')]]])
    feature_rows=[]
    for exp,label in [('E1','A: khí quyển/đối lưu/mây/đất'),('E2','M: hoàn lưu và vận chuyển ẩm'),('E3','C: SST vùng biển')]:
        r=effect(exp)
        feature_rows.append([label,f'{r.estimate:+.3f}',f'[{r.family_adjusted_low:+.3f}; {r.family_adjusted_high:+.3f}]'])
    def mdtable(headers,rows):
        return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |',
            *['| '+' | '.join(row)+' |' for row in rows]])
    summary=(f'TFT ngày 1 có BSS từ {score("tft","E0",1).BSS:.3f} ở E0 lên {score("tft","E7",1).BSS:.3f} ở E7. '
        'Thêm riêng nhóm A hoặc M có bằng chứng cải thiện ngày 1 sau hiệu chỉnh chín phép so. '
        'Chưa có bằng chứng tương tự cho SST hoặc ngày 2–3. Đây là kết quả phát triển hồi cứu, chưa phải kiểm định cuối.')
    amount=(f'RMSE ngày 1 E7: TFT {score("tft","E7",1).RMSE:.2f}, LightGBM {score("lgbm","E7",1).RMSE:.2f} mm/24h. '
        'TFT chưa thể hiện ưu thế tổng thể; phép so này còn khác biểu diễn lịch sử và loss.')
    page1=f'''# Báo cáo nhanh: dự đoán mưa vùng 24–72 giờ

Ngày 21/10/2026 · Dùng để báo cáo tiến độ nghiên cứu · Test 2025 giữ kín

{summary}

## 1. Bài toán và cách làm

Mỗi ngày lúc 07 giờ Việt Nam, dùng 168 giờ khí tượng quá khứ gộp thành 28 bước 6 giờ để dự đoán trực tiếp ba tổng mưa trung bình vùng: [t,t+24), [t+24,t+48), [t+48,t+72). Đơn vị mm/24h; ngày 2/3 là hai tổng riêng. Nhãn GSMaP Gauge; đầu vào khí tượng ERA5. Decoder chỉ biết lịch.

Vùng hiện là 93 ô xấp xỉ, chưa phải giao ranh hành chính chính xác. E0 có 11 biến nền; E7 có 44 biến. Thí nghiệm đủ E0–E7 trên ba fold evaluation 2021–2023, ba seed TFT: 72 lượt TFT và 24 pipeline LightGBM. Mỗi lead có 1.075 mốc đánh giá, 114 sự kiện ≥16 mm; thêm seed không tạo thêm ngày thời tiết độc lập.

## 2. Kết quả cần nhớ

{mdtable(['Ngày','TFT E0','TFT E7','LGBM E0','LGBM E7'],table_rows)}

Bảng là BSS cho xác suất mưa vùng ≥16 mm/24h, so với xác suất khí hậu học từ train: dương có lợi, âm kém hơn tham chiếu. TFT là trung bình dự báo ba seed; LightGBM là một pipeline xác định.

{amount}

Ngày 1 có tín hiệu lặp lại: E7−E0 dương ở cả ba seed TFT, ΔBSS từ +0,068 đến +0,153. Ngày 3, BSS âm ở tất cả cấu hình hiện đã chạy của cả hai họ mô hình. Đây là điểm yếu cần báo cáo.
'''
    page2=f'''## 3. Nhóm đặc trưng nào có tín hiệu?

{mdtable(['Nhóm thêm riêng vào E0','ΔBSS ngày 1','CI99,44% đã hiệu chỉnh'],feature_rows)}

Bootstrap paired 20.000 lần, khối 7 ngày; gia đình chín phép so gồm ba nhóm × ba lead. Hai CI trên 0 thuộc A/M ngày 1; kiểm tra khối 14 ngày cho cùng kết luận. Các CI điều kiện trên mô hình đã fit, chưa bao gồm mọi bất định của nhãn và huấn luyện. Những năm development đã được xem trước.

M có mức tăng lớn hơn khi thêm riêng vào nền; chưa kiểm định trực tiếp M−A. SST hiện chỉ là một biến vùng biển, chưa đại diện đầy đủ ENSO/MJO. Không suy ra yếu tố khí hậu biển không quan trọng từ phép thử này.

## 4. Nhận xét độc lập và bước tiếp theo

- Thêm hết biến chưa chắc tốt nhất: LightGBM ngày 1 có BSS E4=0,315, E7=0,303. Đây là chênh lệch điểm mô tả, chưa kiểm định E4−E7.
- Ngày 2 có kỹ năng xác suất thấp, ngày 3 chưa vượt khí hậu. Chưa thể kết luận nguyên nhân duy nhất là kiến trúc hoặc thiếu forecast.
- Cần chốt ranh/trọng số 93 ô và kiểm chứng nhãn ngày. Nếu thay dữ liệu, chạy lại phần phát triển liên quan; sau khóa thiết kế mới đánh giá năm 2025.
- ERA5 hồi cứu chưa bảo đảm sẵn có lúc phát hành. Chưa có kết luận vận hành, mưa cực đoan tại điểm hay giá trị trên toàn địa giới Đà Nẵng mới.

Kiểm chứng kỹ thuật: 27 kiểm tra đạt; 96/96 lượt tái lập cả calibration/evaluation. Những kiểm tra này xác nhận pipeline, chưa xác nhận nhãn vệ tinh đúng với mưa thực. Informer chưa được chạy trong phép so, nên chưa có kết luận về tính khả thi của Informer.

Nguồn số liệu: scores_ensemble_pooled.csv, paired_bootstrap.csv, checkpoint_audit.json trong thư mục báo cáo mở rộng. Báo cáo đầy đủ, 8 slide và danh mục biến vẫn được giữ nguyên.
'''
    write('BAO_CAO_NHANH.md',page1+'\n'+page2)
    speech=f'''# Lời trình bày 5 phút

Ngày 21/10/2026. Dùng cùng PBL6_TFT_5_PHUT.pptx. Các con số lấy từ ma trận phát triển đã kiểm chứng, không phải test 2025.

## Slide 1 · 0:00–0:55 · Mục tiêu và dữ liệu

“Phần em phụ trách là dự đoán lượng mưa trung bình vùng trong ba ngày tới. Mỗi mẫu đại diện cho một thời điểm của cả vùng, không phải một ô lưới. Nhóm tổng hợp khí tượng và nhãn từ 93 ô trước khi đưa vào mô hình. Hiện vùng này còn xấp xỉ, chưa khẳng định khớp toàn địa giới Đà Nẵng mới.

Đầu vào là bảy ngày khí tượng ERA5 quá khứ, gộp sáu giờ; nhãn là mưa GSMaP Gauge. Đầu ra gồm ba tổng 24 giờ riêng. Mô hình chỉ biết lịch tương lai, không nhận ERA5 của giờ cần dự báo.”

## Slide 2 · 0:55–1:55 · Đã chạy và kết quả tổng thể

“Nhóm đã chạy đủ tám tổ hợp đặc trưng, ba fold phát triển 2021–2023 và ba seed TFT, gồm 72 lượt TFT cùng 24 pipeline LightGBM. Mọi cấu hình được chấm trên cùng các thời điểm hợp lệ.

BSS đánh giá xác suất sự kiện mưa vùng từ 16 mm trong 24 giờ so với khí hậu tham chiếu. Ngày 1, TFT tăng từ 0,135 ở nền E0 lên 0,246 khi thêm toàn bộ biến E7. Tuy nhiên LightGBM E7 đạt 0,303 và có RMSE ngày 1 thấp hơn. Vì vậy em chưa kết luận TFT tốt hơn mô hình đối chứng.

Ngày 3, BSS đều âm ở các cấu hình đã chạy. Đây là kết quả cần ghi nhận, chưa cho phép kết luận chung rằng dự báo 72 giờ không khả thi.”

## Slide 3 · 1:55–3:00 · Đóng góp từng nhóm biến

“Mục tiêu chính là nghiên cứu đặc trưng nên nhóm tách A, M và C. A gồm khí quyển, đối lưu, mây và ẩm đất; M gồm gió, ẩm và vận chuyển ẩm; C hiện chỉ có SST vùng biển.

Khi thêm riêng vào nền E0, ngày 1 có ΔBSS khoảng +0,085 với A và +0,144 với M. Hai khoảng tin cậy sau hiệu chỉnh chín phép so vẫn trên 0, kể cả khi đổi khối bootstrap từ bảy sang mười bốn ngày. Chưa kiểm định trực tiếp M tốt hơn A.

Chưa có bằng chứng tương tự cho SST hoặc ngày 2–3. Kết luận về C chỉ áp dụng cho biến SST này, chưa phải kết luận về ENSO/MJO hay toàn bộ yếu tố khí hậu.”

## Slide 4 · 3:00–4:00 · Ý nghĩa và giới hạn

“Tín hiệu tích cực là lợi ích ngày 1 lặp lại qua ba seed và có đối chứng. Nhưng thêm hết biến không mặc nhiên tốt nhất: LightGBM E4 có điểm BSS cao hơn E7 ở ngày 1, còn chênh lệch đó chưa được kiểm định riêng.

Kết quả đang là development đã được xem trước, với nhãn vệ tinh và vùng xấp xỉ. Nhiều seed không tạo thêm ngày mưa độc lập; checksum hoặc replay đúng cũng không chứng minh nhãn đúng với mưa thật. Các yếu tố này giới hạn phạm vi kết luận.”

## Slide 5 · 4:00–5:00 · Việc tiếp theo

“Ưu tiên tiếp theo là chốt ranh và trọng số vùng, kiểm chứng nhãn ngày bằng đối chứng phù hợp. Nếu thay dữ liệu, nhóm dựng lại dữ liệu và chạy lại phần phát triển liên quan. Sau đó mới khóa thiết kế và đánh giá năm 2025 đang giữ kín.

Nếu cần cải thiện ngày 3, nhóm cân nhắc một nhánh dùng forecast thật đã được phát hành, với đối chứng cùng thông tin và thời điểm sẵn có. Chưa có phép đo Informer trong ma trận này nên em chưa kết luận nó không khả thi. Hướng chính vẫn là nghiên cứu đóng góp đặc trưng cho dự đoán vùng trực tiếp.”

## Ba câu để chốt khi thời gian ngắn

- Đã có pipeline và ma trận phát triển tái lập, với 72 TFT và 24 LightGBM.
- A/M giúp ngày 1 trong phép thử này; SST/ngày 2–3 chưa có bằng chứng tương tự.
- Còn ranh/nhãn và kiểm định độc lập; năm 2025 vẫn giữ kín.
'''
    write('LOI_DAN_5_PHUT.md',speech)
    doc=Document();section=doc.sections[0]
    section.page_width,section.page_height=Cm(21),Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(1.5)
    section.left_margin=section.right_margin=Cm(1.7)
    for name in ['Normal','Title','Heading 1','Heading 2','List Bullet']:
        style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(10)
        style.paragraph_format.space_after=Pt(4);style.paragraph_format.line_spacing=1.05
    doc.styles['Title'].font.size=Pt(18);doc.styles['Heading 1'].font.size=Pt(12)
    section.header.paragraphs[0].text='PBL6 | BÁO CÁO TIẾN ĐỘ TFT VÙNG | 21/10/2026'
    section.header.paragraphs[0].runs[0].font.size=Pt(8)
    section.footer.paragraphs[0].text='Development 2021–2023 · Footprint 93 ô xấp xỉ · Test 2025 giữ kín'
    section.footer.paragraphs[0].runs[0].font.size=Pt(8)
    def add_markdown(text):
        lines=text.splitlines();i=0
        while i<len(lines):
            line=lines[i]
            if not line.strip():i+=1;continue
            if line.startswith('|'):
                rows=[]
                while i<len(lines) and lines[i].startswith('|'):
                    cells=[c.strip() for c in lines[i].strip('|').split('|')]
                    if not all(set(c)<={'-',' '} for c in cells):rows.append(cells)
                    i+=1
                add_table(doc,rows);continue
            style='Title' if line.startswith('# ') else 'Heading 1' if line.startswith('## ') else 'List Bullet' if line.startswith('- ') else 'Normal'
            value=line[2:] if line.startswith(('# ','- ')) else line[3:] if line.startswith('## ') else line
            inline(doc.add_paragraph(style=style),value);i+=1
    add_markdown(page1);doc.add_page_break();add_markdown(page2)
    doc.core_properties.title='Báo cáo nhanh TFT mưa vùng 24–72 giờ'
    doc.core_properties.author='Nhóm PBL6'
    doc.save(OUT/'BAO_CAO_NHANH.docx')
    shutil.copyfile(SOURCE/'primary_family_ci.png',OUT/'primary_family_ci.png')
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    def text(slide,x,y,w,h,value,size=23,bold=False,color='123047'):
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame
        tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        lines=value.split('\n')
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial'
            p.font.size=SlidePt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
            p.space_after=SlidePt(12 if i+1<len(lines) else 0)
    notes=speech.split('## Slide ')[1:]
    def base(title,index):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid()
        slide.background.fill.fore_color.rgb=RGBColor(248,250,252)
        text(slide,.55,.3,12.2,.8,title,29,True)
        text(slide,.55,7.03,12.2,.25,'PBL6 • 21/10/2026 • vùng 93 ô xấp xỉ • development 2021–2023 • test 2025 giữ kín',11,color='516473')
        slide.notes_slide.notes_text_frame.text=notes[index].split('## Ba câu')[0].strip()
        return slide
    slide=base('Dự đoán mưa vùng: ngày 1, ngày 2, ngày 3',0)
    text(slide,.7,1.5,12,1.,'7 ngày khí tượng quá khứ → 3 tổng mưa 24 giờ riêng',27,True,'087E8B')
    text(slide,.7,2.9,12,2.4,'Đầu vào: ERA5 quá khứ, gộp 6 giờ.\nNhãn: GSMaP Gauge trung bình vùng 93 ô.\nCùng mẫu, E0–E7, ba fold; 72 TFT + 24 LightGBM.',25)
    text(slide,.7,5.8,12,.8,'Mục tiêu: đo đóng góp đặc trưng; footprint hiện còn xấp xỉ.',21,color='516473')
    slide=base('Ngày 1 có tín hiệu; ngày 3 còn yếu',1)
    shape=slide.shapes.add_table(4,5,Inches(.7),Inches(1.65),Inches(12),Inches(2.65))
    rows=[['Ngày','TFT E0','TFT E7','LGBM E0','LGBM E7']]+table_rows
    for ri,row in enumerate(rows):
        for col,value in enumerate(row):
            cell=shape.table.cell(ri,col);cell.text=value
            for p in cell.text_frame.paragraphs:
                p.font.name='Arial';p.font.size=SlidePt(22);p.font.bold=ri==0
    text(slide,.7,4.7,12,.8,'BSS ≥16 mm/24h so với khí hậu: dương có lợi, âm kém hơn.',21)
    text(slide,.7,5.8,12,.7,f'RMSE ngày 1 E7: TFT {score("tft","E7",1).RMSE:.2f}; LightGBM {score("lgbm","E7",1).RMSE:.2f} mm/24h.',22,True,'087E8B')
    slide=base('A và M cải thiện ngày 1 trong phép thử này',2)
    slide.shapes.add_picture(str(OUT/'primary_family_ci.png'),Inches(.5),Inches(1.45),width=Inches(12.3))
    text(slide,.7,6.45,12,.35,'20.000 bootstrap • 2/9 CI điều chỉnh trên 0 • khối 14 ngày cho cùng kết luận',17,color='516473')
    slide=base('Nhận xét độc lập và giới hạn',3)
    text(slide,.7,1.45,12,4.8,'Ngày 1: lợi ích E7−E0 lặp lại ở cả ba seed TFT.\nThêm hết biến chưa chắc tốt nhất; TFT chưa có ưu thế tổng thể.\nSST chỉ là một biến vùng, chưa đại diện đầy đủ ENSO/MJO.\nVùng 93 ô và nhãn còn cần kiểm chứng.\nDevelopment đã xem trước; chưa có kết luận holdout hoặc vận hành.',24)
    slide=base('Việc tiếp theo để hoàn thiện nghiên cứu',4)
    text(slide,.7,1.5,12,4.7,'1. Chốt ranh, trọng số và độ phủ vùng.\n2. Kiểm chứng nhãn mưa ngày bằng đối chứng phù hợp.\n3. Nếu thay dữ liệu: chạy lại development liên quan.\n4. Khóa thiết kế, sau đó mới đánh giá năm 2025.\n5. Xem forecast thật như nhánh riêng cho ngày 3.',25)
    prs.save(OUT/'PBL6_TFT_5_PHUT.pptx')
    numbers=dict(source='short_range_ablation_2026_10_21',test2025_read=False,
        development_runs=dict(tft=72,lgbm=24,replayed=96),bss_headers=rows[0],bss_rows=table_rows,
        feature_effect_headers=['group','delta_BSS','adjusted_CI'],feature_effect_rows=feature_rows,
        source_hashes={name:hashlib.sha256((SOURCE/name).read_bytes()).hexdigest() for name in
                       ['scores_ensemble_pooled.csv','paired_bootstrap.csv','checkpoint_audit.json','run.json']})
    write('brief_numbers.json',json.dumps(numbers,ensure_ascii=False,indent=2))
    write('README.md','''# Bộ báo cáo nhanh trong 30 phút

- [Báo cáo Word hai trang thiết kế](BAO_CAO_NHANH.docx), [bản Markdown](BAO_CAO_NHANH.md).
- [5 slide](PBL6_TFT_5_PHUT.pptx), có lời dẫn trong notes.
- [Lời trình bày 5 phút](LOI_DAN_5_PHUT.md).
- [Bộ kết quả đầy đủ và 8 slide](../short_range_ablation_2026_10_21/README.md).
- [PR tích hợp vào main](https://github.com/nibsu05/PBL6/pull/1).

Gợi ý chuẩn bị: 0–5 phút đọc bản ngắn; 5–15 phút xem năm slide và bảng BSS; 15–25 phút tập lời dẫn; 25–30 phút kiểm tra ba thông điệp và các giới hạn. Các tài liệu đã có số liệu thực, không cần chờ thêm lượt huấn luyện.

Kết luận: A/M giúp TFT ngày 1 trong phép thử phát triển; SST/ngày 2–3 chưa có bằng chứng tương tự. Footprint 93 ô xấp xỉ, nhãn cần kiểm chứng, năm 2025 giữ kín. Bộ này chỉ tổng hợp kết quả đã có, không fit/chấm mô hình mới hoặc thay đổi dữ liệu.
''')
    files=[p for p in OUT.iterdir() if p.is_file() and p.name!='build_manifest.json']
    manifest=dict(status='built_from_audited_results',date_local='2026-10-21',slides=5,
        models_fitted_by_this_build=0,test2025_read=False,word_explicit_page_breaks=1,
        source_hashes=numbers['source_hashes'],code_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        files={p.name:dict(bytes=p.stat().st_size,sha256=hashlib.sha256(p.read_bytes()).hexdigest()) for p in files})
    write('build_manifest.json',json.dumps(manifest,ensure_ascii=False,indent=2))
    print('BRIEF_COMPLETE',len(files),'files; 5slides; no new model training')


if __name__=='__main__':main()

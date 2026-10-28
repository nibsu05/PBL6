"""Complete the three-model teaching report from the verified Informer matrix."""
from pathlib import Path
import copy
import hashlib
import json
import sys

ROOT = Path(__file__).resolve().parents[1]; sys.path.insert(0, str(ROOT))
import pandas as pd
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt
from presentation_layout import Deck, TEAL, NAVY, RED
from build_final_regional_plan import word_copy
from build_short_range_three_models import METRICS, HEADERS, MODELS, LABELS, markdown_table

OUT = ROOT / 'reports/short_range_complete_2026_10_28'
OLD = ROOT / 'reports/short_range_three_models_2026_10_21'
MATRIX = ROOT / 'reports/informer_regional_matrix_2026_10_28'
STEM = 'PBL6_SO_SANH_3_MO_HINH_DAY_DU'


def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    run = json.loads((MATRIX/'run.json').read_text(encoding='utf8'))
    validation = json.loads((MATRIX/'validation.json').read_text(encoding='utf8'))
    assert run['status']=='complete' and len(run['models'])==72 and validation['status']=='passed'
    full = pd.read_csv(MATRIX/'scores_ensemble_pooled.csv')
    scores = full[(full.season=='all') & (full.threshold_mm==16)].copy()
    assert len(scores)==72 and set(scores.model)==set(MODELS)
    assert set(scores.n)=={1075} and set(scores.events)=={114}
    paired = pd.read_csv(MATRIX/'paired_bootstrap_exploratory.csv')
    absolute = pd.read_csv(MATRIX/'absolute_skill_ci.csv')
    scene = json.loads((OLD/'slide_scene.json').read_text(encoding='utf8'))
    deck = Deck(); deck.slides = copy.deepcopy(scene)
    def find(model, e, day):
        r=scores[(scores.model==model)&(scores.experiment==e)&(scores.lead_day==day)]
        assert len(r)==1
        return r.iloc[0]
    def style(value, winner, key):
        return dict(bold=bool(winner),color=RED if key=='BSS' and value<0 else TEAL if winner else NAVY)
    def values_and_styles(e, day):
        rows, styles=[],{}
        for i, model in enumerate(MODELS,1):
            r=find(model,e,day); rows.append([LABELS[model],*[f'{r[k]:.3f}' for k in METRICS]])
            for col,key in enumerate(METRICS,1):
                values=[float(find(m,e,day)[key]) for m in MODELS]
                if key=='bias': values=[abs(v) for v in values]
                goal=min(values) if key in ['RMSE','MAE','bias','BS','FAR'] else max(values)
                v=abs(float(r[key])) if key=='bias' else float(r[key])
                styles[f'{i},{col}']=style(float(r[key]),round(v,3)==round(goal,3),key)
        return rows,styles
    replacements={
        'Informer: 18 lượt E0/E7':'Informer: 72 lượt, đủ E0–E7',
        'Phần còn thiếu':'Cần kiểm chứng',
        'Informer chưa có E1–E6\nKhông điền số giả hoặc số 0\nChưa mở test 2025':'Ranh vùng và chất lượng nhãn\nĐộ bền của các kết quả\nChưa mở test 2025',
        'Đã chạy đến đâu? TFT/LightGBM đủ, Informer còn thiếu':'Đã hoàn thành E0–E7 cho cả ba mô hình',
        'Đã có: 72 TFT + 24 pipeline LightGBM + 18 Informer = 114 lượt.':'Đã có: 72 TFT + 24 pipeline LightGBM + 72 Informer = 168 lượt.',
        'Informer mới hai E: 2 × 9 = **18 lượt**.':'Informer đủ tám E: 8 × 9 = **72 lượt**.',
        'Đối chứng cần vượt':'Mô hình tham chiếu',
        'Mới E0/E7':'Đã đủ E0–E7',
        'Đậm: tốt nhất trong hàng đủ ba mô hình · Đỏ: BSS âm · Informer E1–E6 chưa chạy.':'Đậm: tốt nhất trong hàng cùng E/ngày · Đỏ: BSS âm · Cả ba đã đủ E0–E7.',
        '— = Informer chưa chạy E1–E6 · Đậm chỉ khi đủ ba mô hình cùng E/ngày.':'Đậm: tốt nhất cùng E/ngày · Đỏ: BSS âm · So sánh đủ ba mô hình.',
        'TFT/LightGBM: đủ E0–E7\nInformer: E0/E7, khoảng 5,3 phút\n114 lượt đã được kiểm chứng\nCó cả kết quả cải thiện và kết quả âm':'Cả ba đã đủ E0–E7\nInformer: 54 mới + 18 dùng lại\n168 lượt mô hình/pipeline\nCó cả kết quả cải thiện và âm',
        'Còn thiếu':'Việc tiếp theo',
        'Informer E1–E6: 54 lượt\nKiểm định chênh lệch ba pipeline\nRanh vùng và kiểm chứng nhãn\nChốt thiết kế rồi mới chấm 2025':'Ranh vùng và kiểm chứng nhãn\nĐộ bền theo mùa/năm\nKiểm tra bước gộp sáu giờ\nKhóa thiết kế rồi mới chấm 2025',
        'Lượt này cập nhật cách trình bày từ kết quả có sẵn; không huấn luyện mới.':'Đã bổ sung 54 lượt Informer; đủ ma trận ba mô hình, chưa phải nghiên cứu cuối.',
        '4. Informer mới E0/E7; chưa có ma trận ba mô hình hoàn chỉnh.':'4. Đã đủ E0–E7; còn ranh/QC nhãn và kiểm định cuối.',
        '3. Ngày 3: cả ba chưa vượt khí hậu về dự báo xác suất.':'3. E7 ngày 3: cả ba BSS âm; Informer E6 chỉ đạt +0,004.',
        'Phụ lục: đủ ba mô hình trong mỗi E/ngày; ô thiếu ghi rõ chưa chạy.':'Phụ lục: đủ 72 dòng chính = tám E × ba mô hình × ba ngày.'}
    for slide in deck.slides:
        for before,after in replacements.items(): slide['title']=slide['title'].replace(before,after)
        for sh in slide['shapes']:
            if sh['kind']=='text':
                for before,after in replacements.items():sh['text']=sh['text'].replace(before,after)
            if sh['kind']!='table':continue
            tag=sh.get('tag','')
            if tag=='coverage_three_models':
                for row in sh['rows']:row[-1]='9'
            elif tag.startswith('bss_3models_'):
                day=int(tag.rsplit('_',1)[1]);sh['styles']={}
                for i in range(8):
                    e=f'E{i}'; vals=[float(find(m,e,day).BSS) for m in MODELS]
                    sh['rows'][i]=[e,*[f'{v:.3f}' for v in vals]]
                    for col,v in enumerate(vals,1):sh['styles'][f'{i+1},{col}']=style(v,round(v,3)==round(max(vals),3),'BSS')
            elif tag.startswith('e7_3models_'):
                day=int(tag.rsplit('_',1)[1]);sh['rows'],sh['styles']=values_and_styles('E7',day)
            elif tag.startswith('full_3models_'):
                day=int(tag.split('_')[-2]);sh['styles']={}
                for first in range(0,len(sh['rows']),3):
                    e=sh['rows'][first][0]; rows,styles=values_and_styles(e,day)
                    sh['rows'][first:first+3]=[[e,*r] for r in rows]
                    for key,v in styles.items():
                        ri,ci=map(int,key.split(','));sh['styles'][f'{first+ri},{ci+1}']=v
        # Replace notes whose earlier purpose was to disclose missing Informer runs.
        if slide['number']==1:slide['note']='Đã hoàn thành 72 TFT, 24 pipeline LightGBM và 72 Informer. Informer gồm 18 lượt E0/E7 cũ giữ nguyên và 54 lượt E1–E6 mới, cùng giao thức. Tất cả vẫn là đánh giá phát triển trên footprint xấp xỉ; chưa mở test 2025.'
        if slide['number']==6:slide['note']='Mỗi E có ba pipeline LightGBM xác định, chín TFT và chín Informer = ba năm nhân ba seed cho mỗi mạng. Đủ 168 lượt mô hình/pipeline; không có nghĩa số ngày thời tiết tăng theo seed. Độc lập phát lại đủ 72 checkpoint Informer, giữ nguyên 96 lượt TFT/LightGBM đã kiểm chứng.'
        if slide['number'] in [24,25,26]:slide['note']='Mỗi hàng đặt đủ LightGBM/TFT/Informer trong cùng E/ngày. BSS dương hơn khí hậu, âm kém khí hậu. Đậm chỉ là điểm cao nhất trong hàng, không tự là ý nghĩa thống kê. Các khoảng tin cậy mới là mô tả phát triển, chưa kiểm định cuối hoặc hiệu chỉnh nhiều phép.'
        if slide['number']==30:slide['note']='Giữ nguyên phép thử TFT chính đã đăng ký: A/M ngày 1 có CI điều chỉnh gia đình chín phép trên 0, SST chưa có. Informer giờ đã đủ E1–E6; bảng ba mô hình bổ sung dùng CI95% mô tả, không thay phép thử TFT chính hoặc biến phân tích khám phá thành holdout.'
        if slide['number']==31:slide['note']='Đã bổ sung đủ 54 lượt Informer và dùng lại 18 lượt E0/E7. Có đủ ma trận ba mô hình nhưng chưa hoàn tất giao ranh, QC nhãn hoặc kiểm tra 1h/3h/6h. Các CI mới có điều kiện trên mô hình đã fit; không bao trùm toàn bộ bất định nhãn/huấn luyện. Chưa mở test 2025.'
        if slide['number']==32:slide['note']='Giữ kết luận trung thực trên E7: LightGBM tốt hơn theo RMSE/BSS/FAR ngày 1, Informer CSI cao nhất, TFT POD cao nhất. Cả ba E7 ngày 3 BSS âm. Đã đủ E0–E7 của cả ba; xem ma trận để biết thêm biến có lợi ở đâu, không chọn mô hình chỉ từ một điểm đẹp hoặc kết luận vượt holdout.'
        for sh in slide['shapes']:
            if sh.get('tag','').startswith('full_3models_'):
                slide['note']='Bảng đủ LightGBM/TFT/Informer cho từng E/ngày. Đậm xét trong cùng E và độ chính xác hiển thị; Bias gần 0 tốt hơn, các metric khác theo mũi tên. Không có ô Informer thiếu. Toàn năm ở bảng này, chi tiết mùa/năm/ngưỡng trong Excel/CSV. Đậm chưa chứng minh khác biệt có ý nghĩa thống kê.'

    added=Deck()
    s=added.slide('E7: BSS và khoảng tin cậy của cả ba mô hình', 'BẤT ĐỊNH · PHÂN TÍCH KHÁM PHÁ',
        'CI95% mô tả • 20.000 bootstrap khối bảy ngày • Có độ nhạy 14 ngày trong CSV • Chưa mở 2025',
        'Mỗi ô là điểm BSS và CI95% từ lấy mẫu lại các đợt ngày trong từng năm. Khoảng hoàn toàn trên 0 hỗ trợ kỹ năng dương trong phân tích mô tả; chứa 0 chưa phân biệt rõ với khí hậu. Không dùng việc hai CI giao nhau để thay CI chênh lệch. Các CI chênh lệch theo cặp của ba mô hình được xuất riêng; chưa hiệu chỉnh nhiều phép hoặc bao trùm bất định nhãn/huấn luyện.')
    rows=[]
    for model in MODELS:
        cells=[LABELS[model]]
        for day in [1,2,3]:
            q=absolute[(absolute.model==model)&(absolute.lead_day==day)&(absolute.metric=='BSS16')&(absolute.block_days==7)].iloc[0]
            cells.append(f'{q.estimate:+.3f}\n[{q.ci_low:+.3f}; {q.ci_high:+.3f}]')
        rows.append(cells)
    added.table(s,['Mô hình','Ngày 1','Ngày 2','Ngày 3'],rows,y=1.60,row_h=.90,widths=[1.75,3.5,3.5,3.5],size=20,tag='e7_skill_ci')
    added.text(s,.70,5.45,11.95,.5,'Khoảng chứa 0: chưa phân biệt rõ với khí hậu trong phân tích này.',20)
    added.band(s,'CI mới là mô tả phát triển; chưa hiệu chỉnh nhiều phép và chưa phải kiểm định cuối.',size=18)
    s=added.slide('Thêm riêng A/M/C: đối chiếu đủ ba mô hình', 'BẤT ĐỊNH · PHÂN TÍCH KHÁM PHÁ',
        'Ngày 1 • Δ BSS so E0 • CI95% mô tả, khối bảy ngày • Không thay CI điều chỉnh TFT đã công bố',
        'Mỗi ô là BSS khi thêm riêng một nhóm trừ BSS nền E0, cùng mô hình và ngày 1. A=E1, M=E2, C=E3. Đây là CI95% điểm của phân tích phát triển đã xem trước, không hiệu chỉnh nhiều phép. Phép thử TFT chính với gia đình chín CI đồng thời vẫn giữ nguyên ở slide sau. Không đồng nhất CI95% mới với CI99,44% chính.')
    rows=[]
    for model in MODELS:
        cells=[LABELS[model]]
        for e in ['E1','E2','E3']:
            q=paired[(paired.kind=='feature')&(paired.model==model)&(paired.experiment==e)&(paired.lead_day==1)&(paired.metric=='BSS16')&(paired.block_days==7)].iloc[0]
            cells.append(f'{q.estimate:+.3f}\n[{q.ci_low:+.3f}; {q.ci_high:+.3f}]')
        rows.append(cells)
    added.table(s,['Mô hình','Thêm A','Thêm M','Thêm C'],rows,y=1.60,row_h=.90,widths=[1.75,3.5,3.5,3.5],size=20,tag='feature_skill_ci')
    added.text(s,.70,5.45,11.95,.5,'Dương: điểm BSS cải thiện; khoảng chứa 0: còn bất định.',20)
    added.band(s,'Không dùng việc thử nhiều nhóm rồi chọn ô đẹp nhất để khẳng định đóng góp chắc chắn.',size=18)
    deck.slides[29:29]=added.slides
    for i,slide in enumerate(deck.slides,1):
        slide['number']=i
        for sh in slide['shapes']:
            if sh['kind']=='text' and abs(sh['x']-12.05)<.01 and abs(sh['y']-.18)<.01:sh['text']=f'{i:02d}'
    deck.save(OUT/(STEM+'.pptx'),OUT/(STEM+'.html'),OUT/'slide_scene.json')
    # Apply the scene's cell style to every paragraph, including multiline CIs.
    # Keep the frozen shared renderer and all previously published decks intact.
    prs = Presentation(OUT/(STEM+'.pptx'))
    for native, slide in zip(prs.slides, deck.slides):
        for shape, spec in zip(native.shapes, slide['shapes']):
            if spec['kind'] != 'table':
                continue
            for ri, row in enumerate([spec['headers']] + spec['rows']):
                for ci, _ in enumerate(row):
                    cell_style = spec['styles'].get(f'{ri},{ci}', {})
                    for paragraph in shape.table.cell(ri, ci).text_frame.paragraphs:
                        paragraph.alignment = PP_ALIGN.CENTER
                        paragraph.space_before = paragraph.space_after = Pt(0)
                        paragraph.line_spacing = 1.0
                        for run in paragraph.runs:
                            run.font.name = 'Arial'
                            run.font.size = Pt(spec['size'])
                            run.font.bold = ri == 0 or cell_style.get('bold', False)
                            run.font.color.rgb = RGBColor.from_string(
                                'FFFFFF' if ri == 0 else cell_style.get('color', NAVY))
    prs.save(OUT/(STEM+'.pptx'))
    html_path = OUT/(STEM+'.html')
    html = html_path.read_text(encoding='utf8')
    assert '.ct{overflow:visible}' in html
    html_path.write_text(html.replace('.ct{overflow:visible}',
        '.ct{overflow:visible;white-space:pre-line}'), encoding='utf8', newline='\n')
    roles=dict(groups=5,coverage=6,seeds=7,metrics=13,counts=19,bss_day1=24,bss_day2=25,bss_day3=26,
               e7_day1=27,e7_day2=28,e7_day3=29,e7_skill_ci=30,feature_skill_ci=31)
    def write(name,value):(OUT/name).write_text(value.rstrip()+'\n',encoding='utf8',newline='\n')
    write('LOI_DAN.md','# Lời dẫn: đủ ma trận ba mô hình\n\n'+'\n\n'.join(f"## Slide {s['number']}: {s['title']}\n\n{s['note']}" for s in deck.slides))
    export=scores[['experiment','model','lead_day','n','events',*METRICS]].copy()
    export.to_csv(OUT/'KET_QUA_72.csv',index=False,lineterminator='\n')
    groups=json.loads((ROOT/'reports/regional_research_plan_2026_10_20/feature_groups.json').read_text(encoding='utf8'))
    catalog=pd.DataFrame([dict(group=g,feature=name) for g in ['B','A','M','C'] for name in groups[g]])
    assert len(catalog)==44
    sheets={'KET_QUA_72':export,'MUA_VA_NGUONG':full,'TUNG_NAM':pd.read_csv(MATRIX/'scores_ensemble.csv'),
            'INFORMER_SEED':pd.read_csv(MATRIX/'scores_seed_pooled.csv'),'ON_DINH_SEED':pd.read_csv(MATRIX/'seed_stability.csv'),
            'CI_CHENH_LECH':paired,'CI_BSS_E7':absolute,'NGUONG_XAC_SUAT':pd.read_csv(MATRIX/'calibration_thresholds_all_models.csv'),
            'DANH_MUC_44_BIEN':catalog}
    with pd.ExcelWriter(OUT/'BANG_KET_QUA_DAY_DU.xlsx',engine='openpyxl') as writer:
        for name,data in sheets.items():
            data.to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor=NAVY)
            for cells in ws.columns:
                width=min(48,max(12,max(len(str(c.value or '')) for c in list(cells)[:100])+2))
                ws.column_dimensions[cells[0].column_letter].width=width
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value,float):cell.number_format='0.000'
    overview=f'''# Kết quả hoàn chỉnh LightGBM, TFT và Informer — 28/10/2026

Đã chạy thêm **54 lượt Informer E1–E6**, cùng hàm fit và protocol pilot. Dùng lại **18 lượt E0/E7**; tổng **72 Informer**, cùng 72 TFT và 24 pipeline LightGBM đã có. Đủ **168 lượt mô hình/pipeline** và **72 dòng chính** = tám E × ba mô hình × ba ngày. Không mở test 2025, không tuning hoặc thay bước gộp sáu giờ.

Độc lập phát lại đủ 72 checkpoint Informer trên 144 split calibration/evaluation; tái tạo scaler train và calibration. E0/E7 Informer và toàn bộ điểm TFT/LightGBM giữ nguyên. Có cùng 1.075 mốc/hạn, nhãn, tham chiếu Brier và quy trình chọn ngưỡng trên calibration.

## Kết quả E7: cùng 44 biến, đủ ba mô hình

'''
    for day in [1,2,3]:
        rows,_=values_and_styles('E7',day)
        overview+=f'### Ngày {day}\n\n'+markdown_table(['Mô hình',*HEADERS],rows)+'\n\n'
    overview+='''LightGBM E7 ngày 1–2 có nhiều điểm tốt hơn; Informer CSI cao nhất ngày 1, TFT POD cao nhất ngày 1 nhưng báo nhầm nhiều hơn. Cả ba E7 ngày 3 BSS âm. Riêng Informer E6 ngày 3 có BSS khoảng +0,004, tương ứng giảm khoảng 0,4% sai số Brier; điểm dương nhỏ này chưa đủ để khẳng định kỹ năng chắc chắn. Các mô hình có thể có lợi ở E khác nhau; không chọn kiến trúc hoặc nhóm biến chỉ từ một điểm đẹp.

## Đọc bảng nhóm biến và độ bất định

E0=B; E1=B+A; E2=B+M; E3=B+C; E4=B+A+M; E5=B+A+C; E6=B+M+C; E7=B+A+M+C. Số biến lần lượt 11/23/31/12/43/24/32/44. Nhóm C chỉ là SST hộp biển, không phải đã kiểm tra ENSO/MJO đầy đủ.

CI mới dùng 20.000 bootstrap khối bảy ngày và độ nhạy 14 ngày, giữ chung mẫu giữa mô hình, ngưỡng và các ngày; lấy mẫu trong từng năm. Đây là **CI95% mô tả**, chưa hiệu chỉnh nhiều phép và có điều kiện trên mô hình đã fit. Chúng không bao trùm toàn bộ bất định huấn luyện hoặc nhãn. Thêm biến cũng có thể làm đổi số tham số mạng, cần công bố khi diễn giải hiệu ứng nhóm. Phép thử TFT chính với CI điều chỉnh gia đình chín phép vẫn giữ riêng; không đổi phân tích mới thành kiểm định xác nhận. Kiểm định cuối 2025 chưa mở.

## Lịch, chín xác suất và bước sáu giờ

Bốn biến lịch là sin_hour/cos_hour/sin_doy/cos_doy, biểu diễn hai chu kỳ giờ trong ngày và ngày trong năm theo giờ Việt Nam. Decoder chỉ biết lịch tương lai, không biết ERA5 tương lai. 72 giờ tương lai tạo 12 khối sáu giờ × bốn biến lịch.

Đầu ra có ba lượng mm/24h và chín xác suất: ba ngày × ba sự kiện đạt 16/50/100 mm. Các sự kiện lồng nhau, không cộng ba xác suất thành 100%; xác suất nói về tổng mưa trung bình vùng, không phải phần trăm diện tích có mưa. Ngưỡng nghiên cứu này không tự là cấp cảnh báo hành chính.

Gộp sáu giờ là lấy trung bình đặc trưng khí tượng để giảm 168 giờ lịch sử còn 28 bước. Nhãn vẫn cộng 24 giờ để tạo tổng ngày. Chưa có đối chứng cùng điều kiện 1h/3h/6h để chứng minh sáu giờ tối ưu hoặc không mất đỉnh CAPE/gió/ẩm. Các kết quả mới vẫn có điều kiện trên lựa chọn sáu giờ đã khóa.

## Seed, LightGBM và BSS

TFT/Informer có ba seed 42/43/44: cùng dữ liệu/cấu hình nhưng khác khởi tạo và thao tác ngẫu nhiên khi học. Báo cáo dùng trung bình dự đoán ba seed rồi tính metric, không chọn seed đẹp nhất; ba seed không tạo thêm ngày thời tiết. LightGBM hiện chạy xác định nên mỗi E/năm có một pipeline.

LightGBM là mô hình tham chiếu để so sánh, không phải mô hình buộc TFT phải thắng. Trọng tâm là đóng góp đặc trưng và đánh giá các pipeline cùng điều kiện. Cây dùng thống kê lịch sử, mạng dùng chuỗi; chưa cô lập riêng ảnh hưởng kiến trúc.

BSS = 1 − BS_mô_hình/BS_khí_hậu. BSS=0,03 là giảm 3% sai số Brier; BSS=−0,038 là tăng 3,8% sai số Brier, không phải tỷ lệ ngày dự đoán sai. BSS âm chỉ kết luận cho xác suất ở ngưỡng, hạn và tập đang xét. POD cao không tự bảo đảm BSS dương.

## Giới hạn và bước tiếp

Footprint 93 toàn ô, diện tích khoảng 11.074 km², còn xấp xỉ địa giới Đà Nẵng mới. Nhãn GSMaP Gauge và provenance khí tượng cần tiếp tục QC. Đủ ma trận không có nghĩa đã xong kiểm chứng ranh/nhãn hoặc chứng minh vận hành. Khóa thiết kế trước holdout; không sửa ngưỡng hay chọn cấu hình theo test 2025.

'''
    write('GIAI_THICH_KET_QUA_DAY_DU.md',overview);word_copy(OUT/'GIAI_THICH_KET_QUA_DAY_DU.md')
    write('README.md',f'''# Đủ E0–E7: LightGBM, TFT, Informer

- [PowerPoint {len(deck.slides)} slide]({STEM}.pptx), [HTML]({STEM}.html), [lời dẫn](LOI_DAN.md).
- [Word giải thích](GIAI_THICH_KET_QUA_DAY_DU.docx), [Markdown](GIAI_THICH_KET_QUA_DAY_DU.md).
- [Excel chín sheet](BANG_KET_QUA_DAY_DU.xlsx), [72 dòng chính](KET_QUA_72.csv).
- [54 lượt Informer mới, 18 lượt dùng lại và kiểm chứng](../informer_regional_matrix_2026_10_28/README.md).

Đủ 72 Informer + 72 TFT + 24 pipeline LightGBM; ba mô hình cùng E/ngày/mẫu. Có {len(deck.slides)} slide, 34 slide chính và 15 tra cứu. Các bảng E7 ở slide 27–29; BSS và CI ở 30; thêm riêng A/M/C ở 31. Đậm là điểm tốt hơn cùng E/ngày, chưa tự là ý nghĩa thống kê.

54 lượt mới giữ hàm fit cũ; E0/E7 và điểm TFT/LightGBM không đổi. CI mới là mô tả phát triển, chưa hiệu chỉnh nhiều phép; giữ riêng phép thử TFT chính. Không đọc 2025. Còn giới hạn ranh/QC nhãn, bước sáu giờ và kiểm định cuối. Tệp slide cũ do người dùng chỉnh sửa giữ nguyên; đây là phiên bản mới từ scene báo cáo đã công bố.

Tái lập báo cáo: `python scripts/build_short_range_complete.py`, `python scripts/render_short_range_complete.py`, `python scripts/validate_short_range_complete.py`, sau khi hoàn tất huấn luyện/kiểm chứng/đánh giá Informer.
''')
    sources=[OLD/'slide_scene.json',MATRIX/'scores_ensemble_pooled.csv',MATRIX/'paired_bootstrap_exploratory.csv',
             MATRIX/'absolute_skill_ci.csv',MATRIX/'run.json',MATRIX/'validation.json',MATRIX/'evaluation.json',
             ROOT/'reports/regional_research_plan_2026_10_20/feature_groups.json']
    sources.extend(sorted(MATRIX.glob('*.csv')))
    files={p.name:dict(sha256=sha(p),bytes=p.stat().st_size) for p in OUT.iterdir() if p.is_file() and p.name not in ['build_manifest.json','render_manifest.json','validation.json','XEM_TRUOC.png']}
    write('build_manifest.json',json.dumps(dict(status='complete',slides=len(deck.slides),main_slides=34,roles=roles,
        models=MODELS,complete_score_rows=72,missing_score_rows=0,completed_fits=dict(tft=72,lgbm=24,informer=72),
        new_informer_fits=54,reused_informer_fits=18,test2025_read=False,code_sha256=sha(__file__),
        sources={str(p.relative_to(ROOT)).replace('\\','/'):sha(p) for p in sources},files=files),ensure_ascii=False,indent=2))
    print(json.dumps(dict(slides=len(deck.slides),score_rows=72,roles=roles),indent=2))


if __name__=='__main__':main()

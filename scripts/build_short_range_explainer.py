"""Build a full teaching report and workbook from audited development aggregates.

No training, raw data read, 2024 calibration or 2025 holdout evaluation occurs.
"""
from __future__ import annotations
from pathlib import Path
import hashlib
import itertools
import json
import re
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches,Pt
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.hyperlink import Hyperlink
from build_final_regional_plan import word_copy

SOURCE=ROOT/'reports/short_range_ablation_2026_10_21'
OUT=ROOT/'reports/short_range_explained_2026_10_21'
GROUPS={'E0':['B'],'E1':['B','A'],'E2':['B','M'],'E3':['B','C'],
        'E4':['B','A','M'],'E5':['B','A','C'],'E6':['B','M','C'],'E7':['B','A','M','C']}
REFS=[
    ('TFT, bài gốc','https://arxiv.org/abs/1912.09363'),
    ('Informer, AAAI 2021','https://ojs.aaai.org/index.php/AAAI/article/view/17325'),
    ('Informer, mã của tác giả','https://github.com/zhouhaoyi/Informer2020/blob/main/models/model.py'),
    ('LightGBM, bài gốc NeurIPS 2017','https://proceedings.neurips.cc/paper_files/paper/2017/hash/6449f44a102fde848669bdd9eb6b76fa-Abstract.html'),
    ('LightGBM, mô tả thuật toán chính thức','https://lightgbm.readthedocs.io/en/stable/Features.html'),
    ('Brier score, scikit-learn','https://scikit-learn.org/stable/modules/generated/sklearn.metrics.brier_score_loss.html'),
    ('Các chỉ số hồi quy, scikit-learn','https://scikit-learn.org/stable/modules/model_evaluation.html'),
    ('Các phương pháp kiểm chứng khí tượng, CAWCR/WMO','https://www.cawcr.gov.au/projects/verification/Stanski_et_al/VerificationSWBPartI.pdf'),
    ('Kiểm chứng xác suất so khí hậu, Mason','https://www.cawcr.gov.au/projects/verification/Mason/IntegratedVerificationProcedures.pdf')]


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def readable(value):
    protected=[]
    names=['ERA5','SHA256','LightGBM','GSMaP','V1','t2m','td2m','rh850','F1','F2','F3','E0','E1','E2','E3','E4','E5','E6','E7','Y1','Y2','Y3',
           'BSS16','CSI16','POD16','FAR16','p16','p50','p100']
    initial=set(names)
    source_catalog=SOURCE/'feature_catalog.csv'
    if source_catalog.exists():
        c=pd.read_csv(source_catalog)
        names+=c.native_variable.tolist()+c.v1_column.tolist()
    def hold(match):
        protected.append(match.group(0));return '\ue000'+str(len(protected)-1)+'\ue001'
    identifiers=r'(?<![A-Za-z0-9_])(?:'+r'|'.join(re.escape(n) for n in sorted(set(names)-initial,key=len,reverse=True))+r')(?![A-Za-z0-9_])'
    identifiers+=r'|'+r'|'.join(re.escape(n) for n in sorted(initial,key=len,reverse=True))
    token_pattern=r'`[^`]*`|https?://[^\s)]+|[A-Za-z0-9_./-]+\.(?:py|md|csv|json|xlsx|docx|png|pptx)|'+identifiers
    value=re.sub(token_pattern,hold,value)
    replacements={
        'đối-sựkiện':'đối sự kiện','hiệuquả':'hiệu quả','kiếntrúc':'kiến trúc','đốilưu':'đối lưu',
        'khíhậu':'khí hậu','khíhọc':'khí học','sựkiện':'sự kiện','tíchphân':'tích phân',
        'biếnđộng':'biến động','thànhphần':'thành phần','độẩm':'độ ẩm','vậnchuyển':'vận chuyển',
        'dựđoán':'dự đoán','dựbáo':'dự báo','huấnluyện':'huấn luyện','tốiưuhóa':'tối ưu hóa',
        'tốiưungưỡng':'tối ưu ngưỡng','tốiưu':'tối ưu','môhình':'mô hình','hồiquy':'hồi quy',
        'nhómbiến':'nhóm biến','sốbiến':'số biến','biếntĩnh':'biến tĩnh','biểudiễn':'biểu diễn',
        'nănglực':'năng lực','đầuvào':'đầu vào','đầura':'đầu ra','sốlượt':'số lượt',
        'thôngtin':'thông tin','tổnghợp':'tổng hợp','tổngthể':'tổng thể','phạmvi':'phạm vi',
        'trọngtâm':'trọng tâm','nhậnđịnh':'nhận định','thiếtkế':'thiết kế','lựachọn':'lựa chọn',
        'đốichứng':'đối chứng','cùngmẫu':'cùng mẫu','cùngnguồn':'cùng nguồn','hồicứu':'hồi cứu',
        'quákhứ':'quá khứ','tươnglai':'tương lai','thựcnghiệm':'thực nghiệm','kiểmtra':'kiểm tra',
        'kiểmchứng':'kiểm chứng','kiểmđịnh':'kiểm định','thửnghiệm':'thử nghiệm','thínghiệm':'thí nghiệm',
        'cốđịnh':'cố định','sắncó':'sẵn có','pháthành':'phát hành','mẫusố':'mẫu số',
        'khảthi':'khả thi','thấtbại':'thất bại','kếtluận':'kết luận','nghiêncứu':'nghiên cứu',
        'toànnăm':'toàn năm','từngngày':'từng ngày','toànvùng':'toàn vùng','từngissue':'từng issue',
        'sốmẫu':'số mẫu','sốđếm':'số đếm','lớnnhất':'lớn nhất','nhỏnhất':'nhỏ nhất',
        'thamchiếu':'tham chiếu','lượngmm':'lượng mm','lượngmưa':'lượng mưa','xácsuất':'xác suất',
        'báosót':'báo sót','bỏsót':'bỏ sót','báonhầm':'báo nhầm','báocáo':'báo cáo',
        'chỉsố':'chỉ số','dữliệu':'dữ liệu','hạnngày':'hạn ngày','độc lập':'độc lập',
        'độclập':'độc lập','giao thức':'giao thức','giađình':'gia đình','hiệuchỉnh':'hiệu chỉnh',
        'cải thiện':'cải thiện','cảithiện':'cải thiện','điểmcaohơn':'điểm cao hơn','phépso':'phép so',
        'ngàygiữkín':'ngày giữ kín','nămgiữkín':'năm giữ kín','giữkín':'giữ kín','mãhóa':'mã hóa',
        'từngfold':'từng fold','khối7ngày':'khối 7 ngày','tham số':'tham số','thamsố':'tham số',
        'đầyđủ':'đầy đủ','giải thích':'giải thích','giảithích':'giải thích','bằngchứng':'bằng chứng',
        'caođiểm':'cao điểm','hợp lệ':'hợp lệ','hợplệ':'hợp lệ','pháttriển':'phát triển',
        'bảnđồ':'bản đồ','tỉnhthành':'tỉnh thành','ngàyđích':'ngày đích','khôngđổi':'không đổi',
        'cựctrị':'cực trị','cựcđoan':'cực đoan','bổsung':'bổ sung','đúngvới':'đúng với',
        'cầnđọc':'cần đọc','đượcđo':'được đo','đượcchọn':'được chọn','đãchạy':'đã chạy',
        'đangdùng':'đang dùng','suyrộng':'suy rộng','địaphương':'địa phương','địa điểm':'địa điểm',
        'địađiểm':'địa điểm','địagiới':'địa giới','hànhchính':'hành chính','chínhxác':'chính xác',
        'xấpxỉ':'xấp xỉ','chốtgiao':'chốt giao','chốtranh':'chốt ranh','trọngsố':'trọng số',
        'thờigian':'thời gian','bướctiếp':'bước tiếp','phầnsai':'phần sai','saisố':'sai số',
        'nhiềuphép':'nhiều phép','mọicặp':'mọi cặp','thốngkê':'thống kê','đốichiếu':'đối chiếu',
        'mỗingày':'mỗi ngày','mỗilead':'mỗi lead','từngseed':'từng seed','lộtrình':'lộ trình',
        'nútthắt':'nút thắt','bắtđầu':'bắt đầu','kỳvọng':'kỳ vọng','nguyênnhân':'nguyên nhân',
        'vậtlý':'vật lý','chiềutốt':'chiều tốt','điểmtốt':'điểm tốt','khảdụng':'khả dụng',
        'nguồnlực':'nguồn lực','cổngdữliệu':'cổng dữ liệu','khoảngtin':'khoảng tin',
        'lấylại':'lấy lại','cùngthôngtin':'cùng thông tin','tậpnhỏ':'tập nhỏ','kếtquả':'kết quả',
        'cậpnhật':'cập nhật','lýdo':'lý do','nhiềuhạn':'nhiều hạn','dựán':'dự án',
        'tài liệu':'tài liệu','tàiliệu':'tài liệu','hiện tại':'hiện tại','hiệntại':'hiện tại',
        'khaithác':'khai thác','liênhệ':'liên hệ','đượcphânbiệt':'được phân biệt','nguyêntắc':'nguyên tắc',
        'gốctừ':'gốc từ','nguồngốc':'nguồn gốc','bàigốc':'bài gốc','biếnbối':'biến bối',
        'chứngnhận':'chứng nhận','đảm bảo':'đảm bảo','bảođảm':'bảo đảm','tỷlệ':'tỷ lệ',
        'bìnhphương':'bình phương','trungbình':'trung bình','khácnhau':'khác nhau','khácbiệt':'khác biệt',
        'chồngnhau':'chồng nhau','ngẫunhiên':'ngẫu nhiên','lặp lại':'lặp lại','lặplại':'lặp lại',
        'thựcđạtngưỡng':'thực đạt ngưỡng','đạtngưỡng':'đạt ngưỡng','đượcchấm':'được chấm',
        'giờcósẵn':'giờ có sẵn','cósẵn':'có sẵn','dùngchung':'dùng chung','tinhchỉnh':'tinh chỉnh',
        'trướckhi':'trước khi','saukhi':'sau khi','mãsrc/':'mã src/','cácscripts':'các scripts',
        'mỗihàng':'mỗi hàng','từngdòng':'từng dòng','chưacó':'chưa có','cóđiểm':'có điểm',
        'cảihthiện':'cải thiện','chỉđọc':'chỉ đọc','theotháng':'theo tháng','dựa vào':'dựa vào',
        'bỏqua':'bỏ qua','khôngfit':'không fit','coilà':'coi là','đơnvị':'đơn vị',
        'dựđịnh':'dự định','quanmô':'quan mô','chênháp':'chênh áp','giómùa':'gió mùa',
        'đất':'đất','hệthống':'hệ thống','bảngtóm':'bảng tóm','thuậntiện':'thuận tiện',
        'độnhạy':'độ nhạy','đượcđốichiếu':'được đối chiếu','chứckhông':'chứ không',
        'giao thức':'giao thức','chọnmôhình':'chọn mô hình','làmđẹp':'làm đẹp',
        'nguồn kỹthuật':'nguồn kỹ thuật','kỹthuật':'kỹ thuật','vào93':'vào 93',
        'khôbáokhô':'khô báo khô','E0BSS':'E0 BSS','BSSkhí':'BSS khí',
        'forecastthật':'forecast thật','losschung':'loss chung','staticcontext':'static context',
        'staticcovariates':'static covariates','futurecovariates':'future covariates',
        'ProbSparseattention':'ProbSparse attention','decoderLSTM':'decoder LSTM',
        'epochTFT':'epoch TFT','τquyếtđịnh':'τ quyết định','τcalibration':'τ calibration',
        'τchọncalibration':'τ chọn calibration','τtừevaluation':'τ từ evaluation',
        'num_leaves15':'num_leaves 15','learningrate':'learning rate','tuningbudget':'tuning budget',
        'skillscore':'skill score','losschung/riêng':'loss chung/riêng','falsealarmratio':'false alarm ratio',
        'TFTchưa':'TFT chưa','CIkhông':'CI không','CIchứa':'CI chứa','CIđiềuchỉnh':'CI điều chỉnh',
        'sheetHUONG_DAN':'sheet HUONG_DAN','cácsheetSEED':'các sheet SEED','NaNkhông':'NaN không',
        'Buildmanifest':'Build manifest','SHA256nguồn':'SHA256 nguồn','pipelinehọc':'pipeline học',
        'họhọc':'họ học','cùnggiao':'cùng giao','nhiềunhóm':'nhiều nhóm','đầu xácsuất':'đầu xác suất',
        'mô hìnhkhác':'mô hình khác','liênquan':'liên quan','cóartifact':'có artifact','vàreplay':'và replay',
        'đủngưỡng':'đủ ngưỡng','cósai':'có sai','thiếtkế':'thiết kế','hiệu chỉnh':'hiệu chỉnh',
        'nhómC':'nhóm C','M−A':'M−A','ĐàNẵngmới':'Đà Nẵng mới'}
    replacements.update({'vậnchuyểnẩm':'vận chuyển ẩm','độẩmđất':'độ ẩm đất','ẩmđất':'ẩm đất',
        'chuyểnẩm':'chuyển ẩm','mưa thật':'mưa thật','mưathật':'mưa thật','biểnkhông':'biển không',
        'đượclấy':'được lấy','phầnnày':'phần này','thángđích':'tháng đích','đượcfit':'được fit',
        'giờquá':'giờ quá','tậpngày':'tập ngày','từngissue':'từng issue','báotừng':'báo từng',
        'ởđây':'ở đây','đangtrênfootprint':'đang trên footprint','ởmột':'ở một','cóthể':'có thể',
        'đangtrên':'đang trên','củabài':'của bài','củabáo':'của báo','cảihọc':'cải học',
        'đủcả':'đủ cả','nhiềubiến':'nhiều biến','sắpxếp':'sắp xếp','saukhóa':'sau khóa',
        'cácmùa':'các mùa','cácngày':'các ngày','cáccột':'các cột','cáctập':'các tập',
        'đềuchung':'đều chung','đồngđều':'đồng đều','trênđây':'trên đây','đàotạo':'đào tạo',
        'đốitượng':'đối tượng','mưacực':'mưa cực','đợtmưa':'đợt mưa','BSS0,':'BSS 0,',
        'khungphân':'khung phân','đọcnhư':'đọc như','chấmlại':'chấm lại','xácminh':'xác minh',
        'độnhạy':'độ nhạy','chưađủ':'chưa đủ','mởrộng':'mở rộng','chưasocùng':'chưa so cùng',
        'thuộcA/M':'thuộc A/M','ngàyđược':'ngày được','liênquan':'liên quan','làmtrơn':'làm trơn',
        'nhãnvệ':'nhãn vệ','họở':'họ ở','họở':'họ ở','của93':'của 93','toàn93':'toàn 93',
        'tạo3':'tạo 3','nhận12':'nhận 12','thành28':'thành 28','tổng24':'tổng 24',
        'tươngđối':'tương đối','quantrắc':'quan trắc','Giaiđoạn':'Giai đoạn',
        'Staticcovariates':'Static covariates','gradientboosting':'gradient boosting',
        'giátrị':'giá trị','quákhớp':'quá khớp','trong±':'trong ±','cóaccuracy':'có accuracy',
        'τensemble':'τ ensemble','τseed':'τ seed','cácτ':'các τ','họcτ':'học τ',
        'giữfallback':'giữ fallback','vịmm':'vị mm','encoder,12':'encoder, 12',
        'mục13':'mục 13','nămdevelopment':'năm development','kểcả':'kể cả',
        'kiểm địnhmọi':'kiểm định mọi','đểchọn':'để chọn','ràsoátnguồn':'rà soát nguồn',
        'Dựtoán':'Dự toán','chưabenchmark':'chưa benchmark','Tínhiệu':'Tín hiệu',
        'haihọ':'hai họ','BSSâm':'BSS âm','chưaủnghộ':'chưa ủng hộ',
        'tuyênbố':'tuyên bố','nguyênô':'nguyên ô','ngoàihộp':'ngoài hộp',
        'phầnphát':'phần phát','sảnphẩm':'sản phẩm','Biếnmã':'Biến mã',
        'hóaCIN':'hóa CIN','ýnghĩanguồn':'ý nghĩa nguồn','kịpt':'kịp t',
        'cảthờiđiểm':'cả thời điểm','cáchhọc':'cách học','trậnchưa':'trận chưa',
        'côlập':'cô lập','mộtvùng':'một vùng','giảngviên':'giảng viên',
        'chứngminh':'chứng minh','trậnphát':'trận phát','trung bìnhvùng':'trung bình vùng',
        'Thêmriêng':'Thêm riêng','thểhiện':'thể hiện','tổng thểso':'tổng thể so ',
        'theoBSS':'theo BSS','lỗiBrier':'lỗi Brier','Informerkhôngkhả':'Informer không khả',
        'giảthuyết':'giả thuyết','đãthử':'đã thử','Informerchưa':'Informer chưa',
        'nguyênliệu':'nguyên liệu','vàbiểu':'và biểu','tổhợp':'tổ hợp',
        'toànĐà':'toàn Đà','hiệncòn':'hiện còn','cầnchốt':'cần chốt',
        'Chạy3':'Chạy 3','đánhgiá':'đánh giá','vẫn1075':'vẫn 1075',
        'đoổnđịnh':'đo ổn định','khônggiúp':'không giúp','khôngquantrọng':'không quan trọng',
        'cóbằng':'có bằng','chứngcho':'chứng cho','mộtbiến':'một biến',
        'biếttrước':'biết trước','vàTUNG_NAM':'và TUNG_NAM',
        'chứađối':'chứa đối','bấtđịnh':'bất định','Ôtrống':'Ô trống','choNaN':'cho NaN',
        'tựcoi':'tự coi','tiệncopyđưavào':'tiện để sao chép vào ',
        'cộtbảng':'cột bảng','giúpđối':'giúp đối','lênGit':'lên Git',
        'Cácnguồn':'Các nguồn','thuậtđã':'thuật đã','đãđối':'đã đối',
        'phầnnhận':'phần nhận','dụngdự':'dụng dự','đểkiểm':'để kiểm',
        'tươngứng':'tương ứng','đãcó':'đã có','số0':'số 0',
        'bình24':'bình 24','thứ tự28':'thứ tự 28','độ2':'độ 2',
        'EpochTFT':'Epoch TFT','fitPlatt':'fit Platt','Bootstrappaired':'Bootstrap paired',
        'theoBonferroni':'theo Bonferroni','giờGPU':'giờ GPU'})
    for old,new in sorted(replacements.items(),key=lambda p:-len(p[0])):value=value.replace(old,new)
    value=re.sub(r'(?<=\d)(giờ|ngày|biến|mẫu|ô\b|lượt|bước|khối|cột|nghìn|phép|đợt|tổng|tháng|sự kiện|quantrắc|họ|ngưỡng|pipeline|mm\b|km\b|h\b|seed|epoch)',r' \1',value)
    value=re.sub(r'(TFT|LightGBM|LGBM|RMSE|seed|hidden|epoch|ngày|giờ|tháng|khối|ngưỡng|sheet)(?=\d)',r'\1 ',value)
    value=re.sub(r'(có|đủ|thêm|mỗi|trên|bằng|khoảng|từ|đến|với|thiếu|và|là|trước|sau|toàn|tạo|nhận|thành|tổng)(?=\d)',r'\1 ',value)
    value=re.sub(r'(E[0-7])(?=TFT|LightGBM|LGBM|BSS|CSI|có|chưa|ở|so|là)',r'\1 ',value)
    value=re.sub(r'(CI|BSS|CSI|TFT|LGBM)(?=chứa|trên|dưới|khí|là|của|và|khác|nên)',r'\1 ',value)
    value=re.sub(r'(?<=[^\W\d_])(?=\d)|(?<=\d)(?=[^\W\d_])',' ',value)
    value=re.sub(r'(?<=%)(?=[^\W\d_])',' ',value)
    value=re.sub(r'(?<=[^\W\d_])(\ue000\d+\ue001)',r' \1',value)
    value=re.sub(r'(\ue000\d+\ue001)(?=[^\W\d_])',r'\1 ',value)
    value=re.sub(r'\ue000(\d+)\ue001',lambda m:protected[int(m[1])],value)
    return value
def write(name,value):
    if name.endswith('.md'):value=readable(value)
    (OUT/name).write_text(value,encoding='utf-8',newline='\n')
def number(v):
    if pd.isna(v):return '—'
    if isinstance(v,(int,np.integer)):return str(v)
    if isinstance(v,(float,np.floating)):return f'{v:.3f}'
    return str(v)
def table(headers,rows):
    return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |',
        *['| '+' | '.join(number(x) for x in row)+' |' for row in rows]])


def illustrations():
    fig,ax=plt.subplots(figsize=(12.4,3.2));ax.set_xlim(-7.3,3.2);ax.set_ylim(-.75,1.65);ax.axis('off')
    for x,w,color,label in [(-7,7,'#e2e8f0','168 giờ khí tượng quá khứ\n28 bước × 6 giờ'),
                            (0,1,'#bfdbfe','Ngày 1\nY1'),(1,1,'#99f6e4','Ngày 2\nY2'),(2,1,'#fed7aa','Ngày 3\nY3')]:
        ax.add_patch(FancyBboxPatch((x,.25),w,.65,boxstyle='round,pad=.02',fc=color,ec='white',lw=2))
        ax.text(x+w/2,.575,label,ha='center',va='center',fontsize=11,color='#123047')
    for x,label in [(0,'t: 07 giờ VN'),(1,'t+24h'),(2,'t+48h'),(3,'t+72h')]:
        ax.plot([x,x],[.12,1.03],color='#334155',lw=1);ax.text(x,-.04,label,ha='center',va='top',fontsize=10)
    ax.text(-3.5,-.40,'Chỉ dùng khí tượng có mốc trước t trong thí nghiệm hồi cứu',ha='center',fontsize=10)
    ax.text(1.5,1.30,'Ba tổng 24 giờ riêng; không phải một tổng tích lũy 72 giờ',ha='center',fontsize=11,color='#087e8b')
    fig.tight_layout();fig.savefig(OUT/'dinh_nghia_han_du_bao.png',dpi=180);plt.close(fig)
    fig,ax=plt.subplots(figsize=(12.4,3.5));ax.axis('off')
    boxes=[(.01,.23,.21,.55,'93 ô nguồn\nERA5 quá khứ\nGSMaP Gauge nhãn'),(.28,.23,.19,.55,'Tổng hợp trước\nmô hình\nMột chuỗi vùng'),
           (.53,.23,.19,.55,'Cùng cấu hình E\nTFT hoặc LightGBM\nCùng mẫu/nhãn'),(.79,.23,.20,.55,'3 lượng mưa mm\n9 xác suất sự kiện\nChấm cùng chỉ số')]
    for x,y,w,h,label in boxes:
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.005',fc='#e8f0f4',ec='#087e8b',transform=ax.transAxes))
        ax.text(x+w/2,y+h/2,label,ha='center',va='center',fontsize=11,transform=ax.transAxes,color='#123047')
    for x0,x1 in [(.225,.265),(.48,.515),(.735,.775)]:
        ax.annotate('',xy=(x1,.50),xytext=(x0,.50),xycoords='axes fraction',arrowprops=dict(arrowstyle='->',color='#123047',lw=2))
    ax.text(.5,.08,'Dự đoán vùng trực tiếp: tổng hợp trước khi học; không dự đoán từng ô rồi mới gộp',ha='center',fontsize=11,transform=ax.transAxes)
    fig.tight_layout();fig.savefig(OUT/'luong_du_lieu_mo_hinh.png',dpi=180);plt.close(fig)


def workbook(sheets,descriptions):
    with pd.ExcelWriter(OUT/'BANG_KET_QUA_DAY_DU.xlsx',engine='openpyxl') as writer:
        for name,frame in sheets.items():frame.to_excel(writer,sheet_name=name,index=False)
        for ws in writer.book.worksheets:
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.fill=PatternFill('solid',fgColor='123047');cell.font=Font(name='Arial',color='FFFFFF',bold=True)
                cell.alignment=Alignment(wrap_text=True,vertical='center')
            ws.row_dimensions[1].height=30
            for ci,cell in enumerate(ws[1],1):
                label=str(cell.value);length=max([len(label)]+[len(str(ws.cell(r,ci).value or '')) for r in range(2,min(ws.max_row,80)+1)])
                ws.column_dimensions[cell.column_letter].width=min(65,max(12,length+2))
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font=Font(name='Arial',size=10)
                    if isinstance(cell.value,float):cell.number_format='0.000'
            ws.sheet_view.showGridLines=True
        guide=writer.book['HUONG_DAN']
        for row in range(2,guide.max_row+1):
            target=guide.cell(row,1).value
            if target in sheets:
                guide.cell(row,1).hyperlink=Hyperlink(ref=guide.cell(row,1).coordinate,location=f"'{target}'!A1",display=target)
                guide.cell(row,1).style='Hyperlink'


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    run=json.loads((SOURCE/'run.json').read_text(encoding='utf-8'))
    audit=json.loads((SOURCE/'checkpoint_audit.json').read_text(encoding='utf-8'))
    protocol=json.loads((SOURCE/'protocol_frozen.json').read_text(encoding='utf-8'))
    assert run['status']=='complete' and audit['status']=='passed' and audit['full_runs_replayed']==96
    models=run['models']
    expected_tft=set(itertools.product([2021,2022,2023],list(GROUPS),[42,43,44]))
    actual_tft={(m['year'],m['experiment'],m['seed']) for m in models if m['learner']=='tft'}
    assert actual_tft==expected_tft and sum(m['learner']=='tft' for m in models)==72
    assert {(m['year'],m['experiment']) for m in models if m['learner']=='lgbm'}==set(itertools.product([2021,2022,2023],list(GROUPS)))
    assert all(r['calibration_and_evaluation_replayed'] and r['all_metrics_reproduced'] for r in audit['results'])
    scores=pd.read_csv(SOURCE/'scores_ensemble_pooled.csv');by_year=pd.read_csv(SOURCE/'scores_ensemble.csv')
    seeds=pd.read_csv(SOURCE/'scores_seed_pooled.csv');seed_year=pd.read_csv(SOURCE/'scores_seed.csv')
    ci=pd.read_csv(SOURCE/'paired_bootstrap.csv');catalog=pd.read_csv(SOURCE/'feature_catalog.csv')
    primary=scores[(scores.season=='all')&(scores.threshold_mm==16)&scores.model.isin(['tft','lgbm'])].copy()
    primary=primary.sort_values(['model','lead_day','experiment']).reset_index(drop=True)
    assert len(primary)==48 and (primary.n==1075).all() and (primary.events==114).all()
    primary['groups']=primary.experiment.map(lambda e:'+'.join(GROUPS[e]))
    counts=catalog.groupby('group').size().to_dict();assert counts=={'A':12,'B':11,'C':1,'M':20}
    def r(model,exp,lead):return primary[(primary.model==model)&(primary.experiment==exp)&(primary.lead_day==lead)].iloc[0]
    rows=[]
    for exp,groups in GROUPS.items():
        entries=[m for m in models if m['experiment']==exp]
        features=next(m['feature_names'] for m in entries)
        assert len(features)==sum(counts[g] for g in groups)
        tab=next(m['tabular_features'] for m in entries if m['learner']=='lgbm')
        rows.append(dict(experiment=exp,groups='+'.join(groups),variables=len(features),tabular_columns=len(tab),
            tft_runs=sum(m['learner']=='tft' for m in entries),lgbm_pipelines=sum(m['learner']=='lgbm' for m in entries),status='complete'))
    experiment_table=pd.DataFrame(rows)
    run_table=pd.DataFrame([dict(model=m['learner'],year=m['year'],experiment=m['experiment'],seed=m['seed'],
        variables=len(m['feature_names']),best_epoch=m.get('best_epoch'),fitting_seconds=m['seconds'],
        tabular_columns=len(m.get('tabular_features',[])) if m['learner']=='lgbm' else None) for m in models])
    primary.to_csv(OUT/'KET_QUA_CHINH_48_DONG.csv',index=False,lineterminator='\n')
    experiment_table.to_csv(OUT/'TIEN_DO_E0_E7.csv',index=False,lineterminator='\n')
    event=scores[(scores.model=='climatology')&(scores.season=='all')][['lead_day','threshold_mm','n','events']]
    metric_rows=[
        ['RMSE','Lượng mưa','sqrt(mean((f-y)^2))','mm/24h','Nhỏ hơn','Nhạy với lỗi lớn; không phải khoảng dự báo'],
        ['MAE','Lượng mưa','mean(abs(f-y))','mm/24h','Nhỏ hơn','Sai số tuyệt đối trung bình; bổ sung RMSE'],
        ['bias','Lượng mưa','mean(f-y)','mm/24h','Gần 0','Âm: dự báo thiếu; dương: dự báo thừa'],
        ['BS','Xác suất','mean((p-o)^2)','Không đơn vị','Nhỏ hơn','o=1 nếu mưa thực đạt ngưỡng; o=0 nếu không'],
        ['BSS','Xác suất','1-BS/BS_khi_hau','Không đơn vị','Lớn hơn','0 ngang tham chiếu; âm kém hơn; không phải accuracy'],
        ['CSI','Quyết định sự kiện','H/(H+M+F)','0 đến 1','Lớn hơn','Bỏ ngày khô báo khô; ngưỡng xác suất chọn trên calibration'],
        ['POD','Quyết định sự kiện','H/(H+M)','0 đến 1','Lớn hơn','Bắt được bao nhiêu sự kiện thật; đọc cùng FAR'],
        ['FAR','Quyết định sự kiện','F/(H+F)','0 đến 1','Nhỏ hơn','Tỷ lệ báo nhầm trong các lần báo có; khác false alarm rate'],
        ['CC','Lượng mưa','Tương quan Pearson','-1 đến 1','Thường lớn hơn','Không phản ánh đầy đủ sai lệch lượng mm'],
        ['amount_CSI','Quyết định từ mm','CSI của f>=q','0 đến 1','Lớn hơn','Khác CSI chính được tính từ xác suất p>=tau'],
        ['reliability','Hiệu chỉnh xác suất','Tần suất thực theo nhóm xác suất','Không đơn vị','Gần đường chéo','Cần đọc số mẫu trong mỗi bin'],
        ['CI','Bất định hiệu số','Bootstrap paired khối 7/14 ngày','Theo metric','Tùy metric','Độ bất định của hiệu số, không phải chỉ số accuracy']]
    descriptions={
        'E0_E7':'Tám cấu hình và số lượt đã chạy; bắt đầu từ đây.',
        'KET_QUA_48':'48 dòng chính: 8 cấu hình × 2 mô hình × 3 ngày, toàn năm, ngưỡng16.',
        'NGUONG_16_50_100':'144 dòng toàn năm cho hai mô hình, cả ba ngưỡng sự kiện.',
        'POOLED_MUA':'Toàn bộ bảng gộp, cả khí hậu/luôn0 và các mùa; không tự trung bình BSS.',
        'TUNG_NAM':'Điểm từng năm evaluation2021/2022/2023, mùa và ngưỡng.',
        'TUNG_SEED_GOP':'Điểm từng seed TFT; LightGBM chỉ một seed xác định.',
        'TUNG_SEED_NAM':'Điểm từng seed theo từng năm, mùa, ngưỡng.',
        'CI_CHINH':'18 dòng = 9 phép so chính × 2 độ dài khối; CI đã hiệu chỉnh ở cột family_adjusted.',
        'CI_DAY_DU':'252 dòng, ba metric BSS16/CSI16/RMSE; các phép phụ dùng CI95% mô tả.',
        'SU_KIEN':'Số sự kiện16/50/100 theo lead; không cộng các lead thành các đợt độc lập.',
        'DANH_MUC_BIEN':'44 biến gốc, nhóm, đơn vị, tên cột V1 và mô tả nguồn.',
        'NHAT_KY_96':'96 lượt pipeline: năm, cấu hình, seed, bestepoch, thời gian.',
        'RELIABILITY':'Xác suất trung bình, tần suất quan trắc và số mẫu từng bin.',
        'NGUONG_CAL':'Ngưỡng quyết định ensemble chọn trên calibration, không trên evaluation.',
        'ON_DINH_SEED':'Trung bình/độ lệch chuẩn/min/max chỉ số qua ba seed.',
        'METRICS':'Công thức, đơn vị, chiều tốt và cách đọc.',
        'NGUON':'Tài liệu gốc và tài liệu thuật toán để tra cứu.'}
    guide=pd.DataFrame([dict(sheet=k,cach_doc=v) for k,v in descriptions.items()])
    sheets={'HUONG_DAN':guide,'E0_E7':experiment_table,'KET_QUA_48':primary,
        'NGUONG_16_50_100':scores[(scores.season=='all')&scores.model.isin(['tft','lgbm'])],
        'POOLED_MUA':scores,'TUNG_NAM':by_year,'TUNG_SEED_GOP':seeds,'TUNG_SEED_NAM':seed_year,
        'CI_CHINH':ci[ci.primary_family],'CI_DAY_DU':ci,'SU_KIEN':event,'DANH_MUC_BIEN':catalog,
        'NHAT_KY_96':run_table,'RELIABILITY':pd.read_csv(SOURCE/'reliability16.csv'),
        'NGUONG_CAL':pd.read_csv(SOURCE/'ensemble_calibration_thresholds.csv'),
        'ON_DINH_SEED':pd.read_csv(SOURCE/'seed_stability.csv'),
        'METRICS':pd.DataFrame(metric_rows,columns=['metric','nhom','cong_thuc','don_vi','chieu_tot','luu_y']),
        'NGUON':pd.DataFrame(REFS,columns=['tai_lieu','url'])}
    workbook(sheets,descriptions)
    illustrations()
    progress=table(['Cấu hình','Nhóm','Biến gốc','Cột LGBM','TFT','LGBM','Trạng thái'],experiment_table.values.tolist())
    full_tables=[]
    for lead in [1,2,3]:
        full_tables.append(f'### Ngày {lead}\n')
        for model,label in [('tft','TFT: trung bình dự báo ba seed'),('lgbm','LightGBM: một pipeline xác định')]:
            part=primary[(primary.model==model)&(primary.lead_day==lead)]
            full_tables.append(label+'\n\n'+table(['E','RMSE','MAE','Bias','BSS16','CSI16','POD16','FAR16'],
                part[['experiment','RMSE','MAE','bias','BSS','CSI','POD','FAR']].values.tolist())+'\n')
    feature_sections=[]
    for key,label in [('B','Nền: bề mặt và lịch'),('A','Khí quyển, đối lưu, mây, ẩm đất'),('M','Hoàn lưu và vận chuyển ẩm'),('C','SST vùng biển')]:
        part=catalog[catalog.group==key]
        feature_sections.append(f'### Nhóm {key}: {label} ({len(part)} biến)\n\n'+table(
            ['Tên biến gốc','Đơn vị','Nội dung theo mã nguồn'],part[['native_variable','unit','description_from_source_code']].values.tolist()))
    ci_main=ci[(ci.block_days==7)&ci.primary_family].sort_values(['lead_day','contrast'])
    ci_table=table(['Phép so TFT','Ngày','ΔBSS16','CI99,44%','Đọc kết quả'],[
        [x.contrast,x.lead_day,f'{x.estimate:+.3f}',f'[{x.family_adjusted_low:+.4f}; {x.family_adjusted_high:+.4f}]',
         'Trên0' if x.family_adjusted_low>0 else 'Chứa0' if x.family_adjusted_high>=0 else 'Dưới0'] for x in ci_main.itertuples()])
    baselines=scores[(scores.season=='all')&(scores.threshold_mm==16)&scores.model.isin(['climatology','zero'])]
    baseline_table=table(['Đối chứng','Ngày','RMSE','MAE','BSS16','CSI16'],baselines[['model','lead_day','RMSE','MAE','BSS','CSI']].values.tolist())
    relative_bs=(r('tft','E7',1).BSS-r('tft','E0',1).BSS)/(1-r('tft','E0',1).BSS)*100
    metric_table=table(['Chỉ số','Đo gì','Đơn vị','Chiều tốt'],[[x[0],x[1],x[3],x[4]] for x in metric_rows])
    table_bss=table(['E','TFT ngày1','TFT ngày2','TFT ngày3','LGBM ngày1','LGBM ngày2','LGBM ngày3'],[
        [exp,*[r(m,exp,d).BSS for m in ['tft','lgbm'] for d in [1,2,3]]] for exp in GROUPS])
    contents=['1. Đã chạy đủ chưa?','2. Mục tiêu dự đoán chính xác là gì?','3. Dữ liệu được biến thành mẫu học thế nào?',
        '4. E0–E7 thêm những gì?','5. Vì sao chọn TFT; Informer ở đâu?','6. So với LightGBM là so cái gì?',
        '7. Metrics: công thức, ví dụ, ý nghĩa','8. Chia tập và kiểm định như thế nào?',
        '9. Bảng đầy đủ 48 dòng chính','10. Khoảng tin cậy và ý nghĩa cải thiện',
        '11. Khả thi đến đâu và còn thiếu gì?','12. Nội dung có thể dùng để báo cáo',
        '13. Tra cứu đủ 44 đặc trưng','14. Cách dùng bảng Excel và nguồn tham khảo']
    text=f'''# Giải thích đầy đủ nghiên cứu dự đoán mưa vùng 24–72 giờ

Ngày 21/10/2026. Tài liệu này giải thích từ mục tiêu đến dữ liệu, mô hình, chỉ số và kết quả, để thành viên chưa đọc mã cũng theo được. Kết quả được lấy từ ma trận đã chạy; không huấn luyện thêm hoặc mở năm 2025 trong bước tổng hợp.

**Kết luận về trạng thái:** đã chạy đủ E0–E7 cho cả TFT và LightGBM trên ba fold phát triển. TFT có ba seed, tổng 72 lượt; LightGBM có 24 pipeline. Đã tái lập 96/96 lượt từ artifact lưu. **Hoàn thành ma trận phát triển chưa có nghĩa hoàn thành nghiên cứu cuối**, vì vùng hiện còn xấp xỉ, nhãn cần kiểm chứng và test 2025 còn giữ kín.

Lộ trình đọc: {'; '.join(contents)}.

## 1. Đã chạy đủ E0–E7 hay chưa?

Đã chạy đủ. Mỗi E có 9 lượt TFT = 3 năm evaluation × 3 seed, cùng 3 pipeline LightGBM = 3 năm evaluation. Bảng sau được dựng từ run.json, kiểm tra tập hợp năm/cấu hình/seed thực tế; không chỉ đếm tên thư mục hoặc dựa vào kế hoạch.

{progress}

“Seed” là khởi tạo ngẫu nhiên cho mạng neural. Chạy ba seed giúp thấy mức nhạy với tối ưu hóa; ba seed không tạo thêm ba tập thời tiết độc lập. LightGBM hiện chạy xác định, không lấy mẫu hàng/cột ngẫu nhiên, nên không nhân ba bản giống nhau để gọi là ba thí nghiệm. Một “pipeline LightGBM” có ba đầu hồi quy và chín đầu phân loại sự kiện; con số 24 không phải chỉ có 24 cây quyết định.

Bằng chứng gồm: run.json ghi đủ 96 lượt; checkpoint_audit.json ghi replay cả calibration và evaluation; bảng scores chứa đủ tổ hợp. Các fold cùng chấm E0–E7 trên mask E7 và cùng nhãn. Tình trạng “đủ ma trận” ở đây là ba năm phát triển 2021–2023, chưa phải đã chấm năm test 2025.

## 2. Mục tiêu dự đoán chính xác là gì?

### 2.1. Đầu ra nghiệp vụ và câu hỏi nghiên cứu

Đầu ra nghiệp vụ là **lượng mưa trung bình trên một vùng trong ba khoảng 24 giờ tương lai**. Câu hỏi nghiên cứu là: thêm nhóm thông tin A/M/C có làm dự đoán tốt hơn không, ở ngày nào và mùa nào? Mô hình đồng thời xuất xác suất vượt các ngưỡng mưa để đánh giá khả năng nhận diện sự kiện.

Ví dụ minh họa, không phải dự báo thật: một lần phát hành có thể trả lượng mưa [12,35,5] mm/24h và các xác suất vượt ngưỡng cho từng ngày. Ba số là ba khoảng riêng. Một số 35 mm là trung bình vùng cho ngày 2; từ đó không suy được mọi nơi trong vùng đều mưa 35 mm.

### 2.2. Ngày 1/2/3 nghĩa là gì?

Gọi t là giờ phát hành. Thí nghiệm phát hành một lần/ngày lúc 00 UTC, tức 07 giờ Việt Nam:

{table(['Đầu ra','Khoảng giờ sau t','Ví dụ minh họa theo giờ Việt Nam'],[
    ['Y1','[0,24)','07h hôm nay đến trước07h ngày mai'],
    ['Y2','[24,48)','07h ngày mai đến trước07h ngày kế tiếp'],
    ['Y3','[48,72)','07h ngày thứ ba đến trước07h ngày thứ tư']])}

Đây là định nghĩa đã dùng trong ma trận. “Hạn 24–72 giờ” không có nghĩa đã dự đoán riêng từng giờ từ giờ24 đến giờ72; cũng không phải Y2/Y3 là các tổng tích lũy 48/72 giờ. Nếu muốn đầu ra từng giờ hoặc bắt đầu từ 00h địa phương, phải xác định lại nhãn và chạy giao thức phù hợp.

Hình minh họa: dinh_nghia_han_du_bao.png.

### 2.3. “Mưa vùng trực tiếp” là gì?

Từng giờ, mưa của các ô được lấy trung bình diện tích trước mô hình: p_vùng(t)=tổng[w_i×p_i(t)], với tổng w_i=1. Nhãn ngày là tổng24giờ của p_vùng. Vì vậy giữ đơn vị mm trung bình vùng, không cộng mm của93ô với nhau. Khí tượng cũng được tổng hợp trước mô hình; một mẫu là một mốc của vùng.

Hiện sử dụng diện tích toàn93ô, khoảng11.074km². Chưa có giao ranh hành chính chính xác và kiểm tra đầy đủ phần lãnh thổ ngoài hộp nguồn. Do đó tên trung thực là **footprint93ô xấp xỉ khu vực Đà Nẵng mới**, chưa khẳng định trung bình đúng toàn địa giới. Đầu ra này không cung cấp bản đồ từng ô, cực đại mưa tại điểm hay cảnh báo từng xã.

## 3. Dữ liệu được biến thành mẫu học thế nào?

### 3.1. Vai trò của từng nguồn

ERA5 và các đại lượng suy ra từ nó cung cấp khí tượng quá khứ: nhiệt độ, ẩm, áp suất, gió, đối lưu, mây, đất và bối cảnh. Nhãn chính là GSMaP Gauge trung bình vùng, cột y_region_gc_mm. Lịch được tính từ timestamp. ERA5 có ích làm đầu vào không có nghĩa mọi biến đều chính xác; GSMaP Gauge cũng là sản phẩm ước lượng cần kiểm chứng với đối chứng phù hợp.

Mưa ERA5 không phải nhãn trong nhánh này. Các cột mưa vệ tinh quá khứ, mưa ERA5 và ONI cũ không được đưa vào ma trận chính hiện tại. Chưa có forecast khí tượng tương lai trong đầu vào. Vùng có một bộ đặc trưng tĩnh không đổi nên không nhân bản các cột hằng như tọa độ/tỉnh để kỳ vọng chúng tự tạo thông tin biến thiên theo thời gian.

### 3.2. Hình dạng mẫu TFT

Lấy 168giờ trước t, chia thành28khối không chồng lấn, mỗi khối6giờ và tính trung bình khí tượng. Đầu vào quá khứ có dạng [số_mẫu,28,F], F thay đổi theo E. Ví dụ E0 có F=11, E7 có F=44.

Decoder nhận12bước lịch tương lai ×4biến lịch, tính từ timestamp. Đầu lượng mưa trả [số_mẫu,3]. Đầu sự kiện trả [số_mẫu,3,3]: ba ngày × ba ngưỡng16/50/100mm. Dữ liệu lịch tương lai đã biết không phải khí tượng tương lai đã quan trắc; không đưa ERA5 giờ đích vào như forecast.

Một mẫu chỉ được giữ khi có đủ168giờ lịch sử E7 và72giờ nhãn hợp lệ/permanent/đủ trọng số; nhãn không vượt ranh giới năm. E0–E7 dùng chung những mẫu này để tránh cấu hình ít biến được chấm trên tập dễ hơn. Kết quả mô tả các mốc đủ dữ liệu, chưa bảo đảm hệ thống xử lý được mọi khoảng thiếu.

### 3.3. Chừng đó dữ liệu có lớn không?

Giai đoạn2015–2023 tạo3.248issue ngày hợp lệ. Số mẫu train lớn nhất trong ba fold là2.536, không phải93×102.984 mẫu độc lập. Đây là lý do cần mô hình đối chứng đơn giản hơn và kiểm tra quá khớp. Các ô đã gộp trước học; nhiều giờ quá khứ trong một mẫu cũng không tạo thêm nhiều nhãn ngày độc lập.

## 4. E0–E7 bổ sung những đặc trưng gì?

B có11biến:7trường bề mặt và4mã lịch. A có12biến khí quyển/đối lưu/mây/ẩmđất. M có20biến hoàn lưu/vận chuyểnẩm. C có1biến SST vùng. Tất cả là nhóm không chồng lấn theo danh mục hiện tại.

Chữ A là tên quy ước của nhóm bổ sung trong thí nghiệm; các biến B/M cũng mô tả khí quyển. “Không chồng lấn” nghĩa là không đưa cùng một cột vào hai nhóm, không có nghĩa các nhóm độc lập về vật lý hay thống kê. Chẳng hạn độ ẩm bề mặt, hơi nước cột và thông lượng ẩm có thể liên quan. Vì vậy thêm nhiều nhóm chưa chắc thêm nhiều thông tin mới; factorial giúp đo từng tổ hợp, chưa tự chứng minh quan hệ nhân quả hoặc mọi tương tác.

{table(['E','Tổ hợp','Số biến','Câu hỏi được kiểm tra'],[
    ['E0','B',11,'Chỉ nền bề mặt/lịch dự đoán được đến đâu?'],
    ['E1','B+A',23,'Thêm khí quyển/đối lưu/mây/đất có giúp không?'],
    ['E2','B+M',31,'Thêm hoàn lưu/gió/ẩm có giúp không?'],
    ['E3','B+C',12,'Thêm riêng SST vùng có giúp không?'],
    ['E4','B+A+M',43,'A và M đi cùng nhau có giúp không?'],
    ['E5','B+A+C',24,'SST có giúp khi đã có A không?'],
    ['E6','B+M+C',32,'SST có giúp khi đã có M không?'],
    ['E7','B+A+M+C',44,'Thêm toàn bộ A/M/C so với nền thì sao?']])}

Các tên thường gặp: t2m là nhiệt độ2m; td2m là điểm sương; rh là độẩm tươngđối; u/v là các thànhphần gió; q là độẩm riêng; tcwv là hơi nước cột; CAPE/CIN mô tả điều kiện đối lưu; IVT mô tả vận chuyển ẩm tích phân. Danh mục đủ44biến, đơn vị và tên cột nằm ở mục13 và sheet DANH_MUC_BIEN.

Trong V1, trường khí tượng thường có tiền tố mean__, ví dụ mean__rh2m_pct. Các biến ctx/SST/chênháp giữ hỗ trợ không gian riêng theo mã dựng dữ liệu: SST lấy một hộp biển, chênháp lấy các vùng/điểm đã khai báo. Nội suy hoặc lặp lại một đại lượng bối cảnh lên93ô không tạo93quantrắc độc lập.

Tên M nên hiểu là hoàn lưu và vận chuyểnẩm, chưa là chỉ số giómùa đã kiểm định. C hiện là SST vùng, chưa phải cả ENSO/MJO. Bổ sung nhóm cũng làm tăng số tham số đầu vào TFT (E0 khoảng121nghìn, E7 khoảng361nghìn), dù giữ hidden32 và ngân sách epoch. Vì vậy đây là đóng góp trong pipeline hiện tại, chưa cô lập hoàn toàn ảnh hưởng thông tin khỏi năng lực mô hình.

## 5. Vì sao chọn TFT; tại sao chưa dùng Informer làm mô hình chính?

### 5.1. Lý do có cơ sở để thử TFT

Bài TFT gốc kết hợp thành phần xử lý tuần tự, attention, lựa chọn biến và gating cho dự đoán nhiều hạn với các loại covariate khác nhau. [Bài TFT](https://arxiv.org/abs/1912.09363).

Đối với nghiên cứu này, nhận định lựa chọn là: đầu ra có ba hạn và đầu vào nhiều nhóm biến, nên thử cơ chế lựa chọn biến/gating là hợp lý; các lớp tuần tự có thể xử lý biến động ở28bước6giờ, attention kết nối lịch sử. Đây là lý do thiết kế và giả thuyết để đo, chưa phải bằng chứng TFT chắc thắng. Bằng chứng giải thích hiện tại chủ yếu là ablation E0–E7; chưa có phân tích attention/VSN được kiểm chứng để suy nguyên nhân vật lý.

### 5.2. TFT đang chạy có phải nguyên bản bài báo không?

Tên đúng là bản TFT vùng điều chỉnh. Mã hiện có VSN/GRN, encoder/decoderLSTM, attention dùng chung value và gating. Bỏ staticcontext cho vùng đơn; dùng đầu lượng mưa không âm với loss Tweedie p1,5 cộng các đầu xác suất BCE. Bài gốc dùng đầu phân vị; bản này không xuất khoảng phân vị dự báo. Không gọi kết quả này là tái lập mọi benchmark của TFT gốc.

Hai loại đầu ra được học cùng nhau nhưng chấm riêng. Đầu xác suất có hậu xử lý để p16≥p50≥p100. Chưa có một phân phối lượng mưa thống nhất bảo đảm mọi tính chất liên hệ giữa đầu mm và đầu xác suất; không suy phân vị từ ba xác suất này.

Staticcovariates không được khai thác như bài toán nhiều tỉnh; futurecovariates hiện chỉ là lịch. Khả năng nhận forecast thật là hướng mở rộng, chưa phải ưu thế thực nghiệm đã đo của mô hình đang báo cáo.

### 5.3. Informer có không phù hợp hoặc không khả thi không?

Bài Informer tập trung giảm chi phí xử lý chuỗi dài bằng ProbSparseattention, distilling và decoder sinh nhiều bước. [Bài Informer, AAAI2021](https://ojs.aaai.org/index.php/AAAI/article/view/17325). Trong mã của tác giả có x_mark_enc/x_mark_dec cho thông tin thời gian; vì vậy không nói Informer hoàn toàn không nhận được lịch hoặc không thể điều chỉnh covariates. [Mã chính thức](https://github.com/zhouhaoyi/Informer2020/blob/main/models/model.py).

Trong thiết kế hiện tại chỉ có28bước encoder,12bước decoder và3đầu lượng mưa ngày. Nhận định của mình là nút thắt chuỗi cực dài chưa rõ ở kích thước này, nên ưu tiên kiểm tra nhóm biến và chất lượng dữ liệu trước khi mở rộng số kiến trúc. Điều đó không chứng minh Informer dự đoán kém hoặc tốn tài nguyên không khả thi.

**Chưa có lần chạy Informer cùng dữ liệu/nhãn/mẫu/hạn/loss/ngân sách trong ma trận này.** Điểm R² của một mô hình cũ ở dữ liệu điểm hoặc benchmark khác không đủ để loại Informer cho bài toán vùng. Nếu giảng viên muốn so kiến trúc, phải đăng ký phép so cùng input, target, calibration, seed và ngân sách phù hợp; chấp nhận cả trường hợp Informer thắng.

### 5.4. Nên viết lý do lựa chọn thế nào?

“Chọn TFT làm kiến trúc thử nghiệm cho nghiên cứu đặc trưng đa biến và nhiều hạn. LightGBM là đối chứng cần thiết để đo giá trị của pipeline neural. Kết quả hiện chưa chứng minh TFT vượt LightGBM; Informer chưa được so cùng giao thức nên giữ là lựa chọn mở rộng.”

## 6. So sánh với LightGBM là so sánh với cái gì và vì sao?

### 6.1. LightGBM là mô hình học, không phải dữ liệu

LightGBM là triển khai gradientboosting trên cây quyết định. Trong cách hiểu đơn giản, nhiều cây được kết hợp, các cây sau góp phần sửa sai số còn lại. Thuật toán học từ tập train, không phải lấy sẵn dự báo mưa của một cơ quan. [Bài LightGBM](https://proceedings.neurips.cc/paper_files/paper/2017/hash/6449f44a102fde848669bdd9eb6b76fa-Abstract.html).

Trong dự án, đặt hai cách học cạnh nhau: TFT học từ chuỗi khí tượng28bước, LightGBM học từ bảng tóm tắt cùng lịch sử. Cả hai dự đoán cùng ba nhãn lượng mưa vùng và chín xác suất sự kiện, trên cùng tập mốc đánh giá.

### 6.2. LightGBM nhận những cột nào?

Với mỗi trường khí tượng của cấu hình E, lấy giá trị trung bình khối6giờ cuối, trung bình24/72/168giờ, và thay đổi giữa hai khối6giờ cuối. Thêm bốn giá trị lịch của khối cuối và48giátrị lịch tương lai (12khối×4). Ví dụ rh850 đưa vào thống kê độẩm trên các khoảng quákhứ; không đưa độẩm thực ở ngày cần dự báo.

Do mở rộng một biến gốc thành nhiều thống kê, E0 có87cột bảng và E7 có252cột, trong khi TFT có11hoặc44biến trên28bước. Cột tăng không có nghĩa lấy thêm nguồn hoặc lấy dữ liệu sau t. Đây là hai biểu diễn khác nhau của cùng nguyên liệu; LightGBM hiện không thấy nguyên vẹn toàn bộ thứ tự28bước như TFT.

Hồi quy dùng Tweedie; phân loại nhị phân dùng mục tiêu sự kiện cho từng ngày/ngưỡng. Mỗi pipeline có3đầu hồi quy và9đầu sự kiện. Cấu hình giữ num_leaves15, tối đa400cây, learningrate0,04, chọn iteration bằng calibration. Đây là đối chứng đã train thật, không phải một đường tham chiếu được vẽ tùy ý.

### 6.3. Vì sao đối chứng này cần thiết?

Với vài nghìn nhãn ngày, một mạng nhiều tham số có thể không hơn cách học trên bảng. Có đối chứng giúp trả lời: lợi ích đến từ nhóm biến mới có xuất hiện ở cả hai họ học không; và mức phức tạp của TFT có mang lại giá trị so với cách học khác không? Đây là nhận định lựa chọn của dự án, không phải định lý LightGBM luôn tốt với dữ liệu nhỏ. Tài liệu thuật toán cũng cảnh báo cây kiểu leafwise có thể quákhớp khi mẫu nhỏ. [LightGBM Features](https://lightgbm.readthedocs.io/en/stable/Features.html).

Hiện là so **hai pipeline dùng cùng nguyên liệu**, chưa phải kiểm định chỉ thay duy nhất kiến trúc: còn khác biểu diễn lịch sử, loss chung/riêng và số lượt seed. Nếu muốn cô lập kiến trúc hơn, có thể bổ sung cây nhận toàn bộ lag hoặc thiết kế các mô hình với cùng biểu diễn và tuningbudget, trong một phép thử tiếp theo đã khóa.

### 6.4. Hai câu hỏi khác nhau phải giữ riêng

E7TFT−E0TFT đo thay đổi đặc trưng trong họ TFT. E7LightGBM−E0LightGBM đo thay đổi đặc trưng trong họ cây. E7TFT so E7LightGBM đo hai pipeline cùng nguồn đầu vào. Nhóm biến hữu ích không bắt buộc mô hình neural phải thắng cây; nghiên cứu đặc trưng vẫn có kết quả nếu đối chứng đơn giản hơn tốt hơn.

Ngoài LightGBM, khí hậu theo tháng/lead học từ train và luôn0 cũng đã được chấm. Tham chiếu khí hậu dùng để kiểm tra mô hình có hơn một xác suất theo mùa đơn giản hay không. Persistence cần nguồn mưa quákhứ; vì nhánh chính hiện không dùng mưa quákhứ, chưa thêm nó rồi gọi là cùng thông tin.

## 7. Metrics là gì, vì sao dùng và đọc thế nào?

Một con số không trả lời được mọi câu hỏi. Dự đoán lượng mm cần chỉ số lượng; xác suất mưa cần chỉ số xác suất; quyết định báo sự kiện cần chỉ số trúng/bỏsót/báonhầm. Tách ba loại này trước khi đọc bảng.

{metric_table}

Các công thức dưới đây tương ứng mã verification/short_range_evaluation đang dùng. Định nghĩa chung có thể tra ở [hướng dẫn kiểm chứng khí tượng](https://www.cawcr.gov.au/projects/verification/Stanski_et_al/VerificationSWBPartI.pdf) và [hướng dẫn chỉ số hồi quy](https://scikit-learn.org/stable/modules/model_evaluation.html). Ví dụ số là minh họa tự tạo, không phải dự báo mới.

### 7.1. RMSE, MAE và Bias: lượng mưa lệch bao nhiêu mm?

Gọi y là nhãn thực dùng để chấm, f là dự đoán mm. RMSE=sqrt(mean((f−y)^2)); MAE=mean(abs(f−y)); Bias=mean(f−y). RMSE nhấn mạnh lỗi lớn, phù hợp để không bỏ qua các ngày sai nặng; MAE cho sai số tuyệt đối trung bình; Bias cho biết thiên hướng dự báo thiếu/thừa.

Ví dụ nhãn [0,10,40], dự đoán [0,15,25] thì sai số [0,+5,−15]. MAE=6,67mm; RMSE=9,13mm; Bias=−3,33mm. Ba chỉ số bổ sung nhau: sai số lớn bị RMSE phạt mạnh; sai số dương và âm có thể triệt tiêu trong Bias nên Bias gần0 chưa đủ tốt.

RMSE20,36mm/24h là mức lỗi bình phương trung bình đã căn trên tập mẫu; không có nghĩa từng ngày chắc chắn sai trong±20,36mm và chưa phải khoảng tin cậy dự báo. Chỉ so RMSE khi cùng đơn vị/nhãn/hạn; không so RMSE ngày của TFT trực tiếp với RMSE giờ của Bi-LSTM.

### 7.2. BS và BSS: xác suất có ích hơn khí hậu không?

Với sự kiện q=16mm/24h, đặt o=1 nếu y≥16, o=0 nếu không; p là xác suất dự báo. BS=mean((p−o)^2), nhỏ hơn tốt hơn. Ví dụ trời thực đạtngưỡng: p=0,8 bị lỗi0,04; p=0,2 bị lỗi0,64. [Định nghĩa Brier score](https://scikit-learn.org/stable/modules/generated/sklearn.metrics.brier_score_loss.html).

BSS=1−BS_mô_hình/BS_khí_hậu. BSS=0 là ngang tham chiếu, >0 tốt hơn, <0 kém hơn. Nếu BS_model=0,06 và BS_ref=0,08 thì BSS=0,25: giảm25%lỗi Brier so tham chiếu, không phải25%ngày được dự đoán đúng. Khí hậu hiện là xác suất theo tháng đích/lead, chỉ học train và làm trơn; reference từng fold được ghép cùng dự đoán khi chấm. [Kiểm chứng xác suất so khí hậu](https://www.cawcr.gov.au/projects/verification/Mason/IntegratedVerificationProcedures.pdf).

Trong số hiện tại, TFT ngày1 E0 BSS={r('tft','E0',1).BSS:.3f}, E7={r('tft','E7',1).BSS:.3f}. ΔBSS≈+0,111 là tăng đơn vị skillscore; lượng giảm BS tương đối so với E0 khoảng{relative_bs:.1f}%. Hai cách biểu diễn này có mẫu số khác nhau; không gọi là tăng11,1%accuracy. BSS không có đơn vị mm.

### 7.3. CSI, POD và FAR: phát hiện sự kiện và báo nhầm thế nào?

Để chuyển xác suất thành quyết định, so p≥τ, với τ chọn trên calibration. Đây là ngưỡng xác suất, khác q=16mm của sự kiện mưa. Có thể τ=0,25 trong một fold; không mặc định mọi ngày/ngưỡng đều0,5. Ngưỡng số cụ thể có trong bảng calibration.

H là số báo có và thực có; M là số bỏsót; F là số báo có nhưng thực không; TN là khô báo khô. CSI=H/(H+M+F), POD=H/(H+M), FAR=F/(H+F). CSI tổng hợp trúng/bỏsót/báonhầm; POD chỉ tập trung bắt được sự kiện; FAR tập trung các lần báo có bị sai. Cần đọc cùng nhau.

Ví dụ100ngày có10ngày sựkiện: H=7,M=3,F=5,TN=85. CSI=7/15=0,467; POD=0,700; FAR=5/12=0,417. Mô hình luôn báo khô cóaccuracy90% nhưng POD=0 và CSI=0; bởi vậy accuracy đơn thuần dễ gây hiểu nhầm khi ngày khô chiếm nhiều. FAR ở đây là falsealarmratio, không phải tỷ lệ trên tất cả ngày khô.

Trong bảng chính, CSI16/POD16/FAR16 tính từ **đầu xác suất đã hiệu chỉnh và τcalibration**. Cột amount_CSI trong Excel tính trực tiếp từ f≥q, có thể khác. Hai đầu ra mm/xác suất được học riêng nên không tự thay một CSI bằng CSI còn lại. Ô trống ở chỉ số không xác định, ví dụ FAR khi chưa có lần báo có, không có nghĩa điểm0.

### 7.4. Reliability, CC và những chỉ số không nên lạm dụng

Reliability xem trong các mốc dự báo khoảng30%, tần suất thực có gần30%không. Đọc cùng số mẫu; một bin vài điểm không đủ để kết luận đã hiệu chỉnh tốt. CC là tương quan biến động lượng mưa, có thể cao dù dự báo thiếu mm có hệ thống, nên không thay RMSE/Bias.

MAPE dễ gặp phép chia cho0 khi không mưa; accuracy dễ bị số ngày khô chi phối. R² có thể là chỉ số bổ sung trong một thiết kế nhất quán, nhưng không dùng điểm R² từ dữ liệu/hạn khác để kết luận về pipeline này. Các bảng đã chạy không có R²; không tạo lại bộ chỉ số lựa chọn sau khi nhìn điểm để làm đẹp kết quả.

### 7.5. Vì sao ưu tiên BSS/CSI mà vẫn phải báo lượng mm?

Giao thức hiện giữ BSS/CSI làm trọng tâm sự kiện theo định hướng nhóm, với RMSE/MAE/Bias bổ sung. BSS dùng toàn bộ xác suất; CSI cho khía cạnh quyết định khi sự kiện không nhiều. Đầu lượng mm vẫn là đầu ra nghiệp vụ quan trọng và phải báo đầy đủ. Một cải thiện BSS không tự chứng minh cải thiện lượng mm hay dự báo cựcđoan tại điểm.

EpochTFT hiện chọn bằng Brier calibration trung bình ba ngày; hồi quy và xác suất dùng loss chung. Nếu nhóm muốn lượngmm/RMSE trở thành tiêu chí chọn mô hình chính trong giai đoạn tiếp, cần khóa lại giao thức trước chấm test, thay vì đổi tiêu chí sau khi thấy mô hình nào đẹp điểm hơn. Loss huấn luyện, metric chọn epoch và metric báo cáo là ba việc liên quan nhưng không đồng nhất.

### 7.6. Vì sao có ngưỡng 16, 50 và 100 mm?

Đây là các ngưỡng sự kiện đã giữ trong giao thức của nhóm để so sánh nhất quán; không tìm ngưỡng đẹp nhất trên evaluation. Ngưỡng 16 mm được dùng cho phép kiểm định chính hiện tại; 50/100 mm giúp đọc các mức hiếm hơn nhưng chỉ có 37/15 sự kiện mỗi lead trong bộ gộp, nên bất định lớn. Các số này áp dụng cho **trung bình vùng trong 24 giờ**, không tự mang ý nghĩa cấp cảnh báo mưa tại trạm hoặc mọi xã. Trung bình vùng có thể dưới ngưỡng trong khi một phần vùng mưa rất lớn. Nếu mục tiêu chuyển sang mức cảnh báo địa phương, phải xác định lại đại lượng và ngưỡng theo mục tiêu đó.

## 8. Chia tập, hiệu chỉnh và kiểm định như thế nào?

{table(['Fold','Train','Calibration','Evaluation','Số mẫu train/cal/eval'],[
    ['F1','2015–2019','2020','2021','1809 /364 /363'],
    ['F2','2015–2020','2021','2022','2173 /363 /349'],
    ['F3','2015–2021','2022','2023','2536 /349 /363']])}

Train học trọng số/scaler và khíhậu. Calibration chọn epoch/iteration, Platt và τquyếtđịnh; không phải một tập kiểm định độc lập vì được dùng cho nhiều bước. Evaluation đo ngoài giai đoạn fit của từng fold; các năm này đã được dự án xem trước, nên gọi là development. Sau khóa thiết kế mới dùng train2015–2023/cal2024/test2025 theo kế hoạch; chưa làm bước đó trong ma trận hiện tại. Không dùng2026 để chọn mô hình chấm2025.

TFT ensemble lấy trung bình dự đoán ba seed trên cùng issue; τensemble chọn trên trung bình xác suất calibration. Không lấy trung bình cácτseed và không họcτtừevaluation. Hiệu chỉnh xác suất được làm riêng từng ngày/ngưỡng; thiếu20sựkiện hoặc20đối-sựkiện thì khôngfitPlatt, giữfallback. Xác suất được chiếu để cóp16≥p50≥p100.

Pooled có1.075issue/lead. Mỗi lead có114sựkiện≥16,37sựkiện≥50 và15sựkiện≥100. Các lead dịch nhãn và có ngày mưa trùng nhau; không cộng thành342đợt độc lập. Tháng3–4 chỉ4–5sựkiện≥16/lead, nên một điểm BSS cao ở nhóm đó cần đọc thận trọng.

## 9. Bảng kết quả đầy đủ 48 dòng chính

Các bảng dưới bao phủ8E×2họ×3ngày; số được gộp trực tiếp trên dự đoán/ref từng dòng2021–2023, toàn năm. RMSE/MAE/Bias đơn vịmm/24h; BSS/CSI/POD/FAR cho ngưỡng16mm. TFT là ensemble ba seed, LGBM một pipeline xác định. Chi tiết từng seed/năm/mùa/ngưỡng50/100 nằm trong Excel, không bị loại vì kết quả kém.

{chr(10).join(full_tables)}

### Đối chứng tối thiểu để đọc mức điểm

{baseline_table}

CSI của khíhậu/luôn0 dùng ngưỡngquyếtđịnh0,5 cốđịnh trong pipelinebaseline; mô hình học cóτchọncalibration. Vì vậy baselinekhíhậu ở bảng này chủ yếu là thamchiếuBSS; không dùng chênhCSI ấy như một phép so quyếtđịnh đã tốiưungưỡng đồngđều. RMSE của khíhậu là trungbìnhmưa theo tháng/lead học train.

### Bảng gọn nhìn đủ ma trận BSS

{table_bss}

## 10. Khoảng tin cậy: “cải thiện” có chắc không?

Điểm cao hơn có thể do tập ngày được đánh giá. Bootstrappaired lấy lại các khối ngày liên tiếp trong từng năm, dùng chung trọng số cho nhãn, dự báo và reference khíhậu. Khối7ngày là chính,14ngày kiểm tra nhạy;20.000lần. Đây là bất định theo thời gian điều kiện trên các mô hình đã fit, không bao phủ hết bất định đào tạo, nhãn hoặc lựa chọn vùng.

Gia đình chính đã khóa là TFT E1−E0, E2−E0, E3−E0 trên ba lead, tổng9phép BSS16. Dùng CI99,44%cho mỗi phép theoBonferroni để xấp xỉ kiểm soát việc thử nhiều phép; CI95%khác dùng mô tả. CIkhông phải xácsuất môhình dựđoán đúng, cũng không là khoảng dựbáo mm cho một ngày.

{ci_table}

Hai CI trên0 thuộc A/M ngày1, cùng kết luận với khối14ngày. M có mức tăng điểm lớn hơn A; chưa có kiểm định trực tiếp M−A. Với C hoặc ngày2–3, CIchứa0 nghĩa là chưa phân biệt chắc lợi ích trong phép thử này, không chứng minh tác dụng bằng0 trong mọi thiết kế.

E7−E0 đo thêm cả A+M+C, là phép so phụ tiếp nối pilot. E4/E6 có điểm vượt E7 ở một số hàng không tự chứng minh tốt nhất ngoài mẫu; chưa có gia đình kiểm định mọi cặp để chọn cấu hình thắng. Các năm development đã xem trước nên kể cả CI điều chỉnh cũng chưa thay holdout độc lập.

## 11. Khả thi đến đâu và còn thiếu gì?

### 11.1. Khả thi về tính toán: đã đo được

Máy RTX 4060 Laptop 8 GB đã chạy 72 TFT/24 LGBM, tiến trình khoảng {run['elapsed_seconds_this_invocation']/60:.1f} phút, có artifact và replay 96/96. Mô hình vùng và chuỗi 28 bước gọn giúp khả thi. Con số này không bao gồm mọi thời gian rà soát nguồn, xây ranh, viết mã và làm nghiên cứu. Dự toán 8–30 giờ GPU ban đầu chưa benchmark không còn dùng như thời gian thực của cấu hình này.

### 11.2. Khả thi để báo cáo nghiên cứu phát triển: có

Có câu hỏi về nhóm biến, đối chứng, cùng mẫu, chia tập theo thời gian, nhiều seed, CI và kết quả âm. Tín hiệu mạnh nhất hiện là xác suất ngày 1 với A/M. Ngày 2 còn thấp; tất cả cấu hình hai họ ở ngày 3 có BSS âm. Kết quả chưa ủng hộ tuyên bố đã dự đoán tốt cả 72 giờ.

### 11.3. Chưa đủ để coi đã hoàn tất nghiên cứu hoặc vận hành

- Chốt ranh và trọng số: 93 ô nguyên ô chưa là giao ranh hành chính chính xác; có thể phải bổ sung ô biên/ngoài hộp. Nếu thay vùng, cần dựng và chạy lại phần phát triển liên quan.
- Kiểm chứng nhãn: sản phẩm vệ tinh có sai số và hiệu chỉnh; checksum/replay/hữu hạn chưa là bằng chứng lượng mưa thật chính xác. Biến mã hóa CIN/has_lfc và proxy cũng cần xác minh ý nghĩa nguồn.
- Khả dụng lúc phát hành: ERA5 hồi cứu chưa bảo đảm đầu vào đến kịp t; nếu chuyển sang forecast thật cần cả thời điểm phát hành và giờ có sẵn, chấm lại cùng thông tin.
- Hạn ngày 3: có thể cần forecast thật, thông tin bổ sung hoặc cách học khác. Ma trận chưa cô lập nguyên nhân nên không kết luận chỉ do TFT/Informer.
- Phạm vi suy rộng: một vùng, ba năm development, ít đợt mưa cực đoan và mô hình khác biểu diễn. Chưa có holdout 2025 hoặc kiểm tra vùng khác.

Đề xuất độc lập: tiếp tục TFT như một mô hình nghiên cứu, giữ LightGBM đối chứng mạnh và đọc rõ cả lượng mm lẫn xác suất. Ưu tiên ranh/nhãn và khóa thiết kế trước final. Informer chỉ cần khi câu hỏi kiến trúc hoặc yêu cầu giảng viên cần phép so ấy; chạy để đo, không chạy với mục tiêu định trước là chứng minh nó thất bại.

## 12. Nội dung có thể dùng để báo cáo và những câu nên sửa

“Nhóm đã hoàn thành ma trận phát triển E0–E7 cho dự đoán trực tiếp mưa trung bình vùng ba khoảng 24 giờ. Có 72 lượt TFT và 24 pipeline LightGBM, cùng mẫu/nhãn và ba fold thời gian. Thêm riêng A/M cải thiện xác suất ngày 1 trong phép thử này sau hiệu chỉnh nhiều phép so; SST và ngày 2–3 chưa có bằng chứng tương tự. TFT chưa thể hiện ưu thế tổng thể so LightGBM. Kết quả đang trên footprint 93 ô xấp xỉ; tiếp theo chốt ranh/nhãn và kiểm định năm giữ kín.”

{table(['Cách nói dễ gây hiểu nhầm','Cách nói đúng với bằng chứng hiện tại'],[
    ['Đã dự báo mưa tốt 72 giờ','Đã chạy đủ ba lead; ngày 3 hiện còn kém khí hậu theo BSS'],
    ['BSS 0,246 là accuracy 24,6%','BSS 0,246 là giảm 24,6% lỗi Brier so khí hậu tham chiếu'],
    ['TFT được chọn vì Informer không khả thi','TFT là giả thuyết kiến trúc đã thử; Informer chưa so cùng giao thức'],
    ['So LightGBM là so với dữ liệu khác','Cùng nguyên liệu/nhãn/mẫu, khác pipeline học và biểu diễn'],
    ['Càng thêm biến càng tốt','Từng nhóm và tổ hợp phải được đo; E7 chưa chắc có điểm tốt nhất'],
    ['93 ô đúng toàn Đà Nẵng mới','Footprint 93 ô hiện còn xấp xỉ, cần chốt giao ranh'],
    ['Chạy 3 seed có 3225 ngày đánh giá độc lập','Mỗi lead vẫn 1075 issue; 3 seed đo ổn định tối ưu hóa'],
    ['Nhóm C không giúp nên khí hậu biển không quan trọng','Chưa có bằng chứng cho một biến SST vùng đang dùng']])}

## 13. Tra cứu đủ 44 đặc trưng đang dùng

Danh mục sau được lấy từ feature_catalog.csv đã công bố. Đây là mô tả theo mã dựng nguồn, chưa phải chứngnhận độc lập rằng từng trường đều chínhxác. Các biến lịch biếttrước; khí tượng chỉ dùng các mốc quákhứ.

{chr(10).join(feature_sections)}

## 14. Cách dùng bảng Excel, tái lập và nguồn tham khảo

Mở BANG_KET_QUA_DAY_DU.xlsx, bắt đầu sheet HUONG_DAN. KET_QUA_48 có đủ 48 dòng chính và mọi cột đã chấm, kể cả BS/CC/amount_CSI. NGUONG_16_50_100 có 144 dòng toàn năm cho hai họ. POOLED_MUA và TUNG_NAM chứa đối chứng/các mùa; các sheet SEED, CI, RELIABILITY và NGUONG_CAL giúp kiểm tra bất định/hiệu chỉnh. Ô trống cho NaN không được tự coi là 0.

Tệp KET_QUA_CHINH_48_DONG.csv thuận tiện để sao chép vào báo cáo. TIEN_DO_E0_E7.csv ghi số biến, cột bảng và số lượt. Build manifest có SHA256 nguồn/đầu ra và mã dựng, giúp đối chiếu tệp. Các bảng là thống kê tổng hợp, không đưa nhãn/dự báo từng issue hoặc checkpoint lên Git.

Các nguồn kỹ thuật đã đối chiếu; phần nhận định áp dụng dự án được phân biệt với kết luận của bài gốc. Không sao chép bài báo/PDF vào repo công khai:

{chr(10).join('- ['+title+']('+url+')' for title,url in REFS)}

Nguồn dự án để kiểm tra: reports/short_range_ablation_2026_10_21/run.json, checkpoint_audit.json, scores_ensemble_pooled.csv, scores_ensemble.csv, scores_seed*.csv, paired_bootstrap.csv; mã src/short_range_dataset.py, tft_regional.py, short_range_evaluation.py, short_range_ablation.py và các scripts tương ứng. Tài liệu này chỉ đọc kết quả đã có, không fit/chấm 2025.
'''
    write('GIAI_THICH_DAY_DU.md',text)
    target=word_copy(OUT/'GIAI_THICH_DAY_DU.md');doc=Document(target)
    doc.core_properties.title='Giải thích đầy đủ TFT vùng 24–72 giờ, E0–E7 và chỉ số'
    doc.sections[0].header.paragraphs[0].text='PBL6 | NGHIÊN CỨU TFT VÙNG | GIẢI THÍCH VÀ KẾT QUẢ ĐẦY ĐỦ'
    for p in doc.sections[0].footer.paragraphs:
        for rr in p.runs:rr.text=rr.text.replace('20/10/2026','21/10/2026')
    for t in doc.tables:
        if len(t.columns)>=7:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for rr in p.runs:rr.font.size=Pt(8)
    doc.add_page_break();doc.add_heading('Hình minh họa mục tiêu và luồng dữ liệu',1)
    for image in ['dinh_nghia_han_du_bao.png','luong_du_lieu_mo_hinh.png']:
        doc.add_picture(str(OUT/image),width=Inches(6.65));doc.add_paragraph(image)
    doc.save(target)
    source_files=['run.json','checkpoint_audit.json','protocol_frozen.json','scores_ensemble_pooled.csv','scores_ensemble.csv',
        'scores_seed_pooled.csv','scores_seed.csv','paired_bootstrap.csv','feature_catalog.csv','reliability16.csv',
        'ensemble_calibration_thresholds.csv','seed_stability.csv']
    manifest=dict(status='built_from_audited_aggregates',date_local='2026-10-21',training_runs_added=0,test2025_read=False,
        tft_runs_verified=72,lgbm_pipelines_verified=24,checkpoint_runs_replayed=96,main_rows=48,workbook_sheets=list(sheets),
        source_sha256={n:sha(SOURCE/n) for n in source_files},code_sha256=sha(Path(__file__)))
    write('README.md','''# Giải thích đầy đủ và bảng kết quả TFT vùng 24–72 giờ

Đã xác nhận đủ E0–E7: 72 TFT, 24 pipeline LightGBM, 96/96 lượt replay. Đây là development trên footprint 93 ô xấp xỉ; năm2025 giữ kín.

- [Tài liệu Word đầy đủ](GIAI_THICH_DAY_DU.docx), [bản Markdown](GIAI_THICH_DAY_DU.md): mục tiêu, ví dụ, dữ liệu, từng cấu hình, TFT/Informer/LightGBM, công thức và cách đọc metrics, bảng48dòng, CI, giới hạn và đủ44biến.
- [Excel kết quả đầy đủ](BANG_KET_QUA_DAY_DU.xlsx): bắt đầu sheetHUONG_DAN, có18sheet về kết quả/ngưỡng/năm/mùa/seed/CI/calibration và nguồn.
- [Bảng chính48dòng CSV](KET_QUA_CHINH_48_DONG.csv), [tiến độ8cấu hình](TIEN_DO_E0_E7.csv).
- [Bộ5slide và lời dẫn ngắn](../short_range_brief_2026_10_21/README.md), [kết quả gốc](../short_range_ablation_2026_10_21/README.md).

Bản này giải thích kết quả đã có, không thêm lượt huấn luyện hoặc thay dữ liệu. Chạy `python scripts/build_short_range_explainer.py` để dựng lại từ các bảng tổng hợp đã commit; cần requirements-reports và openpyxl trong môi trường dự án.
Chạy `python scripts/validate_short_range_explainer.py` để đối chiếu nguồn, hash, số liệu Excel/Markdown và cấu trúc Word; không đọc dữ liệu hoặc chấm lại mô hình.
''')
    files=[p for p in OUT.iterdir() if p.is_file() and p.name!='build_manifest.json']
    manifest['files']={p.name:dict(bytes=p.stat().st_size,sha256=sha(p)) for p in files}
    write('build_manifest.json',json.dumps(manifest,ensure_ascii=False,indent=2))
    print('EXPLAINER_COMPLETE',len(files),'artifacts;',len(sheets),'workbook sheets; 48primaryrows; no new training')


if __name__=='__main__':main()

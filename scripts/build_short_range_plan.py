"""Xuất Word, lịch và protocol đề xuất 24–72h; không huấn luyện mô hình."""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

from build_final_regional_plan import word_copy

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports/short_range_plan_2026_10_21'


def dump(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8', newline='\n')


def main():
    source = OUT / 'KE_HOACH_TFT_24_72H.md'
    inventory = json.loads((OUT / 'preflight_development.json').read_text(encoding='utf-8'))
    groups_path = ROOT / 'reports/regional_research_plan_2026_10_20/feature_groups.json'
    groups = json.loads(groups_path.read_text(encoding='utf-8'))
    approval_path = OUT / 'approval.json'
    approval = json.loads(approval_path.read_text(encoding='utf-8')) if approval_path.exists() else None
    stages = []
    for line in source.read_text(encoding='utf-8').splitlines():
        if line.startswith('| G') and line[3].isdigit():
            stage, task, gate, hours = [c.strip() for c in line.strip('|').split('|')]
            lower, upper = map(int, hours.split('–'))
            stages.append({'stage': stage, 'task': task, 'completion_gate': gate,
                           'hours_min': lower, 'hours_max': upper})
    assert len(stages) == 8
    minimum = sum(row['hours_min'] for row in stages)
    maximum = sum(row['hours_max'] for row in stages)
    assert (minimum, maximum) == (40, 58)
    with (OUT / 'LICH_THUC_HIEN.csv').open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(stages[0]), lineterminator='\n')
        writer.writeheader(); writer.writerows(stages)
    protocol = {
        'document_date': '2026-10-21', 'status': 'proposed_pending_user_approval',
        'training_authorized_by_this_document': False, 'models_fitted_by_this_build': False,
        'objective': 'Direct regional mean precipitation in three consecutive 24-hour periods',
        'source_dataset_id': inventory['dataset_id'], 'source_sha256': inventory['source_sha256'],
        'source_region': inventory['spatial_support'],
        'region_gate': 'Freeze truthful footprint or verified administrative intersection before main matrix',
        'issue_utc_hour': 0, 'issue_stride_hours': 24, 'history_hours': 168,
        'encoder_bin_hours': 6, 'encoder_steps': 28, 'future_calendar_steps': 12,
        'target_hour_intervals': [[0, 24], [24, 48], [48, 72]], 'label': 'y_region_gc_mm',
        'target_unit': 'mm per 24 hours, regional area mean',
        'future_meteorological_inputs_in_primary': False, 'past_rain_inputs_in_primary': False,
        'static_constant_inputs': False, 'validity_mask': 'Common E7 history and full72h permanent complete label',
        'feature_groups_reference': groups_path.relative_to(ROOT).as_posix(),
        'feature_groups': {k: groups[k] for k in ['B', 'A', 'M', 'C']},
        'experiments': groups['experiments'],
        'development_folds': [{'train_start': 2015, 'train_end': y-2, 'cal': y-1, 'eval': y}
                              for y in [2021, 2022, 2023]],
        'final': {'train': [2015, 2023], 'cal': 2024, 'test': 2025, 'experiments': ['E0', 'E7']},
        'no_2026_model_selection_for_2025': True,
        'tft': {'hidden': 32, 'lstm_layers': 1, 'attention_heads': 4, 'dropout': .1,
                'batch': 64, 'optimizer': 'AdamW', 'lr': .001, 'weight_decay': .0001,
                'gradient_clip': 1., 'max_epochs': 40, 'patience': 6, 'seeds': [42, 43, 44],
                'readout': 'Shared daily heads on each group of four decoder states',
                'amount_loss': 'Tweedie p1.5 normalized by train null deviance for each day',
                'event_loss': 'Mean unweighted BCE, weight1',
                'early_stopping': 'Mean raw calibration Brier score at16mm over3days',
                'pilot_tuning_budget': 'At most hidden16/32 on train2015-2019/cal2020 only'},
        'lgbm': {'development_pipelines': 24, 'final_pipelines': 2, 'heads_per_pipeline': 12,
                 'regression_heads': 3, 'event_heads': 9, 'deterministic_seeds_per_pipeline': 1},
        'calibration': {'platt': 'Regularized per day/threshold, calibration only',
                        'min_positive': 20, 'min_negative': 20, 'fallback': 'Identity',
                        'threshold_grid': [.01, .99, .01], 'threshold_fallback': .5,
                        'monotonicity': 'Fixed per-day projection p16>=p50>=p100 after calibration'},
        'event_thresholds_mm_24h': [16, 50, 100], 'primary_event_threshold': 16,
        'metrics': {'primary': ['BSS16', 'CSI16', 'reliability'],
                    'secondary': ['RMSE', 'bias', 'MAE', 'POD', 'FAR', 'rare50/100']},
        'bootstrap': {'block_days': 7, 'sensitivity_block_days': 14,
                      'replicates_final': 20000, 'replicates_pipeline_check': 2000,
                      'seed': 20261021,
                      'development_family': '3 TFT contrasts x3days on pooled eval2021-2023 BSS16',
                      'development_family_size': 9, 'development_simultaneous_ci': 1-.05/9,
                      'holdout_family': 'E7-E0 x3days BSS16', 'holdout_family_size': 3,
                      'holdout_simultaneous_ci': 1-.05/3,
                      'ensemble': 'Mean per-issue predictions over3seeds; seed scores also reported',
                      'reference_alignment': 'Resample climatology reference jointly with event and prediction'},
        'run_budget': {'development_tft': 72, 'final_tft': 6, 'total_tft': 78,
                       'development_lgbm_pipelines': 24, 'final_lgbm_pipelines': 2,
                       'total_lgbm_output_models_upper_bound': 312, 'pilot_extra_runs_not_in_total': True},
        'time_estimate': {'focused_hours': [minimum, maximum], 'contingency_fraction': .2,
                          'rounded_hours_with_contingency': [48, 70],
                          'assumed_hours_per_workday': 5, 'workdays_with_contingency': [10, 14],
                          'gpu_hours_provisional': [8, 30], 'tft_epoch_time_measured': False,
                          'extra_boundary_data_workdays_if_needed': [2, 5],
                          'gpu': 'NVIDIA GeForce RTX4060 Laptop8GB; CUDA verified',
                          'ram_gib_reported_by_os': 15.2,
                          'hardware_confirmation': 'User and local nvidia-smi on2026-10-21'},
        'optional_excluded': ['Informer', 'V2', 'past rain ablation', 'NWP archive', 'quantiles', 'app'],
        'source_package_verify': {'matched': 2233, 'missing': 0, 'mismatched': 0, 'date': '2026-10-21'},
    }
    dump(OUT / 'protocol_proposed.json', protocol)
    if approval:
        protocol['status'] = approval['status']
        protocol['user_approval'] = approval
        protocol['time_estimate']['gpu_hours_before_pilot'] = protocol['time_estimate']['gpu_hours_provisional']
        protocol['time_estimate']['gpu_hours_provisional'] = [1, 3]
        protocol['time_estimate']['tft_epoch_time_measured'] = True
        protocol['time_estimate']['measured_pilot_report'] = '../short_range_pilot_2026_10_21/run.json'
        dump(OUT / 'protocol_proposed.json', protocol)
    word = word_copy(source)
    from docx import Document
    doc = Document(word)
    doc.core_properties.title = 'Kế hoạch nghiên cứu TFT mưa vùng 24–72 giờ'
    doc.core_properties.subject = 'Đề xuất chờ duyệt, không phải kết quả huấn luyện'
    doc.sections[0].header.paragraphs[0].text = 'PBL6 | KẾ HOẠCH TFT 24–72 GIỜ | ĐỀ XUẤT CHỜ DUYỆT'
    if approval:
        doc.core_properties.subject = 'Kế hoạch đã duyệt; pilot đã có, nghiên cứu đầy đủ chưa hoàn tất'
        doc.sections[0].header.paragraphs[0].text = 'PBL6 | KẾ HOẠCH TFT 24–72 GIỜ | ĐÃ DUYỆT'
    from docx.shared import Cm
    for table in doc.tables:
        if len(table.columns) == 4 and len(table.rows) > 1 and table.cell(1, 0).text == 'G0':
            for i, width in enumerate([1.2, 6.3, 7.5, 2.]):
                table.columns[i].width = Cm(width)
                for row in table.rows:
                    row.cells[i].width = Cm(width)
    footer = doc.sections[0].footer.paragraphs[0]
    for run in footer.runs:
        if '20/10/2026' in run.text:
            run.text = run.text.replace('20/10/2026', '21/10/2026')
    doc.save(word)
    artifacts = [source, word, OUT / 'preflight_development.json', OUT / 'protocol_proposed.json',
                 OUT / 'LICH_THUC_HIEN.csv', OUT / 'README.md']
    if approval:
        artifacts.append(approval_path)
    manifest = {'status': 'proposal_pending_user_approval', 'document_date': '2026-10-21',
                'models_trained': False, 'test_2025_rows_analyzed': False,
                'word_tables': len(doc.tables), 'work_stages': len(stages),
                'files': {p.name: {'sha256': hashlib.sha256(p.read_bytes()).hexdigest(),
                                   'bytes': p.stat().st_size} for p in artifacts}}
    dump(OUT / 'build_manifest.json', manifest)
    if approval:
        manifest['status'] = approval['status']
        manifest['models_trained_by_this_build'] = False
        manifest['pilot_report_reference'] = approval['pilot_report']
        dump(OUT / 'build_manifest.json', manifest)
    print(json.dumps({'word': str(word), 'stages': len(stages), 'focused_hours': [minimum, maximum],
                      'status': manifest['status'], 'models_fitted_by_build': False}, ensure_ascii=False))


if __name__ == '__main__':
    main()

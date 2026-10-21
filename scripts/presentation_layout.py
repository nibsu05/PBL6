"""Small shared scene renderer for editable PowerPoint and HTML slides."""
from __future__ import annotations

import base64
import html
import json
from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

NAVY = "123047"
TEAL = "087E8B"
RED = "AD3434"
MUTED = "536877"
BG = "F7F9FC"
LIGHT = "E9F3F5"


def rich(value):
    return [(part[2:-2], True) if part.startswith("**") and part.endswith("**")
            else (part, False) for part in re.split(r"(\*\*.*?\*\*)", str(value)) if part]


class Deck:
    def __init__(self):
        self.slides = []

    def slide(self, title, section="NỘI DUNG CHÍNH", source="", note=""):
        item = dict(number=len(self.slides)+1, title=title, section=section, source=source,
                    note=note, shapes=[])
        self.slides.append(item)
        self.text(item,.55,.18,11,.22,section,11,True,TEAL)
        self.text(item,12.05,.18,.7,.22,f"{item['number']:02d}",11,False,MUTED,"right")
        self.text(item,.55,.60,12.25,.82,title,29,True)
        self.text(item,.55,6.92,12.25,.25,source or "PBL6 • Development 2021–2023 • Vùng 93 ô xấp xỉ • Test 2025 giữ kín",10.5,False,MUTED)
        return item

    def text(self, slide, x, y, w, h, value, size=21, bold=False, color=NAVY, align="left", url=None):
        shape=dict(kind="text",x=x,y=y,w=w,h=h,text=str(value),size=size,bold=bold,color=color,align=align)
        if url:shape["url"]=url
        slide["shapes"].append(shape)
        return shape

    def rect(self, slide, x,y,w,h,fill=LIGHT):
        slide["shapes"].append(dict(kind="rect",x=x,y=y,w=w,h=h,fill=fill))

    def card(self, slide,x,y,w,h,title,body,size=21):
        self.rect(slide,x,y,w,h)
        self.text(slide,x+.17,y+.14,w-.34,.55,title,22,True,TEAL)
        self.text(slide,x+.17,y+.87,w-.34,h-1.00,body,size)

    def band(self,slide,value,color=TEAL,y=6.12,h=.60,size=19):
        self.rect(slide,.55,y,12.25,h,"EAF1F5")
        self.text(slide,.70,y+.10,11.95,h-.12,value,size,True,color)

    def table(self, slide, headers, rows, x=.55,y=1.65,w=12.25,row_h=.48,widths=None,
              size=19,styles=None,tag=None):
        widths=widths or [w/len(headers)]*len(headers)
        assert abs(sum(widths)-w)<1e-6
        shape=dict(kind="table",x=x,y=y,w=w,h=row_h*(len(rows)+1),headers=headers,
                   rows=[[str(v) for v in row] for row in rows],row_h=row_h,widths=widths,size=size,
                   styles=styles or {})
        if tag:shape["tag"]=tag
        slide["shapes"].append(shape)
        return shape

    def image(self,slide,path,x,y,w,h):
        slide["shapes"].append(dict(kind="image",path=str(path),x=x,y=y,w=w,h=h))

    def save(self, pptx_path, html_path, scene_path):
        prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
        pages=[]
        for slide in self.slides:
            native=prs.slides.add_slide(prs.slide_layouts[6]);native.background.fill.solid()
            native.background.fill.fore_color.rgb=RGBColor.from_string(BG)
            native.notes_slide.notes_text_frame.text=slide["note"]
            elements=[]
            for spec in slide["shapes"]:
                x,y,w,h=[spec[k] for k in ("x","y","w","h")]
                pos=f"left:{x*96:.3f}px;top:{y*96:.3f}px;width:{w*96:.3f}px;height:{h*96:.3f}px;"
                if spec["kind"]=="rect":
                    sh=native.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
                    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(spec["fill"]);sh.line.fill.background()
                    elements.append(f'<div class="shape" style="{pos}background:#{spec["fill"]}"></div>')
                elif spec["kind"]=="image":
                    native.shapes.add_picture(spec["path"],Inches(x),Inches(y),width=Inches(w),height=Inches(h))
                    raw=base64.b64encode(Path(spec["path"]).read_bytes()).decode("ascii")
                    elements.append(f'<img class="shape" style="{pos}" src="data:image/png;base64,{raw}">')
                elif spec["kind"]=="text":
                    sh=native.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
                    tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                    tf.vertical_anchor=MSO_ANCHOR.TOP
                    for pi,line in enumerate(spec["text"].split("\n")):
                        p=tf.paragraphs[0] if pi==0 else tf.add_paragraph()
                        p.alignment={"left":PP_ALIGN.LEFT,"center":PP_ALIGN.CENTER,"right":PP_ALIGN.RIGHT}[spec["align"]]
                        p.space_before=p.space_after=Pt(0);p.line_spacing=1.12
                        for text,bold in rich(line):
                            r=p.add_run();r.text=text;r.font.name="Arial";r.font.size=Pt(spec["size"])
                            r.font.bold=bold or spec["bold"];r.font.color.rgb=RGBColor.from_string(spec["color"])
                            if spec.get("url"):r.hyperlink.address=spec["url"]
                    content="".join(("<b>"+html.escape(t)+"</b>") if b else html.escape(t) for t,b in rich(spec["text"]))
                    if spec.get("url"):content=f'<a href="{html.escape(spec["url"],quote=True)}">{content}</a>'
                    elements.append(f'<div class="shape tx" style="{pos}font-size:{spec["size"]*4/3:.3f}px;font-weight:{700 if spec["bold"] else 400};color:#{spec["color"]};text-align:{spec["align"]}">{content}</div>')
                elif spec["kind"]=="table":
                    all_rows=[spec["headers"]]+spec["rows"]
                    sh=native.shapes.add_table(len(all_rows),len(all_rows[0]),Inches(x),Inches(y),Inches(w),Inches(h))
                    for ci,width in enumerate(spec["widths"]):sh.table.columns[ci].width=Inches(width)
                    rendered=[]
                    for ri,row in enumerate(all_rows):
                        sh.table.rows[ri].height=Inches(spec["row_h"])
                        parts=[]
                        for ci,value in enumerate(row):
                            style=spec["styles"].get(f"{ri},{ci}",{})
                            bold=ri==0 or style.get("bold",False)
                            color="FFFFFF" if ri==0 else style.get("color",NAVY)
                            fill=NAVY if ri==0 else style.get("fill","FFFFFF" if ri%2 else "EDF2F7")
                            cell=sh.table.cell(ri,ci);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(fill)
                            cell.margin_left=cell.margin_right=Inches(.055)
                            cell.margin_top=cell.margin_bottom=Inches(.035)
                            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
                            cell.text=str(value);p=cell.text_frame.paragraphs[0];p.alignment=PP_ALIGN.CENTER
                            p.space_before=p.space_after=Pt(0);p.line_spacing=1.0
                            for r in p.runs:
                                r.font.name="Arial";r.font.size=Pt(spec["size"]);r.font.bold=bold
                                r.font.color.rgb=RGBColor.from_string(color)
                            parts.append(f'<td style="height:{spec["row_h"]*96:.2f}px;width:{spec["widths"][ci]*96:.2f}px;font-weight:{700 if bold else 400};background:#{fill};color:#{color}"><div class="ct">{html.escape(str(value))}</div></td>')
                        rendered.append('<tr>'+''.join(parts)+'</tr>')
                    elements.append(f'<table class="shape tbl" style="{pos}font-size:{spec["size"]*4/3:.2f}px"><tbody>'+''.join(rendered)+'</tbody></table>')
            pages.append(f'<section class="slide" id="s{slide["number"]}" data-slide="{slide["number"]}">'+''.join(elements)+'</section>')
        prs.save(pptx_path)
        css="""*{box-sizing:border-box}body{margin:0;background:#263846;font-family:Arial,sans-serif}a{color:inherit;text-decoration:none}.slide{display:none;position:relative;width:1280px;height:720px;margin:30px auto;background:#F7F9FC;overflow:hidden}.slide.on{display:block}.shape{position:absolute}.tx{white-space:pre-wrap;line-height:1.12}.tbl{border-collapse:collapse;table-layout:fixed}.tbl td{border:1px solid white;text-align:center;padding:3px 5px;line-height:1}.ct{overflow:visible}.all .slide{display:block;margin:0}.all nav{display:none}nav{position:sticky;top:0;background:#123047;color:white;padding:10px;text-align:center;z-index:10}button{font:16px Arial;margin:0 12px}#layout-status{display:none}::-webkit-scrollbar{width:0;height:0}@media print{@page{size:13.333333in 7.5in;margin:0}body{background:white}nav{display:none}.slide{display:block;margin:0;break-after:page}}"""
        js="""let n=1;const all=new URLSearchParams(location.search).has('all');document.body.classList.toggle('all',all);function go(k){n=Math.min(Math.max(k,1),document.querySelectorAll('.slide').length);document.querySelectorAll('.slide').forEach(s=>s.classList.toggle('on',Number(s.dataset.slide)===n));document.querySelector('#page').textContent=n+' / '+document.querySelectorAll('.slide').length}document.addEventListener('keydown',e=>{if(e.key==='ArrowRight')go(n+1);if(e.key==='ArrowLeft')go(n-1)});go(1);document.fonts.ready.then(()=>{let bad=[];document.querySelectorAll('.slide').forEach(s=>{let d=s.style.display;s.style.display='block';s.querySelectorAll('.tx,.ct').forEach(t=>{if(t.classList.contains('tx')&&(t.scrollHeight>t.clientHeight+2||t.scrollWidth>t.clientWidth+2))bad.push({slide:s.dataset.slide,text:t.textContent.slice(0,120),height:t.clientHeight,needed:t.scrollHeight});if(t.classList.contains('ct')){let p=t.parentElement;if(t.scrollWidth>p.clientWidth+2||t.scrollHeight>p.clientHeight+2)bad.push({slide:s.dataset.slide,cell:t.textContent,height:p.clientHeight,needed:t.scrollHeight})}});s.style.display=d});document.querySelector('#layout-status').textContent=JSON.stringify({checked:true,overflows:bad});});"""
        page='<!doctype html><html lang="vi"><meta charset="utf-8"><title>PBL6 — Báo cáo mưa vùng</title><style>'+css+'''</style><body><nav><button onclick="go(n-1)">←</button><span id="page"></span><button onclick="go(n+1)">→</button><button onclick="location.search=location.search?'':'?all=1'">Tất cả / từng slide</button></nav>'''+''.join(pages)+'<pre id="layout-status"></pre><script>'+js+'</script></body></html>'
        Path(html_path).write_text(page,encoding="utf8",newline="\n")
        Path(scene_path).write_text(json.dumps(self.slides,ensure_ascii=False,indent=2),encoding="utf8",newline="\n")

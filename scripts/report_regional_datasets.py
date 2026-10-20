"""Xuất báo cáo đọc được và hình từ kết quả pilot cố định; không huấn luyện lại."""
from pathlib import Path
import sys,json
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'reports/regional_danang_v1_v2'


def table(frame):
    def val(x):
        if isinstance(x,(float,np.floating)): return f'{x:.4f}'
        return str(x)
    rows = ['| '+' | '.join(map(str,frame.columns))+' |',
            '| '+' | '.join(['---']*len(frame.columns))+' |']
    rows += ['| '+' | '.join(val(x) for x in row)+' |' for row in frame.itertuples(index=False,name=None)]
    return '\n'.join(rows)


def word_copy(source):
    try:
        from docx import Document
        from docx.shared import Inches,Pt
        from docx.oxml import OxmlElement
    except ImportError:
        return
    d=Document();sec=d.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.65)
    sec.left_margin=sec.right_margin=Inches(.65)
    d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(11)
    import re
    def add(p,s):
        s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',lambda m:m[1]+' ('+m[2]+')',s)
        for j,c in enumerate(re.split(r'\*\*(.*?)\*\*',s)):
            r=p.add_run(c.replace('`',''));r.bold=bool(j%2)
    lines=source.read_text(encoding='utf-8').splitlines();i=0
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                row=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',c) for c in row):rows.append(row)
                i+=1
            t=d.add_table(rows=0,cols=len(rows[0]));t.style='Light Shading Accent 1'
            for ri,row in enumerate(rows):
                for j,txt in enumerate(row):add(t.add_row().cells[j].paragraphs[0],txt) if j==0 else add(t.rows[-1].cells[j].paragraphs[0],txt)
                t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'));continue
        if line.startswith('!['):
            d.add_picture(str(OUT/'rmse_comparison.png'),width=Inches(6.5))
        elif line.startswith('# '):add(d.add_paragraph(style='Title'),line[2:])
        elif line.startswith('## '):add(d.add_paragraph(style='Heading 1'),line[3:])
        elif line.startswith('- '):add(d.add_paragraph(style='List Bullet'),line[2:])
        else:add(d.add_paragraph(),line)
        i+=1
    d.core_properties.title='Hai phiên bản dữ liệu vùng Đà Nẵng và kết quả pilot'
    target=source.with_suffix('.docx');d.save(target)
    assert Document(target).tables


def main():
    build=json.loads((OUT/'build_manifest.json').read_text(encoding='utf-8'))
    run=json.loads((OUT/'run.json').read_text(encoding='utf-8'))
    if run['status']!='complete' or run['test_2025_scored']:
        raise ValueError('Pilot chưa xong hoặc vi phạm giữ test2025')
    scores=pd.read_csv(OUT/'scores.csv');boots=pd.read_csv(OUT/'paired_bootstrap.csv')
    means=scores[scores.season=='all']
    models=['OLD_COMMON','V1_COMMON','V2_COMMON','V1_FULL','V2_FULL']
    fig,axes=plt.subplots(1,2,figsize=(12,4.3))
    colors=['#777777','#4477aa','#228833']
    plotmodels=['OLD_COMMON','V1_FULL','V2_FULL']
    for ax,h in zip(axes,[3,6]):
        p=means[means.horizon==h].pivot(index='year',columns='source',values='RMSE')
        xx=np.arange(3)
        for j,(m,color) in enumerate(zip(plotmodels,colors)):
            ax.bar(xx+(j-1)*.24,p[m],width=.23,label=m,color=color)
        ax.set_xticks(xx,p.index.astype(str));ax.set_title(f'Next {h} hours: regional mean rainfall')
        ax.set_ylabel('RMSE (mm, lower is better)');ax.grid(axis='y',alpha=.2);ax.set_axisbelow(True)
    axes[0].legend(fontsize=8)
    fig.suptitle('Exploratory LightGBM pilot | same GSMaP Gauge target | 2025 held out',fontsize=11)
    fig.tight_layout();fig.savefig(OUT/'rmse_comparison.png',dpi=180)
    svg = OUT/'rmse_comparison.svg'
    fig.savefig(svg);plt.close(fig)
    svg.write_text('\n'.join(line.rstrip() for line in svg.read_text(encoding='utf-8').splitlines())+'\n',encoding='utf-8')
    pivot=means[means.source.isin(models)].pivot(index=['year','horizon'],columns='source',values='RMSE').reset_index()
    pivot=pivot[['year','horizon','OLD_COMMON','V1_COMMON','V2_COMMON','V1_FULL','V2_FULL']]
    result=pivot.rename(columns={'year':'Năm','horizon':'Tổng giờ'})
    ci=boots[(boots.metric=='RMSE')&(boots.source_a=='V2_FULL')].copy()
    ci=ci[['year','horizon','diff','ci_low','ci_high']].rename(columns={'year':'Năm','horizon':'Tổng giờ','diff':'V2−V1','ci_low':'CI thấp','ci_high':'CI cao'})
    diag=pd.read_csv(OUT/'legacy_field_comparison.csv')
    diag=diag[diag.feature.isin(['mean__t2m_c','mean__sp_hpa','mean__cape','mean__tcwv','mean__swvl1'])]
    diag=diag[['feature','old_mean','v1_mean','correlation']]
    fold=pd.DataFrame([{k:f[k] for k in ['eval','train_start','train_end','cal','n_train','n_cal','n_eval']} for f in run['folds']])
    title='# Hai phiên bản dữ liệu vùng Đà Nẵng: khác biệt và dấu hiệu thực nghiệm'
    text=f'''{title}

Ngày 20/10/2026. Đã tạo hai bộ dữ liệu và chạy thử nghiệm thăm dò (pilot). **Có tín hiệu cải thiện ở một số năm, nhưng chưa ổn định để kết luận V2 luôn tốt hơn V1 hoặc dữ liệu cũ.**

## 1. Hai bộ dữ liệu đã tạo

| Bộ | Nội dung | Quy mô |
| --- | --- | --- |
| V1 | Trung bình diện tích của các biến theo giờ; lịch, nhãn và cờ chất lượng riêng | 102.984 giờ; 49 đặc trưng đầu vào ứng viên; 53 cột kể cả thời gian/nhãn/QC; khoảng 44 MB |
| V2 | Nguyên vẹn V1 + phân bố không gian và tương tác | 102.984 giờ; 138 đặc trưng đầu vào ứng viên; 142 cột kể cả thời gian/nhãn/QC; khoảng 122 MB |

Thời gian: 01/01/2015–30/09/2026, theo UTC. Năm 2026 chưa đủ cả năm. Năm 2025 được tổng hợp cơ học trong bản xuất dữ liệu, **không dùng để huấn luyện, chọn mô hình, chấm điểm hoặc thống kê chẩn đoán pilot**.

Phạm vi là tập 93 ô hiện có của Đà Nẵng mới (gồm Đà Nẵng cũ và Quảng Nam cũ). Mỗi phiên bản tạo một chuỗi theo giờ cho cả vùng. Đầu ra là lượng mưa trung bình vùng, nên hai bộ này phù hợp với bài toán một giá trị cho cả vùng; chúng không cung cấp đầu ra bản đồ mưa từng ô.

- [V1 Parquet](../../data/processed/regional_danang/v1_hourly.parquet).
- [V2 Parquet](../../data/processed/regional_danang/v2_hourly.parquet).
- [Trọng số và tiểu vùng của từng ô](../../data/processed/regional_danang/cell_weights.parquet).
- [Danh mục đặc trưng](feature_catalog.csv), [manifest/nguồn/checksum](build_manifest.json).

V2 thêm 89 cột: 54 thống kê độ lệch chuẩn/phân vị 90%/cực đại của 18 biến; 24 trung bình theo tiểu vùng của 8 biến; 8 chênh lệch ven biển − miền núi; 3 chỉ báo gần đúng về gió, ẩm và địa hình. Phân vị 90% được tính theo diện tích có trọng số.

Tiểu vùng được định trước từ địa hình, không xem nhãn mưa: 23 ô có khoảng cách biển xấp xỉ ≤ 30 km; 26 ô ngoài nhóm này có độ cao trung bình ≥ 600 m; 44 ô nội địa còn lại. Đây là phân tầng kỹ thuật, không phải phân vùng khí hậu chính thức.

## 2. Nhãn và trọng số

Nhãn chung là trung bình diện tích của mưa GSMaP Gauge v8 trên 93 ô. Mỗi hàng t mô tả khoảng [t, t + 1 giờ) UTC. Trung bình giữ đơn vị mm cho khoảng một giờ; không cộng mm của các ô.

Trọng số theo **diện tích toàn ô trên cầu**, tổng diện tích các ô khoảng {build['footprint_full_cells_km2']:,.1f} km². Chưa có phần diện tích giao giữa từng ô và đa giác ranh giới hành chính trong dữ liệu cục bộ; vì vậy đây là trung bình của vùng 93 ô hiện có, chưa là trung bình địa giới Đà Nẵng chính xác. Tổng diện tích các ô không phải diện tích hành chính. Chưa nhân trọng số theo phần đất/biển riêng.

Một biến thiếu ở bất kỳ ô nào thì trung bình/thống kê biến đó là NaN; không tự đổi mẫu không gian hoặc điền 0. Nhãn được tách khỏi các đặc trưng đầu vào. Tám biến tĩnh không được nhân bản thành các cột hằng số cho một vùng; địa hình được dùng để phân tầng và tạo tương tác động với gió/ẩm.

Nguồn gốc vẫn là ERA5/GSMaP/Copernicus DEM. Gộp ô không tạo quan trắc mới hoặc nâng độ phân giải thật của ERA5. Nhãn vẫn là sản phẩm ước lượng; chưa chứng minh mưa thực từng giờ của toàn vùng.

## 3. Khác gì với full_features cũ?

| Nội dung | full_features cũ | V1/V2 |
| --- | --- | --- |
| Không gian | Trộn biến tại một điểm với trung bình các hộp vùng khác nhau | Tổng hợp cùng 93 ô; các biến bối cảnh vùng vẫn giữ ý nghĩa nguồn riêng |
| Cột trong file | 15 biến dữ liệu + datetime | V1 có 49 đầu vào ứng viên; V2 có 138; nhãn/QC tách riêng |
| Nhãn gốc | Mưa Open-Meteo tại tọa độ | Mưa GSMaP Gauge trung bình 93 ô |
| Nguyên liệu | Bề mặt, 850 hPa, CAPE/CIN/TCWV, SST, đất, ONI | Thêm mây, tầng 500 hPa, IVT, chuyển động đứng, lớp biên, chênh áp vùng…; chưa tích hợp ONI |
| Thông tin không gian | Đã rút thành một chuỗi, không còn phân bố từng ô | V1 giữ trung bình; V2 thêm phân bố và tiểu vùng |
| Đơn vị gió/ẩm | Gió 10 m: km/h và góc; q850: kg/kg theo script nguồn | Pilot đổi sang u/v (m/s) và q850 (g/kg) để so |

Ví dụ trung bình 2015–2023 trên các giờ chung, **khác giá trị không đồng nghĩa mới chính xác hơn**:

{table(diag)}

Nhiệt độ vùng thấp hơn và áp suất bề mặt vùng thấp hơn có thể phù hợp khác biệt địa hình/độ cao giữa điểm cũ và vùng rộng. Chưa có quan trắc độc lập để quy chênh lệch thành lỗi của một nguồn.

## 4. Pilot so sánh có kiểm soát

Giữ lịch sử 24 giờ → tổng mưa vùng 3/6 giờ sau, phát hành cách 6 giờ. Chỉ dùng đặc trưng quá khứ, không dùng ERA5 tại giờ đích. Mô hình là LightGBM với hàm mất mát Tweedie. Mỗi đặc trưng được biểu diễn bằng giá trị cuối, trung bình 6 giờ, trung bình 24 giờ và thay đổi sau 6 giờ; biến lịch dùng giá trị cuối. Đây là thử nghiệm thăm dò, chưa phải toàn bộ khung E0–E7 hoặc kết quả Bi-LSTM mới.

{table(fold)}

Mỗi lượt đánh giá dùng cùng hàng huấn luyện/hiệu chỉnh/đánh giá cho mọi cấu hình. Các hàng được lọc chung theo các biến cần dùng; khoảng nhãn 6 giờ không vượt năm. Các năm 2021–2023 đã được khám phá trước đây nên đây là bằng chứng thăm dò, chưa là kiểm định độc lập.

| Tên | Đầu vào | Câu hỏi |
| --- | --- | --- |
| OLD_COMMON | 12 đại lượng khí tượng từ file cũ + 4 biến lịch | Mốc đầu vào cũ khi chấm trên cùng nhãn vùng |
| V1_COMMON | 12 đại lượng tương ứng lấy trung bình vùng + 4 biến lịch | Nguồn/đại diện không gian mới có giúp không? |
| V2_COMMON | V1_COMMON + thống kê không gian từ cùng loại nguyên liệu | Riêng cách biểu diễn không gian có giúp không? |
| V1_FULL | Tất cả đặc trưng khí tượng trung bình vùng và lịch, loại các cột mưa | Thêm nguyên liệu khí tượng |
| V2_FULL | V1_FULL + các đặc trưng không gian không chứa mưa | Thêm biểu diễn không gian trên bộ nguyên liệu rộng |

12 đại lượng chung: nhiệt độ 2 m, độ ẩm tương đối 2 m, áp suất bề mặt, u/v ở 10 m, SST, CAPE, TCWV, độ ẩm đất, u/v/q ở 850 hPa. ONI cũ không dùng vì nguy cơ chứa tháng tương lai; CIN cũ không dùng trong phép so cùng nguyên liệu vì ý nghĩa thiếu/0 chưa khôi phục được. CIN mới có trong FULL. Các cột mưa vẫn được xuất trong V1/V2 nhưng bị loại khỏi **tất cả** mô hình pilot để hạn chế tác động nguồn mưa hồi cứu. Vì vậy OLD_COMMON không phải tái lập nguyên mô hình cũ hoặc toàn bộ 15 cột cũ.

LightGBM dùng một seed cố định và chế độ xác định, không lấy mẫu ngẫu nhiên hàng/biến. Ngân sách tối đa 400 cây, 15 lá, dừng sớm sau 30 vòng không cải thiện trên năm hiệu chỉnh; cùng tham số cho mọi cấu hình. Khoảng tin cậy (CI) không bao gồm bất định do thay kiến trúc, siêu tham số hoặc seed huấn luyện mạng nơ ron.

## 5. Kết quả

RMSE (mm tổng 3/6 giờ), càng thấp càng tốt; toàn năm, cùng mẫu và nhãn:

{table(result)}

![So sánh RMSE](rmse_comparison.png)

Chênh RMSE V2_FULL − V1_FULL và CI 95% từ bootstrap theo khối 5 ngày, 1.000 lần; âm là V2 tốt hơn:

{table(ci)}

Diễn giải:

- So với V1_FULL, V2_FULL giảm RMSE 3 giờ khoảng 3,4% ở năm 2021, tăng khoảng 2,5% ở năm 2022 và giảm khoảng 2,8% ở năm 2023. Chỉ CI năm 2021 nằm hoàn toàn dưới 0; năm 2022/2023 chưa phân biệt chắc.
- V2_COMMON − V1_COMMON cải thiện RMSE năm 2021 ở cả hạn 3 giờ và 6 giờ với CI dưới 0; không lặp lại chắc ở năm 2022/2023. Đây là tín hiệu đáng thử tiếp, chưa xác nhận lợi ích ổn định.
- V1_COMMON − OLD_COMMON có CI chứa 0 ở cả 3 năm và cả 2 hạn. Chuyển từ đại diện cũ sang trung bình vùng chưa chứng minh tăng kỹ năng.
- V1_FULL có RMSE 3 giờ thấp hơn OLD_COMMON ở cả 3 năm về điểm ước lượng, nhưng khác biệt gồm cả thêm biến, thay nguồn/cách xử lý và phạm vi không gian. Không quy toàn bộ lợi ích cho gộp lưới.
- Ở hạn 6 giờ, kết quả dao động hơn; chưa chọn V2 làm cấu hình thắng chung.
- Đối chứng giữ nguyên mưa Gauge quá khứ (persistence) có RMSE 3 giờ thấp hơn các mô hình khí tượng ở cả 3 năm. Nó có thêm thông tin mưa quá khứ và là đối chứng hồi cứu lý tưởng hóa, chưa là nguồn vận hành. Điều này gợi ý lượt tiếp theo nên đo thêm mưa quá khứ với điều kiện thông tin rõ ràng; không kết luận kiến trúc nào thua từ phép so khác đầu vào.
- Bảng theo mùa trong scores.csv cho thấy năm 2022, tháng 5–10, FULL xấu hơn OLD_COMMON. Cần trình bày cả những mùa/năm bất lợi. Nhãn mùa đông vẫn có bất định riêng.

CI là có điều kiện theo các mô hình đã fit, chưa hiệu chỉnh nhiều phép so. Dấu hiệu dương ở một năm chưa đủ để kết luận đóng góp vật lý, kết quả toàn Việt Nam hoặc chất lượng mưa thực độc lập.

## 6. Khuyến nghị sử dụng

Giữ V1 làm bộ nền dễ diễn giải; V2 là nhánh nghiên cứu để kiểm tra giá trị phân bố không gian. Giữ full_features cũ làm đối chứng sau chuẩn hóa đơn vị/thời gian và cùng nhãn vùng. Tiếp tục giữ kín tập kiểm tra năm 2025.

Bước tiếp có căn cứ: chốt đầu ra/hạn nghiên cứu; bổ sung đối chứng mưa quá khứ phù hợp; sau đó dùng hai phiên bản để thêm/bớt nhóm đặc trưng trên cùng mẫu và mô hình. Nếu mục tiêu cuối là trung bình địa giới chính xác, cần sửa trọng số theo phần giao ranh giới trước khi đóng băng dữ liệu.

## 7. Tái lập và kiểm tra

- Đã kiểm tra nguồn: 2.233/2.233 tệp khớp; thiếu 0, sai 0.
- Tạo dữ liệu: `.venv\\Scripts\\python.exe -X utf8 scripts/build_regional_datasets.py`.
- Pilot: `.venv\\Scripts\\python.exe -X utf8 scripts/pilot_regional_datasets.py`.
- Báo cáo/hình: `.venv\\Scripts\\python.exe -X utf8 scripts/report_regional_datasets.py`.
- Đã chạy 22 kiểm thử liên quan, gồm phép tổng hợp, đơn vị, cửa sổ không nhìn tương lai và việc đặc trưng không phụ thuộc nhãn: `tests/test_regional_dataset.py`, `tests/test_verification.py`, `tests/test_dataset_v2.py`.
- [Scores theo năm/mùa](scores.csv), [CI theo cặp](paired_bootstrap.csv), [run/siêu tham số](run.json), [so sánh cột cũ/mới](legacy_field_comparison.csv).

Dữ liệu Parquet, trọng số, checkpoint và dự báo từng dòng nằm ngoài Git. Chỉ code, manifest/danh mục và thống kê tổng hợp được giữ trong báo cáo công khai. Các file gốc không bị ghi đè.
'''
    target=OUT/'README.md';target.write_text(text,encoding='utf-8');word_copy(target)
    print(target)


if __name__=='__main__':main()

"""Create an evidence-based development report without changing the pilot report."""
from pathlib import Path
import ast
import importlib.metadata
import platform
import json
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches as DocInches
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from build_final_regional_plan import word_copy
from report_short_range_pilot import write,table,num
from src.short_range_dataset import feature_groups,file_hash
from src.tft_regional import RegionalTFT

REPORT=ROOT/'reports/short_range_ablation_2026_10_21'
PRIVATE=ROOT/'data/operational/short_range/ablation_2026_10_21'
GROUPS={'E0':'B','E1':'B+A','E2':'B+M','E3':'B+C','E4':'B+A+M','E5':'B+A+C','E6':'B+M+C','E7':'B+A+M+C'}
NAMES={'lgbm':'LightGBM','tft':'TFT (ensemble 3 seed)'}


def make_figures(pooled,seed_scores,ci):
    main=pooled[(pooled.season=='all')&(pooled.threshold_mm==16)]
    fig,axes=plt.subplots(1,2,figsize=(12.4,5.0),sharey=True)
    for ax,model in zip(axes,['lgbm','tft']):
        pivot=main[main.model==model].pivot(index='experiment',columns='lead_day',values='BSS').reindex(GROUPS)
        image=ax.imshow(pivot.to_numpy(),vmin=-.15,vmax=.40,cmap='RdYlGn',aspect='auto')
        for i in range(8):
            for j in range(3):ax.text(j,i,f'{pivot.iloc[i,j]:.3f}',ha='center',va='center',fontsize=11)
        ax.set_xticks(range(3),['Ngày 1','Ngày 2','Ngày 3'])
        ax.set_yticks(range(8),[e+' '+GROUPS[e] for e in GROUPS]);ax.set_title(NAMES[model])
    fig.colorbar(image,ax=axes.ravel().tolist(),fraction=.025,pad=.03,label='BSS16: dương tốt hơn khí hậu')
    fig.suptitle('Toàn bộ E0–E7 • development 2021–2023 • cùng 1.075 issue/lead',fontsize=13)
    fig.subplots_adjust(left=.12,right=.87,bottom=.09,top=.87,wspace=.15)
    fig.savefig(REPORT/'bss_factorial.png',dpi=180);plt.close(fig)

    fig,axes=plt.subplots(1,2,figsize=(12.4,4.2))
    subset=seed_scores[(seed_scores.model=='tft')&(seed_scores.season=='all')&(seed_scores.threshold_mm==16)]
    colors={42:'#2563eb',43:'#ea580c',44:'#15803d'}
    for ax,column,ylabel in zip(axes,['BSS','RMSE'],['BSS16','RMSE (mm/24h)']):
        for seed in [42,43,44]:
            for exp in ['E0','E7']:
                p=subset[(subset.seed==seed)&(subset.experiment==exp)].sort_values('lead_day')
                ax.plot(p.lead_day,p[column],marker='o',color=colors[seed],ls='--' if exp=='E0' else '-',label=f'{exp} seed {seed}')
        if column=='BSS':ax.axhline(0,color='#64748b',lw=.8)
        ax.set(xlabel='Ngày dự báo',ylabel=ylabel,xticks=[1,2,3]);ax.grid(alpha=.2)
    handles,labels=axes[0].get_legend_handles_labels()
    fig.legend(handles,labels,loc='lower center',ncol=3,fontsize=9)
    fig.suptitle('Độ ổn định tối ưu hóa TFT: E0/E7, ba seed; không nhân số mẫu',fontsize=13)
    fig.tight_layout(rect=[0,.17,1,.93]);fig.savefig(REPORT/'seed_stability.png',dpi=180);plt.close(fig)

    primary=ci[(ci.block_days==7)&ci.primary_family].sort_values(['lead_day','contrast'])
    fig,ax=plt.subplots(figsize=(12.4,4.8))
    for i,row in enumerate(primary.itertuples()):
        ax.plot([row.family_adjusted_low,row.family_adjusted_high],[i,i],color='#64748b',lw=2)
        ax.plot([row.ci_low,row.ci_high],[i,i],color='#087e8b',lw=5)
        ax.scatter(row.estimate,i,color='#123047',s=30,zorder=3)
    ax.set_yticks(range(len(primary)),[f'Ngày {r.lead_day} · {r.contrast}' for r in primary.itertuples()])
    ax.invert_yaxis();ax.axvline(0,color='#b91c1c',lw=1);ax.grid(axis='x',alpha=.2)
    ax.set_xlabel('ΔBSS16: dương có lợi; thanh đậm CI95%, thanh mảnh CI99,44% cho 9 phép so')
    fig.suptitle('TFT ensemble: từng nhóm thêm vào nền B • 20.000 bootstrap khối 7 ngày',fontsize=13)
    fig.tight_layout();fig.savefig(REPORT/'primary_family_ci.png',dpi=180);plt.close(fig)

    fig,axes=plt.subplots(1,3,figsize=(12.4,4.3),sharey=True)
    seasonal=pooled[(pooled.threshold_mm==16)&pooled.season.isin(['may_oct','nov_feb','mar_apr'])]
    for ax,lead in zip(axes,[1,2,3]):
        for model,exp,color,ls in [('lgbm','E0','#60a5fa','--'),('lgbm','E7','#1d4ed8','-'),('tft','E0','#fb923c','--'),('tft','E7','#c2410c','-')]:
            p=seasonal[(seasonal.model==model)&(seasonal.experiment==exp)&(seasonal.lead_day==lead)].set_index('season').reindex(['may_oct','nov_feb','mar_apr'])
            ax.plot(range(3),p.BSS,'o',color=color,ls=ls,label=f'{model} {exp}')
        ax.axhline(0,color='#64748b',lw=.8);ax.grid(alpha=.2)
        reference=seasonal[(seasonal.model=='climatology')&(seasonal.lead_day==lead)].set_index('season')
        labels=[f'{label}\n{int(reference.loc[season,"events"])} sự kiện' for season,label
                in [('may_oct','Th5–10'),('nov_feb','Th11–2'),('mar_apr','Th3–4')]]
        ax.set(title=f'Ngày {lead}',xticks=range(3),xticklabels=labels)
    axes[0].set_ylabel('BSS16');h,l=axes[0].get_legend_handles_labels()
    fig.legend(h,l,loc='lower center',ncol=4,fontsize=9)
    fig.suptitle('Tách mùa theo tháng issue giờ Việt Nam • số sự kiện ≥16 mm/24h cạnh trục',fontsize=12)
    fig.tight_layout(rect=[0,.17,1,.92]);fig.savefig(REPORT/'seasonal_bss.png',dpi=180);plt.close(fig)


def slides(claims,primary_count):
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    def box(s,x,y,w,h,text,size=22,bold=False,color='123047'):
        shape=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=shape.text_frame
        tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        for i,line in enumerate(text.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial'
            p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
            p.space_after=Pt(10 if i+1<len(text.split('\n')) else 0)
    def base(title,notes):
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(248,250,252)
        box(s,.55,.3,12.2,.7,title,29,True)
        box(s,.55,7.06,12.2,.22,'PBL6 • 21/10/2026 • vùng 93 ô xấp xỉ • development 2021–2023 • chưa mở 2025',11,color='516473')
        s.notes_slide.notes_text_frame.text=notes
        return s
    s=base('Dự đoán mưa vùng 24–72 giờ: kết quả mở rộng',
        'Bộ mở rộng thay lời mô tả pilot một seed khi trình bày số liệu mới. Pilot gốc giữ nguyên. '
        'Các đầu ra là ba tổng 24 giờ riêng. Đây là development trên vùng 93 ô xấp xỉ; năm 2025 vẫn giữ kín. Chưa có kết luận vận hành.')
    box(s,.7,1.4,12,1.,'7 ngày khí tượng quá khứ → mô hình → 3 tổng mưa 24 giờ',27,True,'087E8B')
    box(s,.7,2.8,12,2.5,'72 lượt TFT: 8 cấu hình × 3 fold × 3 seed\n24 pipeline LightGBM: 8 cấu hình × 3 fold\nCùng 1.075 thời điểm evaluation cho mỗi ngày dự báo',25)
    box(s,.7,5.7,12,.8,'Dự đoán trực tiếp vùng; không dùng ERA5 tương lai hoặc mưa quá khứ.',21,color='516473')
    s=base('Giữ phép so công bằng và có thể kiểm chứng',
        'F1: train 2015–2019, cal 2020, eval 2021; F2: train 2015–2020, cal 2021, eval 2022; '
        'F3: train 2015–2021, cal 2022, eval 2023. Cùng dữ liệu V1, mask E7, lịch sử 168 giờ gộp 6 giờ, '
        'hidden 32 và tối đa 40 epoch; chọn epoch trên calibration. Lấy trung bình dự báo ba seed trên từng issue. '
        'Ngưỡng CSI ensemble chọn từ trung bình dự báo calibration. Cả 96 lượt đều được replay trên calibration và evaluation.')
    box(s,.7,1.4,12,3.4,'B: 11 biến nền và lịch\nA: 12 biến khí quyển, đối lưu, mây, ẩm đất\nM: 20 biến gió, ẩm, vận chuyển ẩm và proxy hoàn lưu\nC: 1 biến SST vùng; chưa phải ENSO/MJO đầy đủ',24)
    box(s,.7,5.4,12,1.,'Mẫu không tăng khi chạy thêm seed.\nNgưỡng quyết định chọn trên calibration; năm 2025 giữ kín.',22,True,'087E8B')
    for title,filename,note in [
        ('Đóng góp đặc trưng: xem đủ E0–E7','bss_factorial.png','Mỗi ô là BSS16 trên cùng 1.075 issue cho mỗi lead, gộp theo lỗi từng dòng và reference từng fold. BSS dương có kỹ năng so khí hậu. Chấm đầy đủ các cấu hình, báo cả kết quả âm. SST vùng chưa đại diện cho ENSO.'),
        ('Kết quả có ổn định qua ba seed không?','seed_stability.png','Các đường mô tả từng seed; seed không phải mẫu khí tượng độc lập. Ensemble lấy trung bình ba seed trên từng issue. Bảng seed_stability.csv có trung bình, độ lệch chuẩn và khoảng min/max. Các đường này không thay CI theo thời gian.'),
        ('Từng nhóm A/M/C: khoảng tin cậy đã hiệu chỉnh','primary_family_ci.png','Gia đình chính có 9 phép so: TFT ensemble E1−E0/E2−E0/E3−E0 trên ba lead, BSS16. Thanh mảnh là CI99,44% theo Bonferroni; thanh đậm là CI95% mô tả. Có 20.000 bootstrap paired khối 7 ngày, kiểm tra độ nhạy 14 ngày. Các năm phát triển đã được xem trước; chưa có xác nhận trên holdout độc lập.'),
        ('Kỹ năng còn phụ thuộc mùa và hạn dự báo','seasonal_bss.png','Tách mùa theo tháng issue ở giờ Việt Nam; không khẳng định đây là phân loại khí hậu tối ưu hoặc cơ chế mưa. Sự kiện 50/100 mm hiếm, cần đọc số đếm events. Phép chấm ngày này không tự thay bằng chứng về nhãn giờ mùa đông của nhánh cũ.')]:
        s=base(title,note)
        s.shapes.add_picture(str(REPORT/filename),Inches(.5),Inches(1.3),width=Inches(12.3))
        box(s,.7,6.52,12,.32,'Bảng đầy đủ và điều kiện diễn giải trong BAO_CAO_MO_RONG.md.',15,color='516473')
    s=base('Điều có thể kết luận từ đợt mở rộng',
        'Đọc nguyên văn từ số liệu sinh báo cáo; không suy nguyên nhân duy nhất do kiến trúc. TFT vàLightGBM là hai pipeline khácdạngbiểudiễnvàloss. '
        f'Có {primary_count}/9 CI điều chỉnh nằm trên 0 ở khối 7 ngày; đọc độ nhạy 14 ngày trong báo cáo. '
        'E7−E0 là phép so phụ, CI95% chưa hiệu chỉnh đa so sánh.')
    box(s,.7,1.4,12,4.8,'\n'.join(claims),23)
    s=base('Bước tiếp theo: làm chắc dữ liệu và kiểm định cuối',
        'Đã hoàn thành ma trận development trên footprint xấp xỉ đã khóa; nghiên cứu còn các bước dữ liệu và kiểm định. '
        'Cần chốt ranh, độ phủ và kiểm chứng nhãn trước mở 2025. Nếu đổi vùng/nhãn phải dựng và chạy lại development. '
        'Chưa có phép đo để kết luận Informer không khả thi. Nếu so kiến trúc phải dùng cùng dữ liệu và hạn. Forecast thật là một nhánh nghiên cứu riêng.')
    box(s,.7,1.4,12,4.7,'1. Chốt ranh và trọng số: 93 ô hiện là vùng xấp xỉ.\n2. Kiểm chứng nhãn ngày với nguồn độc lập phù hợp.\n3. Khóa thiết kế, sau đó mới mở holdout 2025.\n4. Ngày 3 còn yếu: cân nhắc forecast thật trong nhánh riêng.\n5. Hoàn thiện bài viết; giữ cả kết quả âm.',24)
    path=REPORT/'PBL6_TFT_24_72h_mo_rong_2026_10_21.pptx';prs.save(path)
    return path


def main():
    run=json.loads((REPORT/'run.json').read_text(encoding='utf-8'))
    audit=json.loads((REPORT/'checkpoint_audit.json').read_text(encoding='utf-8'))
    evaluation=json.loads((REPORT/'evaluation.json').read_text(encoding='utf-8'))
    if run['status']!='complete' or audit['status']!='passed' or audit['full_runs_replayed']!=96:
        raise ValueError('Complete audited factorial required')
    import torch
    environment=dict(python=platform.python_version(),cuda_build=torch.version.cuda,
        cudnn=torch.backends.cudnn.version(),device=run['device'],gpu=run['gpu'],
        packages={name:importlib.metadata.version(name) for name in
                  ['torch','numpy','pandas','pyarrow','lightgbm','scikit-learn','matplotlib','python-docx','python-pptx']})
    (REPORT/'environment.json').write_text(json.dumps(environment,indent=2),encoding='utf-8',newline='\n')
    group_spec=json.loads((ROOT/'reports/regional_research_plan_2026_10_20/feature_groups.json').read_text(encoding='utf-8'))
    tree=ast.parse((ROOT/'scripts/build_dataset_v2.py').read_text(encoding='utf-8'))
    definitions=next(ast.literal_eval(n.value) for n in tree.body if isinstance(n,ast.AnnAssign)
                     and isinstance(n.target,ast.Name) and n.target.id=='FEATURES')
    units={row[0]:(row[1],row[3]) for row in definitions}
    catalog=[]
    for group in ['B','A','M','C']:
        for native in group_spec[group]:
            unit,description=units[native]
            catalog.append(dict(group=group,native_variable=native,
                v1_column=native if native.startswith(('sin_','cos_')) else 'mean__'+native,
                unit=unit or 'dimensionless',description_from_source_code=description,
                input_availability='known calendar' if native.startswith(('sin_','cos_')) else 'past reanalysis only'))
    if len(catalog)!=44:raise ValueError('Unexpected feature catalogue size')
    pd.DataFrame(catalog).to_csv(REPORT/'feature_catalog.csv',index=False,lineterminator='\n')
    scores=pd.read_csv(REPORT/'scores_ensemble_pooled.csv')
    by_year=pd.read_csv(REPORT/'scores_ensemble.csv')
    seeds=pd.read_csv(REPORT/'scores_seed_pooled.csv')
    ci=pd.read_csv(REPORT/'paired_bootstrap.csv')
    main_scores=scores[(scores.season=='all')&(scores.threshold_mm==16)]
    main_seeds=seeds[(seeds.season=='all')&(seeds.threshold_mm==16)&(seeds.model=='tft')]
    dry=scores[(scores.model=='climatology')&(scores.threshold_mm==16)&(scores.season=='mar_apr')].sort_values('lead_day')
    dry_counts='/'.join(str(int(v)) for v in dry.events)
    year_rows=[]
    annual=by_year[(by_year.season=='all')&(by_year.threshold_mm==16)]
    for year in [2021,2022,2023]:
        for lead in [1,2,3]:
            values=[]
            for model,exp in [('lgbm','E0'),('lgbm','E7'),('tft','E0'),('tft','E7')]:
                value=annual[(annual.year==year)&(annual.lead_day==lead)&(annual.model==model)&(annual.experiment==exp)].iloc[0]
                values.append(num(value.BSS))
            year_rows.append([year,lead,*values])
    def row(model,exp,lead):return main_scores[(main_scores.model==model)&(main_scores.experiment==exp)&(main_scores.lead_day==lead)].iloc[0]
    make_figures(scores,seeds,ci)
    primary=ci[(ci.block_days==7)&ci.primary_family].sort_values(['lead_day','contrast'])
    positive=primary[primary.family_adjusted_low>0]
    robust=[]
    for r in positive.itertuples():
        sensitivity=ci[(ci.block_days==14)&ci.primary_family&(ci.contrast==r.contrast)&(ci.lead_day==r.lead_day)].iloc[0]
        if sensitivity.family_adjusted_low>0:robust.append(f'{r.contrast}, ngày {r.lead_day}')
    gb1,tf1=row('lgbm','E7',1),row('tft','E7',1)
    seed_deltas=[]
    for lead in [1,2,3]:
        for seed in [42,43,44]:
            a=main_seeds[(main_seeds.lead_day==lead)&(main_seeds.seed==seed)&(main_seeds.experiment=='E0')].iloc[0]
            b=main_seeds[(main_seeds.lead_day==lead)&(main_seeds.seed==seed)&(main_seeds.experiment=='E7')].iloc[0]
            seed_deltas.append(dict(lead_day=lead,seed=seed,delta_BSS=float(b.BSS-a.BSS),delta_RMSE=float(b.RMSE-a.RMSE)))
    pd.DataFrame(seed_deltas).to_csv(REPORT/'paired_seed_deltas.csv',index=False,lineterminator='\n')
    day1=[r['delta_BSS'] for r in seed_deltas if r['lead_day']==1]
    claims=[f'E7 ngày 1: BSS LightGBM {gb1.BSS:.3f}; TFT ensemble {tf1.BSS:.3f}.',
        f'TFT, ΔBSS E7−E0 ngày 1 dương ở {sum(v>0 for v in day1)}/3 seed; khoảng [{min(day1):+.3f}; {max(day1):+.3f}].',
        f'Thêm riêng A hoặc M giúp ngày 1: {len(positive)}/9 CI chính hiệu chỉnh trên 0.',
        f'E7 ngày 3: BSS LightGBM {row("lgbm","E7",3).BSS:.3f}; TFT ensemble {row("tft","E7",3).BSS:.3f}.',
        'Đây là bằng chứng development; chưa chứng minh tốt trên holdout hoặc vận hành.']
    tables=[]
    for lead in [1,2,3]:
        rows=[]
        for exp in GROUPS:
            a,b=row('lgbm',exp,lead),row('tft',exp,lead)
            rows.append([exp+' '+GROUPS[exp],num(a.BSS),num(b.BSS),num(a.RMSE),num(b.RMSE),num(b.CSI)])
        tables.append(f'### Ngày {lead}\n\n'+table(['Cấu hình','BSS LGBM','BSS TFT','RMSE LGBM','RMSE TFT','CSI TFT'],rows))
    primary_rows=[[r.contrast,r.lead_day,f'{r.estimate:+.3f}',f'[{r.family_adjusted_low:+.3f}; {r.family_adjusted_high:+.3f}]',
        'Trên 0' if r.family_adjusted_low>0 else ('Dưới 0' if r.family_adjusted_high<0 else 'Chứa 0')] for r in primary.itertuples()]
    secondary=ci[(ci.block_days==7)&(ci.contrast=='E7-E0')&ci.metric.isin(['BSS16','RMSE'])]
    secondary_rows=[[r.model,r.lead_day,r.metric,f'{r.estimate:+.3f}',f'[{r.ci_low:+.3f}; {r.ci_high:+.3f}]'] for r in secondary.itertuples()]
    parameter_rows=[]
    for exp,features in feature_groups().items():
        params=sum(p.numel() for p in RegionalTFT(len(features)).parameters())
        parameter_rows.append([exp,GROUPS[exp],len(features),params])
    marginal=[]
    pairs={'A':[('E1','E0'),('E4','E2'),('E5','E3'),('E7','E6')],
           'M':[('E2','E0'),('E4','E1'),('E6','E3'),('E7','E5')],
           'C':[('E3','E0'),('E5','E1'),('E6','E2'),('E7','E4')]}
    for model in ['lgbm','tft']:
        for group,pair in pairs.items():
            marginal.append([model,group,*[f'{np.mean([row(model,a,d).BSS-row(model,b,d).BSS for a,b in pair]):+.3f}' for d in [1,2,3]]])
    pd.DataFrame(marginal,columns=['model','group','day1','day2','day3']).to_csv(REPORT/'factorial_marginal_bss.csv',index=False,lineterminator='\n')
    tfts=[m for m in run['models'] if m['learner']=='tft']
    epoch_times=[h['seconds'] for m in tfts for h in m['history']]
    pilot_comparison=(f"{audit['seed42_pilot_matches']}/12 pipeline seed 42 E0/E7 tái lập kết quả pilot cũ."
        if audit.get('pilot_comparison_available',False) else
        'Không có dự báo riêng tư của pilot để đối chiếu; kiểm chứng 96 lượt của ma trận vẫn được thực hiện đầy đủ.')
    report=f'''# Mở rộng nghiên cứu dự đoán mưa vùng 24–72 giờ

Ngày 21/10/2026. **Hoàn thành ma trận phát triển E0–E7 trên vùng xấp xỉ đã khóa; chưa phải nghiên cứu cuối hoặc kết quả test 2025.** Bộ này bổ sung [pilot một seed](../short_range_pilot_2026_10_21/README.md), giữ nguyên pilot để truy vết.

## 1. Kết quả mới để báo cáo

{chr(10).join('- '+c for c in claims)}

Đã chạy **72 lượt TFT** (8 cấu hình × 3 fold × 3 seed 42/43/44) và **24 pipeline LightGBM xác định** (8 cấu hình × 3 fold, không nhân seed giả). Đã nạp lại **96/96 lượt**, kiểm tra cả calibration và evaluation, tái lập dự báo/chỉ số từ artifact lưu. {pilot_comparison} Tổng cộng 27 kiểm tra đã đạt, gồm các kiểm tra mới chống tăng giả số mẫu, lệch nhãn/reference và chọn ngưỡng trên evaluation.

## 2. Phần nào giống pilot, phần nào mở rộng

Không đổi nguồn V1, 93 ô, trọng số toàn ô, nhãn GSMaP Gauge, mốc issue 00 UTC, lịch sử 168 giờ gộp 6 giờ, ba tổng mưa 24 giờ riêng. E0–E7 dùng cùng mask E7, cùng nhãn và mẫu. Scaler/khí hậu học từ train; chọn epoch và Platt trên calibration. Decoder chỉ biết lịch; không có ERA5 giờ đích, mưa quá khứ hoặc nhãn trong predictors. Cấu hình hidden 32, batch 64, tối đa 40 epoch, patience 6 được giữ nguyên. Thêm seed và tách A/M/C, không dò siêu tham số dựa trên điểm evaluation.

Vùng hiện tại là **93 ô đầy đủ, khoảng 11.074 km²**, chưa là giao ranh hành chính Đà Nẵng mới chính xác. Trước chạy ma trận đã khóa lựa chọn dùng footprint xấp xỉ và gọi đúng tên trong protocol_frozen.json. Đây là nghiên cứu phát triển có giới hạn phạm vi địa lý; phải giải quyết ranh và kiểm chứng nhãn trước tuyên bố nghiên cứu cuối. Nếu đổi ranh/nhãn phải dựng và chạy lại phần phát triển liên quan.

{table(['Cấu hình','Nhóm','Số biến','Số tham số TFT'],parameter_rows)}

B: bề mặt và lịch; A: khí quyển/đối lưu/mây/ẩm đất; M: gió, ẩm và proxy hoàn lưu/vận chuyển ẩm; C: SST vùng. Các trường khí tượng/đại lượng dẫn xuất là ERA5; nhãn mưa là GSMaP Gauge. C không đại diện đầy đủ ENSO/MJO; M chưa phải chỉ số gió mùa đã kiểm định. Danh mục từng cột trong feature_groups.json của phương án hợp nhất. Cùng kiến trúc/hidden/budget nhưng thêm biến làm tăng số tham số đầu vào; không gọi đây là thí nghiệm giữ tuyệt đối số tham số.

[Danh mục 44 đặc trưng, đơn vị và mô tả lấy từ mã dựng dữ liệu](feature_catalog.csv). Các biến bối cảnh giữ hỗ trợ không gian của nguồn: SST là trung bình một hộp biển; ctx là lân cận nút ERA5; các chênh áp lấy các địa điểm/vùng đã khai báo. Không diễn giải chúng như 93 quan trắc mới tại địa phương. Biến mã hóa như CIN/has_lfc và các proxy cũng cần kiểm tra ý nghĩa vật lý; hữu hạn không đồng nghĩa đã được xác nhận đúng.

F1 train 2015–2019/cal 2020/eval 2021: 1.809/364/363 issue; F2 train 2015–2020/cal 2021/eval 2022: 2.173/363/349; F3 train 2015–2021/cal 2022/eval 2023: 2.536/349/363. Mỗi lead gộp có **1.075 issue, 114 sự kiện ≥16 mm/24h**. Ba seed không tạo 3.225 mẫu thời tiết độc lập. Các lead cũng có ngày nhãn chồng nhau, không cộng thành 342 sự kiện độc lập. Chỉ đọc dữ liệu trước 2024; 2025 tiếp tục giữ kín.

## 3. Kết quả toàn bộ E0–E7

BSS16 so với khí hậu theo tháng đích/lead học từ train từng fold; BSS dương có lợi. RMSE mm/24h thấp hơn có lợi. TFT trong bảng là **trung bình dự báo ba seed trên từng issue**, LightGBM là một pipeline xác định. CSI16 dùng xác suất và ngưỡng chọn từ calibration. Đây là hai pipeline dùng cùng nguyên liệu nhưng khác biểu diễn lịch sử/loss, chưa cô lập tuyệt đối ảnh hưởng kiến trúc. Các giá trị pooled tính trực tiếp trên lỗi và reference từng dòng, không trung bình cơ học BSS của năm hoặc seed.

{chr(10).join(tables)}

Toàn bộ MAE, bias, CSI lượng mm, POD/FAR và ngưỡng 50/100 nằm trong scores_ensemble*.csv. Số lượng sự kiện cực đoan ít, nhất là ≥100 mm; không dùng bảng ít sự kiện để khẳng định dự báo cực đoan tốt.

Thêm toàn bộ biến không bảo đảm cho điểm cao nhất: LightGBM ngày 1 có BSS E4={row('lgbm','E4',1).BSS:.3f}, E7={row('lgbm','E7',1).BSS:.3f}. Đây là so điểm mô tả trên development; chưa có kiểm định E4−E7 riêng để tuyên bố ưu thế chắc chắn hoặc chọn lại cấu hình holdout.

## 4. Độ ổn định qua seed và cách gộp

{table(['Ngày','Seed','ΔBSS E7−E0','ΔRMSE E7−E0'],[[r['lead_day'],r['seed'],f"{r['delta_BSS']:+.3f}",f"{r['delta_RMSE']:+.3f}"] for r in seed_deltas])}

Sau hiệu chỉnh xác suất riêng từng seed, lấy trung bình xác suất và lượng mm trên cùng issue. Ngưỡng CSI ensemble được chọn lại trên **trung bình dự báo calibration**, không lấy trung bình các ngưỡng seed và không dùng nhãn evaluation. Nếu calibration không đủ 20 sự kiện/counter-event thì giữ ngưỡng 0,5. Không fit một Platt mới trên evaluation. Báo cả chỉ số từng seed, mean/std/min/max trong seed_stability.csv; ensemble không thay bảng seed.

## 5. Nhóm A/M/C có bằng chứng cải thiện không?

Gia đình chính đã khóa trước chạy ma trận: TFT ensemble, E1−E0 / E2−E0 / E3−E0 × ba lead, BSS16 gộp development. Bootstrap paired theo khối lịch 7 ngày, resample trong từng năm, 20.000 lần; nhãn, dự báo và khí hậu reference cùng trọng số. Dùng CI99,44% cho mỗi phép so theo Bonferroni 9 phép; CI95% chỉ mô tả. Không đổi gia đình sau khi nhìn điểm.

{table(['So sánh','Ngày','ΔBSS16','CI99,44% cho 9 phép','Diễn giải'],primary_rows)}

Trong ma trận này, A và M khi thêm riêng vào nền B có bằng chứng cải thiện ở ngày 1; M có mức tăng BSS lớn hơn. Chưa có CI chính điều chỉnh nằm trên 0 ở ngày 2–3 hoặc với nhóm C. Việc thêm A/M có lợi riêng không bảo đảm tổ hợp đầy đủ E7 tốt nhất: tương tác, dư thừa thông tin và tối ưu hóa mô hình đều có thể ảnh hưởng; chưa tách được nguyên nhân từ ma trận này. Không suy rằng khí hậu biển không quan trọng chỉ từ một biến SST vùng.

Có {len(positive)}/9 CI điều chỉnh nằm trên 0 ở khối 7 ngày. Các phép còn trên 0 khi dùng khối 14 ngày: **{'; '.join(robust) if robust else 'không có'}**. CI này điều kiện trên các mô hình đã fit, không bao trùm mọi bất định huấn luyện, nguồn nhãn hoặc cách chọn vùng. Những năm development đã được dự án xem trước; điều chỉnh đa so sánh không biến chúng thành holdout độc lập. Khoảng bootstrap percentile hữu hạn cũng không phải bảo đảm tuyệt đối mức sai lầm thống kê.

Đóng góp trung bình của một nhóm qua bốn cặp factorial (điểm ước lượng mô tả, chưa có CI/hiệu chỉnh riêng):

{table(['Pipeline','Nhóm','ΔBSS ngày 1','ΔBSS ngày 2','ΔBSS ngày 3'],marginal)}

Ví dụ nhóm A lấy trung bình E1−E0, E4−E2, E5−E3, E7−E6. Đây là đóng góp trong pipeline và biểu diễn hiện tại; không chứng minh nhân quả khí tượng. Vị trí xếp hạng cấu hình trên development cũng không phải bằng chứng đã chọn được cấu hình tốt nhất ngoài mẫu độc lập.

## 6. E7−E0: phép so phụ tiếp nối pilot

{table(['Pipeline','Ngày','Chỉ số','ΔE7−E0','CI95% mô tả'],secondary_rows)}

Phép so này mô tả cả A+M+C; không dùng nó thay kiểm định riêng từng nhóm. Xem paired_bootstrap.csv để có CSI16 và độ nhạy khối 14 ngày. Độ ổn định seed và CI thời gian trả lời hai câu hỏi khác nhau; không xem ba seed như ba đợt mưa độc lập.

## 7. Mùa, xác suất và diễn giải báo cáo

Đọc thêm từng năm để không che khuất khác biệt bằng số gộp. Bảng dưới giữ E0/E7 đã định trước; toàn bộ tổ hợp có trong scores_ensemble.csv.

{table(['Năm','Ngày','BSS LGBM E0','BSS LGBM E7','BSS TFT E0','BSS TFT E7'],year_rows)}

Đã chấm riêng tháng 5–10, tháng 11–2 và tháng 3–4 theo tháng issue ở giờ Việt Nam; xem seasonal_bss.png và bảng scores_ensemble.csv. Các nhóm mùa là phân tầng giao thức, chưa khẳng định tối ưu theo cơ chế khí hậu Đà Nẵng. Reliability16.csv có số mẫu và tần suất từng bin; không kết luận hiệu chỉnh tốt chỉ từ vài bin ít mẫu.

Tháng 3–4 chỉ có {dry_counts} sự kiện ≥16 mm tương ứng ngày 1/2/3 trên ba năm evaluation; không diễn giải một BSS cao ở nhóm ít sự kiện này như bằng chứng mạnh về dự báo mưa lớn.

Có thể trình bày: “Nhóm đã hoàn thành ma trận phát triển nhánh dự đoán vùng 24–72 giờ, với ba seed TFT và đối chứng LightGBM. Nghiên cứu so đủ các tổ hợp đặc trưng, báo cả kết quả không cải thiện, kèm bootstrap theo thời gian và hiệu chỉnh đa so sánh. Đây là bằng chứng hồi cứu trên footprint xấp xỉ; còn phải chốt ranh/kiểm chứng nhãn và kiểm định trên năm giữ kín.”

Không dùng kết quả này để nói TFT luôn tốt hơn LightGBM, Informer không khả thi, đã dự báo chính xác toàn Đà Nẵng mới, hoặc đã có hệ thống vận hành. Chưa chạy Informer trong phép so. Ngày 3 yếu là kết quả cần báo; muốn kiểm tra dự báo số tương lai phải có nhánh dùng forecast thật đã phát hành và đối chứng cùng thông tin.

## 8. Tái lập, thời gian và giới hạn

Tổng fitting/save/calibration của 96 lượt khoảng {run['total_fitting_seconds']/60:.1f} phút; thời gian tiến trình ma trận {run['elapsed_seconds_this_invocation']/60:.1f} phút. Epoch TFT median {np.median(epoch_times):.2f} giây, khoảng [{min(epoch_times):.2f}; {max(epoch_times):.2f}] giây; peak allocated PyTorch cao nhất {max(m['peak_allocated_mib'] for m in tfts):.0f} MiB. GPU {run['gpu']}. Đánh giá ensemble/bootstrap khoảng {evaluation['seconds']:.1f} giây; replay kiểm chứng {audit['seconds']:.1f} giây. Đây là số đo thực trên mô hình nhỏ/dữ liệu vùng; không đồng nhất với tổng thời gian nghiên cứu.

Artifact/checkpoint/dự báo từng dòng lưu ngoài Git trong data/operational/short_range/ablation_2026_10_21. Run.json có code/source SHA256, artifact SHA256, thứ tự biến, best epoch, lịch sử train và hiệu chỉnh; environment.json ghi phiên bản phần mềm/CUDA. Checkpoint_audit.json ghi replay cả calibration/evaluation; source verify 2.233/2.233 từ lần kiểm tra cùng ngày không thay kiểm chứng nhãn mưa độc lập.

Giới hạn còn lại: vùng xấp xỉ; nhãn vệ tinh; số đợt mưa độc lập ít; ERA5 hồi cứu chưa bảo đảm có sẵn tại issue; calibration vừa chọn epoch vừa fit xác suất; development đã xem trước; số tham số tăng theo số biến; LightGBM nhận thống kê còn TFT nhận chuỗi; bootstrap không gồm mọi nguồn bất định. Không so RMSE mm/24h nhánh này trực tiếp với RMSE giờ của Bi-LSTM; hai nhánh còn khác lịch sử 168/24 giờ.

## 9. Công việc tiếp theo

1. Khóa ranh và trọng số phù hợp đúng câu hỏi “mưa vùng Đà Nẵng”, kiểm chứng nhãn ngày độc lập.
2. Nếu thay vùng/nhãn, dựng lại dữ liệu và chạy lại ma trận development liên quan; giữ báo cáo hiện tại như thí nghiệm footprint xấp xỉ.
3. Khóa thiết kế cuối, sau đó mới train 2015–2023/cal 2024/test 2025 E0/E7 theo kế hoạch; chưa làm trong đợt này.
4. Với ngày 3, cân nhắc nhánh forecast thật khi có dữ liệu lịch sử phát hành phù hợp, không đưa ERA5 giờ đích vào như forecast.
5. Hoàn thiện bài viết và phần phương pháp. Informer/V2/app không cần thêm chỉ để báo cáo ngày mai.

Tài liệu phương pháp: [bài TFT gốc](https://arxiv.org/abs/1912.09363). Mô hình này là bản điều chỉnh vùng, giữ VSN/GRN/LSTM/attention/gating, bỏ static context và dùng đầu Tweedie/BCE thay đầu phân vị; chưa phải tái lập toàn bộ benchmark bài gốc.
'''
    write(REPORT/'BAO_CAO_MO_RONG.md',report)
    path=word_copy(REPORT/'BAO_CAO_MO_RONG.md');doc=Document(path)
    doc.core_properties.title='Nghiên cứu TFT vùng 24–72 giờ: E0–E7, ba seed'
    doc.sections[0].header.paragraphs[0].text='PBL6 | TFT VÙNG 24–72 GIỜ | MA TRẬN PHÁT TRIỂN'
    for p in doc.sections[0].footer.paragraphs:
        for run_text in p.runs:run_text.text=run_text.text.replace('20/10/2026','21/10/2026')
    doc.add_page_break();doc.add_heading('Hình kết quả',1)
    for name in ['bss_factorial.png','seed_stability.png','primary_family_ci.png','seasonal_bss.png']:
        doc.add_picture(str(REPORT/name),width=DocInches(6.65));doc.add_paragraph(name)
    doc.save(path)
    deck=slides(claims,len(positive))
    write(REPORT/'README.md',f'''# Kết quả mở rộng TFT vùng 24–72 giờ

Đã hoàn thành 72 lượt TFT (ba seed, E0–E7, ba fold) và 24 pipeline LightGBM; 96/96 lượt replay cả calibration/evaluation. Giữ nguyên [pilot](../short_range_pilot_2026_10_21/README.md), không mở 2025. Footprint 93 ô xấp xỉ.

- [Báo cáo Word](BAO_CAO_MO_RONG.docx), [bản Markdown](BAO_CAO_MO_RONG.md).
- [8 slide có lời dẫn]({deck.name}), [câu hỏi để chuẩn bị bảo vệ](CAU_HOI_BAO_VE.md).
- [Bảng ensemble gộp](scores_ensemble_pooled.csv), [từng seed](scores_seed_pooled.csv), [độ ổn định seed](seed_stability.csv).
- [44 đặc trưng, đơn vị và mô tả nguồn](feature_catalog.csv).
- [CI paired 20.000 lần, khối 7/14 ngày](paired_bootstrap.csv), [kiểm chứng checkpoint](checkpoint_audit.json).
- [Giao thức đã khóa trước chạy](protocol_frozen.json), [cấu hình và lịch sử](run.json).

{chr(10).join('- '+c for c in claims)}

Dự báo từng dòng, Parquet và checkpoint nằm ngoài Git. Đây là ma trận development trên vùng xấp xỉ; còn ranh/nhãn và holdout để hoàn thiện nghiên cứu.

Tái lập trên đúng dữ liệu và mã đã khóa: chạy lần lượt `python scripts/run_short_range_matrix.py --stage all`, `python scripts/audit_short_range_matrix.py`, `python scripts/evaluate_short_range_matrix.py`, `python scripts/report_short_range_matrix.py`, `python scripts/validate_short_range_matrix.py`. Runner tiếp tục từ artifact đã có nếu checksum khớp; không tự chạy lại hoặc đổi giao thức. Dùng Python trong `.venv` của dự án.
''')
    files=[p for p in REPORT.iterdir() if p.is_file() and p.name!='build_manifest.json']
    manifest=dict(status='built',date='2026-10-21',test2025_read=False,slides=8,
        files={p.name:dict(bytes=p.stat().st_size,sha256=file_hash(p)) for p in files},
        report_code_sha256=file_hash(Path(__file__)))
    (REPORT/'build_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8',newline='\n')
    print('REPORT_COMPLETE',len(files),'artifacts',flush=True)


if __name__=='__main__':main()

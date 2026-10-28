"""Create the pilot report, figures, Word and editable slides from frozen scores."""
from pathlib import Path
import hashlib
import json
import re
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches as DocInches
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

from build_final_regional_plan import word_copy

REPORT=ROOT/'reports/short_range_pilot_2026_10_21'
PRIVATE=ROOT/'data/operational/short_range/pilot_2026_10_21'
LABELS={('climatology','BASE'):'Khí hậu',('zero','BASE'):'Luôn 0',
        ('lgbm','E0'):'LightGBM E0',('lgbm','E7'):'LightGBM E7',
        ('tft','E0'):'TFT E0',('tft','E7'):'TFT E7'}
COLORS={('climatology','BASE'):'#64748b',('zero','BASE'):'#a1a1aa',
        ('lgbm','E0'):'#60a5fa',('lgbm','E7'):'#1d4ed8',
        ('tft','E0'):'#fb923c',('tft','E7'):'#c2410c'}


def write(path,text):
    if path.suffix=='.md':
        parts=re.split(r'(\[[^\]]*\]\([^)]+\)|`[^`]*`|https?://\S+)',text)
        text=''.join(part if part.startswith(('[','`','http://','https://')) else readable(part)
                     for part in parts)
    path.write_text(text,encoding='utf-8',newline='\n')


def readable(text):
    substitutions={'Ngày21':'Ngày 21','ngày21':'ngày 21','ngày22':'ngày 22',
                   'ngày1':'ngày 1','ngày2':'ngày 2','ngày3':'ngày 3',
                   'Ngày1':'Ngày 1','Ngày2':'Ngày 2','Ngày3':'Ngày 3',
                   '93ô':'93 ô','11cột':'11 cột','44cột':'44 cột','seed42':'seed 42',
                   '24–72giờ':'24–72 giờ','24h':'24 h','168h':'168 h','72h':'72 h',
                   '48h':'48 h','6h':'6 h','7ngày':'7 ngày','7giờ':'7 giờ',
                   '6lượt':'6 lượt','6pipeline':'6 pipeline','12lượt':'12 lượt',
                   '12kiểm':'12 kiểm','2.000bootstrap':'2.000 bootstrap','2000lần':'2000 lần',
                   '14ngày':'14 ngày','footprintxấp xỉ':'footprint xấp xỉ',
                   'toànĐàNẵngmới':'toàn địa giới Đà Nẵng mới','TFT vùng24':'TFT vùng 24',
                   'train2015':'train 2015','cal2020':'cal 2020','cal2021':'cal 2021','cal2022':'cal 2022',
                   'eval2021':'eval 2021','eval2022':'eval 2022','eval2023':'eval 2023',
                   'cal2024':'cal 2024','test2025':'test 2025','evaluation2021':'evaluation 2021',
                   'mở2025':'mở 2025','khối7ngày':'khối 7 ngày','khối7 ngày':'khối 7 ngày',
                   'TFT24':'TFT 24','soLightGBM':'so với LightGBM','forecastECMWF':'forecast ECMWF',
                   '20pseudo':'20 pseudo','fitPlatt':'fit Platt','Referencekhíhậu':'Reference khí hậu',
                   'tham chiếukhíhậu':'tham chiếu khí hậu','bỏ staticcontext':'bỏ static context',
                   'attentiongiátrịchung':'attention dùng chung value','đầuquantile':'đầu phân vị',
                   'đầungàymuTweedie':'đầu ngày Tweedie','trong.venv':'trong .venv',
                   '2233/2233':'2.233/2.233','V1footprint':'V1 footprint'}
    for old,new in substitutions.items(): text=text.replace(old,new)
    return text


def table(headers,rows):
    return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |',
                      *['| '+' | '.join(map(str,row))+' |' for row in rows]])


def num(value):
    return f'{value:.3f}' if np.isfinite(value) else 'không xác định'


def figure_scores(scores,column,filename,ylabel):
    fig,axes=plt.subplots(1,3,figsize=(12.4,3.8),sharey=True)
    subset=scores[(scores.season=='all')&(scores.threshold_mm==16)]
    for ax,year in zip(axes,[2021,2022,2023]):
        for key,label in LABELS.items():
            part=subset[(subset.model==key[0])&(subset.experiment==key[1])&(subset.year==year)].sort_values('lead_day')
            style='--' if key[1]=='E0' else '-'
            if key[0]=='zero': style=':'
            ax.plot(part.lead_day,part[column],marker='o',markersize=4,label=label,color=COLORS[key],ls=style)
        if column=='BSS': ax.axhline(0,color='#334155',lw=.7)
        ax.set(title=str(year),xlabel='Ngày dự báo',xticks=[1,2,3]); ax.grid(alpha=.2)
    axes[0].set_ylabel(ylabel)
    handles,labels=axes[0].get_legend_handles_labels()
    fig.legend(handles,labels,loc='lower center',ncol=3,fontsize=9)
    fig.suptitle('Pilot V1 vùng • E0/E7 • seed42 • chưa mở2025',fontsize=12)
    fig.tight_layout(rect=[0,.16,1,.94]); fig.savefig(REPORT/filename,dpi=180); plt.close(fig)


def figure_ci(ci):
    fig,axes=plt.subplots(1,2,figsize=(12.4,3.9))
    main=ci[ci.block_days==7]
    for ax,metric,label in zip(axes,['RMSE','BSS16'],['ΔRMSE (mm/24h), âm có lợi','ΔBSS16, dương có lợi']):
        rows=main[main.metric==metric].sort_values(['lead_day','model']).reset_index(drop=True)
        for i,row in rows.iterrows():
            color='#1d4ed8' if row.model=='lgbm' else '#c2410c'
            ax.plot([row.ci_low,row.ci_high],[i,i],color=color,lw=2)
            ax.scatter(row.estimate,i,color=color,s=28,zorder=3)
        ax.set_yticks(range(len(rows)),[f'{"LightGBM" if r.model=="lgbm" else "TFT"} · ngày{r.lead_day}' for r in rows.itertuples()])
        ax.axvline(0,color='#475569',lw=1); ax.invert_yaxis(); ax.grid(axis='x',alpha=.2); ax.set_xlabel(label)
    fig.suptitle('E7−E0: CI95% thăm dò • bootstrap khối7ngày • 2021–2023 gộp',fontsize=12)
    fig.tight_layout(); fig.savefig(REPORT/'feature_effect_ci.png',dpi=180); plt.close(fig)


def figure_reliability(reliability):
    fig,axes=plt.subplots(1,3,figsize=(12.4,3.7))
    for ax,lead in zip(axes,[1,2,3]):
        ax.plot([0,1],[0,1],color='#94a3b8',ls=':',label='Hoàn hảo')
        for key in [('climatology','BASE'),('lgbm','E7'),('tft','E7')]:
            part=reliability[(reliability.model==key[0])&(reliability.experiment==key[1])&(reliability.lead_day==lead)]
            ax.plot(part.mean_prob,part.obs_freq,'o-',color=COLORS[key],label=LABELS[key],markersize=4)
        ax.set(title=f'Ngày{lead}',xlabel='Xác suất dự báo',xlim=(0,1),ylim=(0,1)); ax.grid(alpha=.2)
    axes[0].set_ylabel('Tần suất mưa vùng ≥16 mm/24h')
    h,l=axes[0].get_legend_handles_labels(); fig.legend(h,l,loc='lower center',ncol=4,fontsize=9)
    fig.suptitle('Reliability trên evaluation2021–2023; mỗi bin cần đọc cùng số mẫu',fontsize=12)
    fig.tight_layout(rect=[0,.12,1,.94]); fig.savefig(REPORT/'reliability16.png',dpi=180); plt.close(fig)


def figure_example(predictions):
    # Fixed calendar period, not selected by error or apparent visual performance.
    frame=predictions[(predictions.year==2023)&(predictions.lead_day==1)].copy()
    frame=frame[(frame.issue>=pd.Timestamp('2023-10-01',tz='UTC'))&(frame.issue<pd.Timestamp('2023-12-01',tz='UTC'))]
    fig,ax=plt.subplots(figsize=(12.4,3.6))
    obs=frame[(frame.model=='climatology')].sort_values('issue')
    ax.plot(obs.issue,obs.obs,color='#0f172a',label='Nhãn GSMaP Gauge vùng',lw=1.6)
    for key in [('climatology','BASE'),('lgbm','E7'),('tft','E7')]:
        part=frame[(frame.model==key[0])&(frame.experiment==key[1])].sort_values('issue')
        ax.plot(part.issue,part.amount,label=LABELS[key],color=COLORS[key],alpha=.85)
    ax.set(ylabel='mm/24h',title='Ví dụ cố định: ngày1, issue tháng10–11/2023; không chọn đoạn theo điểm tốt')
    ax.grid(alpha=.2); ax.legend(ncol=2,fontsize=9); fig.tight_layout()
    fig.savefig(REPORT/'fixed_example_2023.png',dpi=180); plt.close(fig)


def slides(summary,claims):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    dark='123047'; teal='087E8B'; gray='516473'
    def textbox(slide,x,y,w,h,text,size=22,color=dark,bold=False):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=shape.text_frame; tf.word_wrap=True
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        lines=text.split('\n')
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=readable(line)
            p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
            p.space_after=Pt(12 if i+1<len(lines) else 0)
        return shape
    def base(title,notes):
        slide=prs.slides.add_slide(prs.slide_layouts[6]); fill=slide.background.fill; fill.solid(); fill.fore_color.rgb=RGBColor(248,250,252)
        textbox(slide,.55,.3,12.2,.7,title,30,bold=True)
        textbox(slide,.55,7.02,12.2,.28,'PBL6 • pilot21/10/2026 • footprint93ô xấp xỉ • seed42 • eval2021–2023 • chưa mở2025',11,gray)
        slide.notes_slide.notes_text_frame.text=notes
        return slide
    s=base('Dự đoán mưa vùng: ngày 1, ngày 2, ngày 3',
           'Báo cáo sơ bộ cho ngày22/10. Nghiên cứu đặc trưng, không phải app hoặc cảnh báo chính thức. '
           'Ba đầu ra là các tổng24h riêng, không phải tổng tích lũy24/48/72h. Nguồn: run.json, scores.csv, checkpoint_audit.json trong bộ báo cáo.')
    textbox(s,.65,1.4,12,.75,'7 ngày khí tượng quá khứ → TFT / LightGBM → 3 tổng mưa 24 giờ',27,teal,True)
    textbox(s,.75,2.65,11.8,2.6,'Ngày1: [t, t+24h)\nNgày2: [t+24h, t+48h)\nNgày3: [t+48h, t+72h)',25)
    textbox(s,.75,5.55,11.8,.8,'Dự đoán trực tiếp trung bình vùng; decoder chỉ biết lịch, không dùng ERA5 tương lai.',21,gray)
    s=base('Thiết kế pilot và kiểm chứng pipeline',
           'V1 hiện là trung bình diện tích toàn93ô, chưa là giao ranh hành chính chính xác. Chỉ kiểm kê/training2015–2023. '
           'F1 train2015–2019/cal2020/eval2021; F2 train2015–2020/cal2021/eval2022; F3 train2015–2021/cal2022/eval2023. '
           '23 kiểm tra code đạt (12 mới và11 bộ chỉ số chung),12 lượt checkpoint tái lập đầy đủ trên thiết bị gốc. Có sai khác số học nhỏ CPU/GPU ghi trong audit.')
    textbox(s,.7,1.3,12,1.1,'E0: 11 đặc trưng nền     |     E7: 44 đặc trưng\nCùng mask, nhãn, mốc giờ, split và nguyên liệu cho mỗi phép so.',24)
    textbox(s,.7,2.8,12,2.2,'Evaluation2021: 363 mẫu ngày, 33 sự kiện ≥16mm\nEvaluation2022: 349 mẫu ngày, 31 sự kiện ≥16mm\nEvaluation2023: 363 mẫu ngày, 50 sự kiện ≥16mm',23)
    textbox(s,.7,5.4,12,1.,'6 lượt TFT + 6 pipeline LightGBM; một seed42.\nĐã nạp lại và tái lập chỉ số của cả12 lượt.',22,teal,True)
    for title,image,note in [
        ('Lượng mưa: so RMSE theo ngày và năm','rmse_by_year.png','RMSE thấp hơn là tốt hơn; giữ cả khí hậu và luôn0 làm đối chứng. Không so mm/24h với RMSE nhánh6h. RMSE mô tả độ đúng lượng mm, không thay BSS/CSI.'),
        ('Xác suất: ngày1 có tín hiệu, ngày3 còn yếu','bss_by_year.png','BSS16>0 nghĩa là tốt hơn xác suất khí hậu train theo tháng/lead. Platt và ngưỡng quyết định chỉ fit calibration. Theo số gộp, E7 cảhai họ có BSS âm ở ngày3; không suy rằng mọi mô hình chỉ dùng quá khứ đều bất khả thi.'),
        ('Thêm đặc trưng: điểm ước lượng và bất định','feature_effect_ci.png','Δ=E7−E0. RMSE âm có lợi; BSS dương có lợi. CI95% pilot2.000bootstrapkhối7ngày, cố định một seed, chưa hiệu chỉnh đa so sánh. Không gọi hiệu quả từng nhóm A/M/C vì mới có E0/E7.')]:
        s=base(title,note)
        s.shapes.add_picture(str(REPORT/image),Inches(.5),Inches(1.35),width=Inches(12.3))
        textbox(s,.7,6.12,11.9,.55,'Kết quả hồi cứu trên nhãn GSMaP Gauge vùng; đọc cùng CI và giới hạn nhãn.',17,gray)
    s=base('Điều có thể báo cáo và việc tiếp theo',
           'Không có tuyên bố TFT thắng LightGBM hoặc V1 đã đại diện chính xác toàn địa giới. Ngày3 yếu là kết quả cần báo, không phải lý do chỉnh test. '
           'Bước tiếp theo: khóa ranh/nhãn, ba seed, factorialE0–E7, sau đó mới holdout2025. TFT không có Informer đối chứng trong pilot, không thể dùng pilot bác bỏ Informer.')
    textbox(s,.7,1.35,12,3.05,'\n'.join(claims),22)
    textbox(s,.7,4.8,12,1.5,'Còn thiếu: giao ranh chính xác, kiểm chứng nhãn độc lập đầy đủ, nhiều seed và E0–E7.\nGiữ kín2025; chưa có kết luận vận hành hoặc cảnh báo.',22,teal,True)
    s=base('Dự phòng: reliability xác suất mưa ≥16 mm/24h',
           'Đường chéo là hiệu chỉnh hoàn hảo. Mỗi bin có số mẫu khác nhau, số đợt mưa độc lập ít. Xem reliability16.csv để đọc n từng bin. Không đọc đường vài điểm là bằng chứng chắc chắn đã hiệu chỉnh tốt.')
    s.shapes.add_picture(str(REPORT/'reliability16.png'),Inches(.5),Inches(1.4),width=Inches(12.3))
    textbox(s,.7,6.,12,.6,'Bảng số mẫu từng bin nằm trong reliability16.csv; hình không thay khoảng tin cậy.',18,gray)
    path=REPORT/'PBL6_TFT_24_72h_pilot_2026_10_21.pptx'; prs.save(path)
    return path


def main():
    run=json.loads((REPORT/'run.json').read_text(encoding='utf-8'))
    audit=json.loads((REPORT/'checkpoint_audit.json').read_text(encoding='utf-8'))
    assert run['status']=='complete' and audit['status']=='passed' and audit['full_runs_replayed']==12
    scores=pd.read_csv(REPORT/'scores.csv'); pooled=pd.read_csv(REPORT/'scores_pooled.csv')
    ci=pd.read_csv(REPORT/'paired_bootstrap.csv'); reliability=pd.read_csv(REPORT/'reliability16.csv')
    predictions=pd.read_parquet(PRIVATE/'predictions.parquet')
    main_scores=pooled[(pooled.season=='all')&(pooled.threshold_mm==16)]
    def row(model,experiment,lead):
        return main_scores[(main_scores.model==model)&(main_scores.experiment==experiment)&(main_scores.lead_day==lead)].iloc[0]
    figure_scores(scores,'RMSE','rmse_by_year.png','RMSE (mm/24h), thấp hơn tốt hơn')
    figure_scores(scores,'BSS','bss_by_year.png','BSS ≥16mm/24h, cao hơn tốt hơn')
    figure_ci(ci); figure_reliability(reliability); figure_example(predictions)
    gb1,tf1=row('lgbm','E7',1),row('tft','E7',1)
    claims=[f'Ngày1, E7: BSS LightGBM={gb1.BSS:.3f}; TFT={tf1.BSS:.3f}.',
            f'Ngày1: RMSE LightGBM={gb1.RMSE:.2f}; TFT={tf1.RMSE:.2f}mm/24h.',
            f'Ngày3, E7: BSS LightGBM={row("lgbm","E7",3).BSS:.3f}; TFT={row("tft","E7",3).BSS:.3f}.',
            'Mới là pilot một seed; chưa chứng minh TFT tốt hơn đối chứng.']
    pooled_rows=[]
    for lead in [1,2,3]:
        for key in [('climatology','BASE'),('lgbm','E0'),('lgbm','E7'),('tft','E0'),('tft','E7')]:
            r=row(*key,lead)
            pooled_rows.append([lead,LABELS[key],num(r.RMSE),num(r.BSS),num(r.CSI)])
    ci_rows=[[r.model,int(r.lead_day),r.metric,num(r.estimate),f'[{num(r.ci_low)}; {num(r.ci_high)}]']
             for r in ci[ci.block_days==7].itertuples()]
    summary_table=table(['Ngày','Cấu hình','RMSE mm/24h','BSS16','CSI16'],pooled_rows)
    ci_table=table(['Họ','Ngày','Chỉ số ΔE7−E0','Ước lượng','CI95%'],ci_rows)
    epoch_times=[h['seconds'] for m in run['models'] if m['learner']=='tft' for h in m['history']]
    tft_seconds=sum(m['seconds'] for m in run['models'] if m['learner']=='tft')
    peak=max(m['peak_allocated_mib'] for m in run['models'] if m['learner']=='tft')
    findings=[]
    for model in ['lgbm','tft']:
        for lead in [1,2,3]:
            r=ci[(ci.model==model)&(ci.lead_day==lead)&(ci.metric=='BSS16')&(ci.block_days==7)].iloc[0]
            meaning='CI nằm trên0' if r.ci_low>0 else ('CI nằm dưới0' if r.ci_high<0 else 'CI chứa0')
            findings.append(f'- {"LightGBM" if model=="lgbm" else "TFT"}, ngày{lead}: ΔBSS16={r.estimate:+.3f}, CI95%[{r.ci_low:+.3f};{r.ci_high:+.3f}], {meaning}; thăm dò, chưa hiệu chỉnh đa so sánh.')
    report=f'''# Báo cáo pilot dự đoán mưa vùng 24–72 giờ

Ngày21/10/2026, chuẩn bị cho báo cáo ngày22/10. **Kết quả sơ bộ, chưa phải nghiên cứu cuối hoặc kiểm định trên2025.**

## 1. Điều đã thực hiện

Đã triển khai loader vùng168h→ba tổng24h, TFT vùng điều chỉnh và LightGBM. Đã chạy E0/E7 trên ba fold evaluation2021/2022/2023, seed42, cùng mask và nhãn; hoàn tất6lượtTFT và6pipelineLightGBM.23 kiểm tra code đạt (12 mới,11 bộ chỉ số chung). Nạp lại toàn bộ12lượt từ artifact đã lưu tái lập chỉ số evaluation trên thiết bị gốc; không tái huấn luyện trong audit.

{chr(10).join('- '+x for x in claims)}

Đây là kết quả của các mô hình đã chạy, không phải bảng dự toán. Không có mô hình Informer trong phép so này, nên không thể kết luận Informer không khả thi.

## 2. Dữ liệu và giao thức

V1: nhãn GSMaP Gauge trung bình diện tích toàn93ô, footprintxấp xỉ khoảng11.074km²; chưa là giao ranh chính xác toànĐàNẵngmới. Nghiên cứu dự đoán vùng trực tiếp, không dự đoán93ô rồi gộp sau mô hình. Khí tượng ERA5 quá khứ; không dùng ERA5 giờ đích, mưa quá khứ hoặc nhãn làm predictor. E0=11cột nền, E7=44cột B/A/M/C, SST chỉ là một biến vùng, chưa là ENSO/MJO đầy đủ.

Issue00UTC=07giờVN, mỗi ngày; encoder28bước6h của7ngày trước issue; decoder chỉ biết lịch. Y1/Y2/Y3 là các tổng24h riêng [t,t+24),[t+24,t+48),[t+48,t+72). Không phải tổng tích lũy48/72h. Đơn vị mm/24h, trung bình vùng, không phải mức mưa tại điểm hoặc cảnh báo chính thức.

F1train2015–2019/cal2020/eval2021:1809/364/363issue. F2train2015–2020/cal2021/eval2022:2173/363/349. F3train2015–2021/cal2022/eval2023:2536/349/363. Mask chung E7, đủ168h lịch sử và72h nhãn permanent; bỏ cửa sổ nhãn qua ranh năm. Scaler và khí hậu chỉ học train; epoch/iteration, Platt và ngưỡng quyết định dùng calibration. Ngưỡng50/100 quá ít sự kiện nên thường không fitPlatt; fallback được lưu trong metadata. Calibration không phải tập kiểm định độc lập.

Pooled mỗi lead có1075issue và114sự kiện≥16mm. Các lead là các ngày được dịch, không được cộng thành342sự kiện độc lập. Năm2025 không được đọc để phân tích/fit/chấm;2026 không dùng lựa chọn.

## 3. Kết quả gộp evaluation2021–2023

{summary_table}

BSS so với xác suất khí hậu theo tháng đích/lead học từ train, làm trơn20pseudo-observation. BSS>0 có kỹ năng so với reference; <0 kém hơn reference trong phép đo này. CSI16 dùng threshold xác suất chọn trên calibration; không phải CSI từ lượng mm. CSI từ lượng mm, POD/FAR, MAE/bias và các ngưỡng50/100 có trong scores.csv. PooledBSS tính trên toàn bộ lỗi ghép và reference từng fold, không trung bình cơ học ba BSS năm.

## 4. Tác dụng thêm toàn bộ đặc trưng E7 so E0

{ci_table}

{chr(10).join(findings)}

ΔRMSE âm có lợi; ΔBSS dương có lợi. Các CI là pairedbootstrapkhối7ngày,2000lần, gộp theo năm nhưng resample trong từng năm, một seed cố định; có độ nhạy14ngày trong paired_bootstrap.csv. Referencekhíhậu được resample cùng obs/pred. Chưa có hiệu chỉnh đa so sánh hoặc bất định do seed, không dùng bảng này như kiểm định xác nhận của ma trận E0–E7 đầy đủ.

Ngày1 có tín hiệu giá trị từ lịch sử/đặc trưng thêm; ở ngày3 cả hai E7 có BSS âm. TFT chưa tốt hơn LightGBM về cả lượng mm và xác suất. Không suy nguyên nhân duy nhất là kiến trúc hay thiếu forecast; cần thêm phép so có kiểm soát. E0/E7 chỉ đo tổng thể thêm A+M+C, chưa phân biệt nhóm nào tạo lợi ích.

## 5. Hình để trình bày

- rmse_by_year.png: tất cả ba năm và ba lead, có khí hậu/luôn0.
- bss_by_year.png: kỹ năng xác suất riêng theo năm, không chỉ đọc số gộp.
- feature_effect_ci.png: hiệu số E7−E0 và CI.
- reliability16.png, reliability16.csv: độ hiệu chỉnh và số mẫu từng bin.
- fixed_example_2023.png: ví dụ ngày1 cố định issue tháng10–11/2023, không chọn đoạn theo lỗi nhỏ.

## 6. Kiểm tra kỹ thuật và thời gian đo

Tổng chạy pipeline mô hình khoảng{run['elapsed_seconds']:.1f}giây; tổng thời gian6lượtTFT gồm calibration/save khoảng{tft_seconds:.1f}giây. EpochTFT median{np.median(epoch_times):.2f}giây, khoảng[{min(epoch_times):.2f};{max(epoch_times):.2f}]giây; peak memory allocated PyTorch cao nhất{peak:.0f}MiB, không phải tổngVRAM của hệ thống. GPU là{run.get('gpu','CPU')}.

Dự toán8–30giờGPU trước pilot là dự phòng khi chưa có benchmark. Số đo hiện tại cho thấy chuỗi vùng nhỏ chạy nhanh hơn nhiều; ma trận đầy đủ cần cập nhật bằng số run/epoch thực tế, không tiếp tục coi8–30giờ là số đã đo. Thời gian viết code, kiểm tra dữ liệu, ranh, kiểm chứng nhãn và viết nghiên cứu không nằm trong thời gian huấn luyện vài phút này.

Checkpointaudit:12/12lượt tái lập trên thiết bị gốc. CPU/GPU có sai khác số học nhỏ, max lượng mưa{max(r.get('cpu_max_amount_drift',0) for r in audit['results']):.6f}mm/24h, ghi riêng trong audit; không trộn làm kết quả huấn luyện khác. Cấu hình, codeSHA256, sourceSHA256, calibration và bestepoch được lưu. Nguồn2233/2233 đãverify trong lượt lập kế hoạch ngày21/10; không phải kiểm chứng độc lập nhãn mưa.

## 7. Những điều chưa được chứng minh

- Ranh93ô xấp xỉ chưa bảo đảm trung bình toàn địa giới; nhãn vệ tinh chưa có kiểm chứng vùng độc lập đầy đủ. Không đồng nhất “pipeline chạy đúng” với “mưa thật chính xác”.
- Một seed42; chưa có ma trận E0–E7, chưa đánh giá độ ổn định tối ưu hóa, chưa mở test2025.2021–2023 là development đã được dự án xem nhiều lần.
- Chỉ có4–6ngày≥100mm mỗi năm evaluation; không khẳng định dự báo cực đoan tốt từ bảng nhỏ.
- ERA5 hồi cứu không có bảo đảm sẵn tại issue; không phải kết quả vận hành/live, không có forecastECMWF thật trong đầu vào.
- Biến chọn/attention chỉ giúp diễn giải mô hình, không chứng minh quan hệ nhân quả khí tượng.
- Bi-LSTM nhánh6h dùng lịch sử24h, nhánh ngày này dùng lịch sử168h. Không quy khác biệt giữa hai nhánh chỉ cho hạn hoặc kiến trúc khi cửa sổ đầu vào chưa giống nhau.
- TFT nhận chuỗi còn LightGBM nhận thống kê lịch sử từ cùng nguyên liệu; loss/đầu ra là joint so với các đầu riêng. So sánh là giữa hai pipeline, chưa cô lập tuyệt đối tác dụng kiến trúc.

## 8. Nội dung có thể nói trong báo cáo ngày mai

“Nhóm đã thống nhất dự đoán vùng trực tiếp và triển khai pilot nhánh24–72giờ. Từ7ngày khí tượng quá khứ, mô hình dự đoán ba tổng mưa24giờ và xác suất sự kiện. Trên ba năm phát triển, ngày1 có tín hiệu kỹ năng so khí hậu; ngày3 còn yếu. TFT chưa thể hiện ưu thế rõ soLightGBM. Kết quả hiện là pilot một seed trên footprint93ô; bước tiếp theo là chốt ranh/kiểm chứng nhãn, chạy nhiều seed và tách đóng góp từng nhóm đặc trưng, giữ kín2025.”

## 9. Việc tiếp theo

Ưu tiên ranh/QC nhãn, xác nhận protocol đã duyệt, chạy ba seed và factorialE0–E7. Không điều chỉnh thiết kế để làm đẹp điểm2025. Chỉ sau khóa mới train2015–2023/cal2024/test2025E0/E7. Forecastthật, Informer và V2 là nhánh riêng, không bắt buộc thêm để có báo cáo pilot.

Artifact: scores.csv, scores_pooled.csv, paired_bootstrap.csv, reliability16.csv, checkpoint_audit.json, run.json; checkpoint/dự báo từng dòng giữ ngoàiGit tại data/operational/short_range/pilot_2026_10_21. Code: short_range_dataset, tft_regional, short_range_evaluation và các scriptpilot/audit/evaluate/report. Kiến trúc tham khảo bàiTFT https://arxiv.org/abs/1912.09363; bản này điều chỉnh bỏ staticcontext, giữVSN/GRN/LSTM/attentiongiátrịchung/gating, thay đầuquantile bằngđầungàymuTweedie/BCE. Chưa phải tái lập mọi benchmark của bài gốc.
'''
    write(REPORT/'BAO_CAO_PILOT.md',report)
    word=word_copy(REPORT/'BAO_CAO_PILOT.md'); doc=Document(word)
    doc.core_properties.title='Báo cáo pilot dự đoán mưa vùng24–72giờ'
    doc.sections[0].header.paragraphs[0].text='PBL6 | PILOT TFT24–72GIỜ | SỐ LIỆU21/10/2026'
    for p in doc.sections[0].footer.paragraphs:
        for run_text in p.runs: run_text.text=run_text.text.replace('20/10/2026','21/10/2026')
    doc.add_page_break(); doc.add_heading('Hình kết quả pilot',1)
    for filename in ['rmse_by_year.png','bss_by_year.png','feature_effect_ci.png','reliability16.png','fixed_example_2023.png']:
        doc.add_picture(str(REPORT/filename),width=DocInches(6.65)); doc.add_paragraph(filename)
    doc.save(word)
    deck=slides(main_scores,claims)
    write(REPORT/'README.md',f'''# Kết quả pilot24–72giờ để báo cáo ngày mai

Số liệu21/10/2026, chuẩn bị cho ngày22/10. **Pilot một seed, chưa mở2025.**

- [Báo cáo đầy đủ](BAO_CAO_PILOT.md) / [Word](BAO_CAO_PILOT.docx).
- [PowerPoint7slide chỉnh sửa được]({deck.name}):6slide chính,1reliability dự phòng; có ghi chú trình bày.
- [Khung4giờ](PILOT_TRONG_4_GIO.md), [bảng theo năm](scores.csv), [bảng gộp](scores_pooled.csv), [CI](paired_bootstrap.csv).
- [Kiểm tra checkpoint](checkpoint_audit.json), [log/cấu hình](run.json).

{chr(10).join('- '+x for x in claims)}

Đã hoàn tất6lượtTFT và6pipelineLightGBM,23 kiểm tra code đạt,12/12lượt tái lập. Dữ liệu V1 hiện là footprint93ô xấp xỉ; chưa là kết quả toàn địa giới hoặc vận hành. Dự báo từng dòng/checkpoint nằm ngoàiGit. Không dùng pilot để bác bỏInformer hoặc khẳng định TFT chắc chắn tốt hơn.

Chạy lại theo thứ tự: scripts/pilot_short_range.py → scripts/audit_short_range_pilot.py → scripts/evaluate_short_range_pilot.py → scripts/report_short_range_pilot.py, dùng Python trong.venv. XuấtWord/slide cầnpython-docx/python-pptx. Không chạy pilot lại chỉ để đọc báo cáo đã có.
''')
    files=[p for p in REPORT.iterdir() if p.is_file() and p.name!='build_manifest.json']
    manifest=dict(status='completed_exploratory_pilot',date_local='2026-10-21',
                  models_replayed=12,code_checks_passed=23,test2025_read=False,
                  report_slides=7,private_predictions_included=False,
                  files={p.name:dict(sha256=hashlib.sha256(p.read_bytes()).hexdigest(),bytes=p.stat().st_size) for p in files})
    write(REPORT/'build_manifest.json',json.dumps(manifest,ensure_ascii=False,indent=2))
    print('REPORT_READY',str(word),str(deck),flush=True)


if __name__=='__main__': main()

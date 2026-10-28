"""Validate all full-matrix report values and preserve prior/user artifacts."""
from pathlib import Path
import hashlib
import json
import re
import subprocess
import zipfile

import pandas as pd
from pptx import Presentation
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'reports/short_range_complete_2026_10_28'
REPORT=ROOT/'reports/informer_regional_matrix_2026_10_28'
PRIVATE=ROOT/'data/operational/short_range/informer_matrix_2026_10_28'
STEM='PBL6_SO_SANH_3_MO_HINH_DAY_DU'
MODELS=['lgbm','tft','informer'];LABELS=['LightGBM','TFT','Informer']
METRICS=['RMSE','MAE','bias','BS','BSS','CSI','POD','FAR']


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    m=json.loads((OUT/'build_manifest.json').read_text(encoding='utf8'))
    assert m['slides']==49 and m['main_slides']==34 and m['complete_score_rows']==72 and m['missing_score_rows']==0
    for name,digest in m['sources'].items():assert sha(ROOT/name)==digest,name
    for name,meta in m['files'].items():assert sha(OUT/name)==meta['sha256'],name
    assert sha(ROOT/'scripts/build_short_range_complete.py')==m['code_sha256']
    validation=json.loads((REPORT/'validation.json').read_text(encoding='utf8'))
    assert validation['status']=='passed' and validation['informer_checkpoints_replayed']==72
    assert validation['old_e0_e7_unchanged'] and validation['tft_lgbm_unchanged'] and not validation['test2025_read']
    frozen=0;missing_worktree=[]
    for name in ['plan','pilot','ablation','brief','explained','presentation','revised','three_models']:
        d=ROOT/f'reports/short_range_{name}_2026_10_21'
        old=json.loads((d/'build_manifest.json').read_text(encoding='utf8'))
        for filename,meta in old['files'].items():
            p=d/filename
            if p.exists():assert sha(p)==meta['sha256'],str(p)
            else:
                # User renamed/edited their working slide. Verify the published input
                # from Git without restoring, overwriting or committing their edit.
                rel=p.relative_to(ROOT).as_posix()
                assert rel=='reports/short_range_three_models_2026_10_21/PBL6_SO_SANH_3_MO_HINH.pptx'
                raw=subprocess.check_output(['git','-c','safe.directory=D:/PBL6_v2','show','HEAD:'+rel])
                assert hashlib.sha256(raw).hexdigest()==meta['sha256']
                missing_worktree.append(rel)
            frozen+=1
    pilot=ROOT/'reports/informer_regional_pilot_2026_10_21'
    old=json.loads((pilot/'publication_manifest.json').read_text(encoding='utf8'))
    for filename,meta in old['files'].items():assert sha(pilot/filename)==meta['sha256'];frozen+=1
    source=pd.read_csv(REPORT/'scores_ensemble_pooled.csv')
    source=source[(source.season=='all')&(source.threshold_mm==16)]
    data=pd.read_csv(OUT/'KET_QUA_72.csv');assert len(data)==72 and data[METRICS].notna().all().all()
    assert set(data.model)==set(MODELS) and not data.duplicated(['model','experiment','lead_day']).any()
    checked_csv=0
    for row in data.itertuples():
        r=source[(source.model==row.model)&(source.experiment==row.experiment)&(source.lead_day==row.lead_day)].iloc[0]
        for key in METRICS:assert abs(getattr(row,key)-r[key])<1e-12;checked_csv+=1
    scene=json.loads((OUT/'slide_scene.json').read_text(encoding='utf8'));prs=Presentation(OUT/(STEM+'.pptx'))
    assert len(prs.slides)==len(scene)==49
    tables={};specs={}
    all_text=[];styled_paragraphs=0
    for sl,sp in zip(prs.slides,scene):
        assert len(sl.shapes)==len(sp['shapes']) and sl.notes_slide.notes_text_frame.text==sp['note']
        all_text.append(sp['note'])
        for shape,item in zip(sl.shapes,sp['shapes']):
            assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=prs.slide_width+100 and shape.top+shape.height<=prs.slide_height+100
            if item['kind']=='text':
                assert shape.text==item['text'].replace('**','') and '\ufffd' not in shape.text
                all_text.append(item['text'])
            if item['kind']!='table':continue
            for ri,row in enumerate([item['headers']]+item['rows']):
                for ci,v in enumerate(row):
                    cell=shape.table.cell(ri,ci)
                    assert cell.text==str(v)
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            assert run.font.name=='Arial' and abs(run.font.size.pt-item['size'])<1e-6
                        styled_paragraphs+=1
            labels={v for row in [item['headers']]+item['rows'] for v in row if v in LABELS}
            if len(labels)>=2:assert labels==set(LABELS),(sp['number'],labels)
            if item.get('tag'):tables[item['tag']]=shape.table;specs[item['tag']]=item
    joined='\n'.join(all_text).lower()
    assert 'chưa chạy' not in joined and 'mới e0/e7' not in joined and 'đối chứng cần vượt' not in joined
    assert len(tables['metrics_complete'].rows)==9 and len(tables['groups_complete'].rows)==5
    assert [tables['coverage_three_models'].cell(i,5).text for i in range(1,9)]==['9']*8
    checked=0
    for day in [1,2,3]:
        t=tables[f'e7_3models_{day}']
        assert [t.cell(i,0).text for i in range(1,4)]==LABELS
        for ri,model in enumerate(MODELS,1):
            r=source[(source.model==model)&(source.experiment=='E7')&(source.lead_day==day)].iloc[0]
            for col,key in enumerate(METRICS,1):assert t.cell(ri,col).text==f'{r[key]:.3f}';checked+=1
        bss=tables[f'bss_3models_{day}']
        for i in range(8):
            for ci,model in enumerate(MODELS,1):
                r=source[(source.model==model)&(source.experiment==f'E{i}')&(source.lead_day==day)].iloc[0]
                assert bss.cell(i+1,ci).text==f'{r.BSS:.3f}';checked+=1
        for block in [1,2,3]:
            spec=specs[f'full_3models_{day}_{block}']
            for ri,row in enumerate(spec['rows'],1):
                model=MODELS[LABELS.index(row[1])];exp=row[0]
                r=source[(source.model==model)&(source.experiment==exp)&(source.lead_day==day)].iloc[0]
                for ci,key in enumerate(METRICS,2):
                    assert row[ci]==f'{r[key]:.3f}';checked+=1
                    values=source[(source.experiment==exp)&(source.lead_day==day)][key].to_numpy()
                    if key=='bias':values=abs(values)
                    goal=min(values) if key in ['RMSE','MAE','bias','BS','FAR'] else max(values)
                    value=abs(float(r[key])) if key=='bias' else float(r[key])
                    assert bool(spec['styles'][f'{ri},{ci}']['bold'])==(round(value,3)==round(float(goal),3))
    absolute=pd.read_csv(REPORT/'absolute_skill_ci.csv');paired=pd.read_csv(REPORT/'paired_bootstrap_exploratory.csv')
    for ri,model in enumerate(MODELS,1):
        for ci,day in enumerate([1,2,3],1):
            r=absolute[(absolute.model==model)&(absolute.lead_day==day)&(absolute.metric=='BSS16')&(absolute.block_days==7)].iloc[0]
            assert tables['e7_skill_ci'].cell(ri,ci).text==f'{r.estimate:+.3f}\n[{r.ci_low:+.3f}; {r.ci_high:+.3f}]'
        for ci,exp in enumerate(['E1','E2','E3'],1):
            r=paired[(paired.kind=='feature')&(paired.model==model)&(paired.experiment==exp)&(paired.lead_day==1)&(paired.metric=='BSS16')&(paired.block_days==7)].iloc[0]
            assert tables['feature_skill_ci'].cell(ri,ci).text==f'{r.estimate:+.3f}\n[{r.ci_low:+.3f}; {r.ci_high:+.3f}]'
    assert len(paired)==432 and len(absolute)==54 and (paired.ci_low<=paired.ci_high).all()
    assert not paired.confirmatory.any()
    # Regression evidence for the shared bootstrap: all previous TFT/LightGBM
    # point estimates and pointwise CIs must agree, not merely table formatting.
    prior_ci=pd.read_csv(ROOT/'reports/short_range_ablation_2026_10_21/paired_bootstrap.csv')
    bootstrap_values=0
    for r in paired[(paired.kind=='feature')&paired.model.isin(['lgbm','tft'])].itertuples():
        old=prior_ci[(prior_ci.model==r.model)&(prior_ci.contrast==r.experiment+'-E0')&
            (prior_ci.lead_day==r.lead_day)&(prior_ci.metric==r.metric)&(prior_ci.block_days==r.block_days)]
        assert len(old)==1
        for key in ['estimate','ci_low','ci_high']:
            assert abs(getattr(r,key)-old.iloc[0][key])<1e-9;bootstrap_values+=1
    for r in paired[paired.kind=='architecture'].itertuples():
        key={'BSS16':'BSS','CSI16':'CSI','RMSE':'RMSE'}[r.metric]
        a=source[(source.model==r.model)&(source.experiment=='E7')&(source.lead_day==r.lead_day)].iloc[0]
        b=source[(source.model==r.reference_model)&(source.experiment=='E7')&(source.lead_day==r.lead_day)].iloc[0]
        assert abs(r.estimate-(a[key]-b[key]))<1e-11
    workbook=load_workbook(OUT/'BANG_KET_QUA_DAY_DU.xlsx',read_only=True,data_only=True)
    assert len(workbook.sheetnames)==9 and workbook['KET_QUA_72'].max_row==73
    sheet=list(workbook['KET_QUA_72'].values)
    assert list(sheet[0])==list(data.columns)
    for a,b in zip(sheet[1:],data.itertuples(index=False,name=None)):
        for va,vb in zip(a,b):
            if isinstance(vb,float):assert abs(va-vb)<1e-11
            else:assert va==vb
    workbook.close()
    render=json.loads((OUT/'render_manifest.json').read_text(encoding='utf8'))
    assert render['status']=='passed' and render['slides']==49 and not render['overflows']
    for name,digest in render['files'].items():assert sha(OUT/name)==digest
    for path in OUT.glob('*.md'):
        for target in re.findall(r'\]\(([^)]+)\)',path.read_text(encoding='utf8')):
            if not target.startswith(('http://','https://','#')):assert (OUT/target).exists(),target
    for name in [STEM+'.pptx','GIAI_THICH_KET_QUA_DAY_DU.docx','BANG_KET_QUA_DAY_DU.xlsx']:
        with zipfile.ZipFile(OUT/name) as z:assert z.testzip() is None
    user_snapshot=PRIVATE/'user_slide_snapshot.json';preserved=[]
    if user_snapshot.exists():
        for item in json.loads(user_snapshot.read_text(encoding='utf8')):
            assert sha(ROOT/item['path'])==item['sha256'],'User slide changed during work'
            preserved.append(item['path'])
    result=dict(status='passed',slides=49,main_slides=34,complete_score_rows=72,missing_score_rows=0,
        informer_checkpoints_replayed=72,csv_metric_values_checked=checked_csv,slide_metric_values_checked=checked,
        ci_cells_checked=18,styled_table_paragraphs_checked=styled_paragraphs,
        bootstrap_regression_values_checked=bootstrap_values,
        comparison_tables_include_all_three=True,excel_sheets=9,legacy_files_verified=frozen,
        historical_worktree_missing_verified_from_git=missing_worktree,user_slides_preserved=preserved,
        html_text_overflows=0,native_office_rendered=False,test2025_read=False)
    (OUT/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf8',newline='\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':main()

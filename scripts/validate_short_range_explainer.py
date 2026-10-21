"""Check report provenance, complete result coverage and export consistency.

This reads public aggregates only. It does not load observations, checkpoints or
the 2025 holdout, and does not train or evaluate a forecasting model.
"""
from __future__ import annotations

import hashlib
import itertools
import json
from pathlib import Path
import re
import zipfile

import numpy as np
import pandas as pd
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports/short_range_ablation_2026_10_21"
OUT = ROOT / "reports/short_range_explained_2026_10_21"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def same_frame(left, right):
    pd.testing.assert_frame_equal(
        left.reset_index(drop=True), right.reset_index(drop=True),
        check_dtype=False, rtol=1e-12, atol=1e-12,
    )


def main():
    frozen_artifacts = 0
    for folder in ["plan", "pilot", "ablation", "brief"]:
        directory = ROOT / f"reports/short_range_{folder}_2026_10_21"
        manifest = json.loads((directory / "build_manifest.json").read_text(encoding="utf8"))
        for name, meta in manifest["files"].items():
            assert sha(directory / name) == meta["sha256"], (folder, name)
            frozen_artifacts += 1

    manifest = json.loads((OUT / "build_manifest.json").read_text(encoding="utf8"))
    assert manifest["training_runs_added"] == 0 and not manifest["test2025_read"]
    for name, meta in manifest["files"].items():
        assert sha(OUT / name) == meta["sha256"], name
        assert (OUT / name).stat().st_size == meta["bytes"], name
    for name, digest in manifest["source_sha256"].items():
        assert sha(SOURCE / name) == digest, name
    assert sha(ROOT / "scripts/build_short_range_explainer.py") == manifest["code_sha256"]
    run = json.loads((SOURCE / "run.json").read_text(encoding="utf8"))
    for name, digest in run["code_hashes"].items():
        assert sha(ROOT / name) == digest, name

    scores = pd.read_csv(SOURCE / "scores_ensemble_pooled.csv")
    expected = scores[(scores.season == "all") & (scores.threshold_mm == 16)
                      & scores.model.isin(["tft", "lgbm"])].sort_values(
                          ["model", "lead_day", "experiment"]).reset_index(drop=True)
    exported = pd.read_csv(OUT / "KET_QUA_CHINH_48_DONG.csv")
    same_frame(exported[expected.columns], expected)
    assert len(exported) == 48 and (exported.n == 1075).all()
    assert set(zip(exported.model, exported.experiment, exported.lead_day)) == set(
        itertools.product(["tft", "lgbm"], [f"E{i}" for i in range(8)], [1, 2, 3]))

    book_path = OUT / "BANG_KET_QUA_DAY_DU.xlsx"
    frames = pd.read_excel(book_path, sheet_name=None)
    assert list(frames) == manifest["workbook_sheets"] and len(frames) == 18
    same_frame(frames["KET_QUA_48"], exported)
    same_frame(frames["E0_E7"], pd.read_csv(OUT / "TIEN_DO_E0_E7.csv"))
    direct_sources = {
        "POOLED_MUA": "scores_ensemble_pooled.csv",
        "TUNG_NAM": "scores_ensemble.csv", "TUNG_SEED_GOP": "scores_seed_pooled.csv",
        "TUNG_SEED_NAM": "scores_seed.csv", "CI_DAY_DU": "paired_bootstrap.csv",
        "DANH_MUC_BIEN": "feature_catalog.csv", "RELIABILITY": "reliability16.csv",
        "NGUONG_CAL": "ensemble_calibration_thresholds.csv",
        "ON_DINH_SEED": "seed_stability.csv",
    }
    for sheet, filename in direct_sources.items():
        same_frame(frames[sheet], pd.read_csv(SOURCE / filename))
    same_frame(frames["NGUONG_16_50_100"], scores[
        (scores.season == "all") & scores.model.isin(["tft", "lgbm"])])
    assert len(frames["NHAT_KY_96"]) == 96 and len(frames["CI_CHINH"]) == 18
    assert set(frames["NHAT_KY_96"].year) == {2021, 2022, 2023}
    assert frames["E0_E7"].variables.tolist() == [11, 23, 31, 12, 43, 24, 32, 44]
    assert frames["E0_E7"].tabular_columns.tolist() == [87, 147, 187, 92, 247, 152, 192, 252]
    assert (frames["E0_E7"].tft_runs == 9).all()
    assert (frames["E0_E7"].lgbm_pipelines == 3).all()

    book = load_workbook(book_path)
    for sheet in book:
        assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref
    for row in list(book["HUONG_DAN"].rows)[1:]:
        assert row[0].hyperlink.location == f"'{row[0].value}'!A1"

    # Check every main-table cell as printed, not just workbook row counts.
    text = (OUT / "GIAI_THICH_DAY_DU.md").read_text(encoding="utf8")
    main = text.split("## 9.", 1)[1].split("### ", 1)[1].split(
        "### " + "\u0110\u1ed1i", 1)[0]
    lines = [line for line in main.splitlines() if re.match(r"\| E[0-7] \|", line)]
    assert len(lines) == 48
    order = [expected[(expected.model == model) & (expected.lead_day == lead)]
             for lead in [1, 2, 3] for model in ["tft", "lgbm"]]
    ordered = pd.concat(order, ignore_index=True)
    metrics = ["RMSE", "MAE", "bias", "BSS", "CSI", "POD", "FAR"]
    for line, (_, row) in zip(lines, ordered.iterrows()):
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        assert cells[0] == row.experiment
        np.testing.assert_allclose([float(v) for v in cells[1:]], row[metrics].astype(float),
                                   rtol=0, atol=0.000501)

    for filename in ["README.md", "GIAI_THICH_DAY_DU.md"]:
        value = (OUT / filename).read_text(encoding="utf8")
        assert "\ufffd" not in value and "\ue000" not in value and "\ue001" not in value
        assert "LightGBM" in value and "Light GBM" not in value
        for target in re.findall(r"\]\(([^)]+)\)", value):
            if target.startswith(("https://", "http://", "#")):
                continue
            assert (OUT / target.split("#", 1)[0]).exists(), target
    catalog = pd.read_csv(SOURCE / "feature_catalog.csv")
    for name in catalog.native_variable:
        assert f"| {name} |" in text, name
    assert "attention" in text and "a tt ention" not in text
    assert "Platt" in text and "Pla tt" not in text
    doc_path = OUT / "GIAI_THICH_DAY_DU.docx"
    with zipfile.ZipFile(doc_path) as package:
        assert package.testzip() is None
    doc = Document(doc_path)
    assert len(doc.inline_shapes) == 2 and len(doc.tables) >= 15
    assert doc.core_properties.title.startswith("Gi\u1ea3i th\u00edch")
    width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    assert all(p.width <= width for p in doc.inline_shapes)
    print(json.dumps(dict(status="passed", frozen_artifacts_verified=frozen_artifacts,
                          workbook_sheets=len(frames), main_rows_checked=48,
                          main_printed_values_checked=48 * len(metrics),
                          feature_identifiers_verified=44, word_tables=len(doc.tables),
                          word_images=len(doc.inline_shapes), training_runs_added=0,
                          test2025_read=False), indent=2))


if __name__ == "__main__":
    main()

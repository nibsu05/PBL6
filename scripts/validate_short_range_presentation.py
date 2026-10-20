"""Validate exported numbers, emphasis, provenance and slide structure."""
from pathlib import Path
import hashlib
import json
import re
import subprocess
import zipfile

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/"reports/short_range_presentation_2026_10_21"
SOURCE=ROOT/"reports/short_range_ablation_2026_10_21"


def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def main():
    manifest=json.loads((REPORT/"build_manifest.json").read_text(encoding="utf8"))
    assert manifest["slides"]==34 and manifest["main_slides"]==24 and manifest["appendix_slides"]==10
    assert manifest["training_runs_added"]==0 and not manifest["test2025_read"]
    for group in ["source_sha256","code_sha256"]:
        for name,digest in manifest[group].items():assert sha(ROOT/name)==digest,name
    for name,meta in manifest["files"].items():assert sha(REPORT/name)==meta["sha256"],name
    run=json.loads((SOURCE/"run.json").read_text(encoding="utf8"))
    for name,digest in run["code_hashes"].items():assert sha(ROOT/name)==digest,name
    # Check all prior short-range publication artifacts, including the explainer.
    frozen=0
    for folder in ["plan","pilot","ablation","brief","explained"]:
        directory=ROOT/f"reports/short_range_{folder}_2026_10_21"
        m=json.loads((directory/"build_manifest.json").read_text(encoding="utf8"))
        for name,meta in m["files"].items():assert sha(directory/name)==meta["sha256"],(folder,name);frozen+=1
    scores=pd.read_csv(SOURCE/"scores_ensemble_pooled.csv")
    scores=scores[(scores.season=="all")&(scores.threshold_mm==16)&scores.model.isin(["tft","lgbm"])]
    scene=json.loads((REPORT/"slide_scene.json").read_text(encoding="utf8"))
    prs=Presentation(REPORT/"PBL6_TFT_DAY_DU.pptx")
    assert len(scene)==len(prs.slides)==34
    tables={}
    for index,(spec,slide) in enumerate(zip(scene,prs.slides),1):
        assert len(slide.shapes)==len(spec["shapes"]),(index,len(slide.shapes),len(spec["shapes"]))
        assert len(slide.notes_slide.notes_text_frame.text)>70,index
        assert slide.notes_slide.notes_text_frame.text==spec["note"]
        for sh,ss in zip(slide.shapes,spec["shapes"]):
            assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=prs.slide_width+100 and sh.top+sh.height<=prs.slide_height+100,index
            if ss["kind"]=="text":
                assert sh.text==ss["text"].replace("**","")
                assert "\ufffd" not in sh.text,index
            if ss["kind"]=="table":
                rows=[ss["headers"]]+ss["rows"]
                for ri,row in enumerate(rows):
                    for ci,value in enumerate(row):assert sh.table.cell(ri,ci).text==str(value),(index,ri,ci)
                if ss.get("tag"):tables[ss["tag"]]=(index,sh.table,ss)
    metrics=["RMSE","MAE","bias","BSS","CSI","POD","FAR"]
    checked=0;bold_checked=0
    for model in ["tft","lgbm"]:
        for day in [1,2,3]:
            _,table,spec=tables[f"full_{model}_{day}"]
            part=scores[(scores.model==model)&(scores.lead_day==day)].sort_values("experiment")
            assert len(part)==8
            for ri,(_,row) in enumerate(part.iterrows(),1):
                assert table.cell(ri,0).text==row.experiment
                for ci,metric in enumerate(metrics,1):
                    cell=table.cell(ri,ci);assert cell.text==f'{row[metric]:.3f}'
                    rounded=part[metric].round(3).abs() if metric=="bias" else part[metric].round(3)
                    goal=rounded.min() if metric in ["RMSE","MAE","bias","FAR"] else rounded.max()
                    value=abs(round(row[metric],3)) if metric=="bias" else round(row[metric],3)
                    runs=[r for p in cell.text_frame.paragraphs for r in p.runs]
                    assert all(bool(r.font.bold)==(value==goal) for r in runs),(model,day,metric,row.experiment)
                    if metric=="BSS" and row[metric]<0:
                        assert all(str(r.font.color.rgb)=="AD3434" for r in runs)
                    checked+=1;bold_checked+=1
    counts=pd.read_csv(REPORT/"E7_CONTINGENCY_COUNTS.csv")
    for row in counts.itertuples():
        r=scores[(scores.model==row.model)&(scores.experiment=="E7")&(scores.lead_day==row.lead_day)].iloc[0]
        H,M,F,TN=row.hits,row.miss,row.false_alarm,row.correct_nonevent
        assert H+M+F+TN==r.n and H+M==r.events
        np.testing.assert_allclose([H/(H+M),F/(H+F),H/(H+M+F)],r[["POD","FAR","CSI"]].astype(float),rtol=0,atol=1e-12)
    audit=json.loads((REPORT/"dev_metrics_audit.json").read_text(encoding="utf8"))
    for path,meta in audit["files"].items():
        raw=subprocess.check_output(["git","-c","safe.directory="+ROOT.as_posix(),"show",audit["ref"]+":"+path],cwd=ROOT)
        assert hashlib.sha256(raw).hexdigest()==meta["sha256"],path
    refs=json.loads((REPORT/"literature.json").read_text(encoding="utf8"))["sources"]
    assert [r["id"] for r in refs[:3]]==["R1","R2","R3"]
    assert refs[0]["doi"]=="10.1016/j.knosys.2022.109125"
    assert refs[1]["doi"]=="10.2197/ipsjjip.32.873"
    assert refs[2]["doi"]=="10.1016/j.atmosres.2023.106961"
    with zipfile.ZipFile(REPORT/"PBL6_TFT_DAY_DU.pptx") as package:
        assert package.testzip() is None
    for path in REPORT.glob("*.md"):
        for target in re.findall(r"\]\(([^)]+)\)",path.read_text(encoding="utf8")):
            if not target.startswith(("http://","https://","#")):assert (REPORT/target).exists(),target
    render_path=REPORT/"render_manifest.json"
    rendered=False
    if render_path.exists():
        render=json.loads(render_path.read_text(encoding="utf8"))
        assert render["status"]=="passed" and not render["text_overflows"]
        assert render["html_sha256"]==sha(REPORT/"PBL6_TFT_DAY_DU.html")
        assert render["pptx_sha256"]==sha(REPORT/"PBL6_TFT_DAY_DU.pptx")
        assert render["scene_sha256"]==sha(REPORT/"slide_scene.json")
        assert render["preview_sha256"]==sha(REPORT/"XEM_TRUOC.png")
        rendered=True
    result=dict(status="passed",slides=34,main_slides=24,appendix_slides=10,
                full_result_rows=48,result_values_checked=checked,numeric_bold_rules_checked=bold_checked,
                contingency_rows_checked=6,dev_blobs_verified=len(audit["files"]),
                frozen_artifacts_verified=frozen,training_source_files_unchanged=len(run["code_hashes"]),
                html_layout_render_verified=rendered,native_office_rendered=False,
                training_runs_added=0,test2025_read=False,
                pptx_sha256=sha(REPORT/"PBL6_TFT_DAY_DU.pptx"))
    (REPORT/"validation.json").write_text(json.dumps(result,indent=2),encoding="utf8",newline="\n")
    print(json.dumps(result,indent=2))


if __name__=="__main__":main()

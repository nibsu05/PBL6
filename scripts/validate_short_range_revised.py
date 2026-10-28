"""Check the revised deliverables against original results and pilot aggregates."""
from pathlib import Path
import hashlib,json,sys,re,zipfile
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
import pandas as pd
from pptx import Presentation


def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def main():
    out=ROOT/'reports/short_range_revised_2026_10_21';source=ROOT/'reports/short_range_ablation_2026_10_21'
    manifest=json.loads((out/'build_manifest.json').read_text(encoding='utf8'))
    assert manifest['slides']==42 and manifest['main_slides']==30 and manifest['informer_runs']==18
    assert not manifest['test2025_read'] and manifest['completed_seeds']==[42,43,44]
    for name,digest in manifest['sources'].items():assert sha(ROOT/name)==digest,name
    for name,meta in manifest['files'].items():assert sha(out/name)==meta['sha256'],name
    assert sha(ROOT/'scripts/build_short_range_revised.py')==manifest['code_sha256']
    frozen=0
    for name in ['plan','pilot','ablation','brief','explained','presentation']:
        d=ROOT/f'reports/short_range_{name}_2026_10_21';m=json.loads((d/'build_manifest.json').read_text(encoding='utf8'))
        for filename,meta in m['files'].items():assert sha(d/filename)==meta['sha256'],filename;frozen+=1
    scene=json.loads((out/'slide_scene.json').read_text(encoding='utf8'));prs=Presentation(out/'PBL6_BAO_CAO_DE_HIEU.pptx')
    assert len(scene)==len(prs.slides)==42
    tables={}
    for no,(sl,sp) in enumerate(zip(prs.slides,scene),1):
        assert len(sl.shapes)==len(sp['shapes'])
        assert sl.notes_slide.notes_text_frame.text==sp['note']
        for shape,item in zip(sl.shapes,sp['shapes']):
            assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=prs.slide_width+100 and shape.top+shape.height<=prs.slide_height+100,no
            if item['kind']=='text':assert shape.text==item['text'].replace('**','') and '\ufffd' not in shape.text
            if item['kind']=='table':
                for ri,row in enumerate([item['headers']]+item['rows']):
                    for ci,value in enumerate(row):assert shape.table.cell(ri,ci).text==str(value)
                if item.get('tag'):tables[item['tag']]=shape.table
    assert len(tables['metrics_complete'].rows)==9 and len(tables['groups_complete'].rows)==5
    assert [tables['groups_complete'].cell(r,1).text for r in range(1,5)]==['11','12','20','1']
    metric_names=[tables['metrics_complete'].cell(r,0).text for r in range(1,9)]
    assert metric_names==['RMSE','MAE','Bias','BS','BSS','CSI','POD','FAR']
    assert manifest['roles']['groups']==5 and manifest['roles']['metrics']==13 and manifest['roles']['seeds']==7
    original=pd.read_csv(source/'scores_ensemble_pooled.csv');metrics=['RMSE','MAE','bias','BSS','CSI','POD','FAR'];checked=0
    for model in ['tft','lgbm']:
        for day in [1,2,3]:
            t=tables[f'full_{model}_{day}']
            for ri,e in enumerate([f'E{i}' for i in range(8)],1):
                row=original[(original.model==model)&(original.experiment==e)&(original.lead_day==day)&(original.season=='all')&(original.threshold_mm==16)].iloc[0]
                assert t.cell(ri,0).text==e
                for ci,key in enumerate(metrics,1):assert t.cell(ri,ci).text==f'{row[key]:.3f}';checked+=1
    pilot_path=ROOT/'reports/informer_regional_pilot_2026_10_21'
    validation=json.loads((pilot_path/'validation.json').read_text(encoding='utf8'))
    assert validation['status']=='passed' and validation['independent_full_replayed_runs']==18 and validation['tft_scores_match_original']
    scores=pd.read_csv(pilot_path/'scores_pooled.csv');table=tables['informer_matched'];new_checked=0
    for day in [1,2,3]:
        for offset,model in enumerate(['tft','informer']):
            ri=day*2-1+offset
            r=scores[(scores.model==model)&(scores.experiment=='E7')&(scores.lead_day==day)&(scores.season=='all')&(scores.threshold_mm==16)].iloc[0]
            for col,key in enumerate(['RMSE','BSS','CSI','POD','FAR'],2):assert table.cell(ri,col).text==f'{r[key]:.3f}';new_checked+=1
    rules=pd.read_csv(out/'DECISION_RULES_E7.csv');t=tables['rule_comparison']
    for ri,rule in enumerate(['amount','probability'],1):
        r=rules[(rules.model=='tft')&(rules.lead_day==1)&(rules.rule==rule)].iloc[0]
        assert t.cell(ri,1).text==str(int(r.hits)) and t.cell(ri,2).text==str(int(r.false_alarm))
        for col,key in enumerate(['POD','FAR','CSI'],3):assert t.cell(ri,col).text==f'{r[key]:.3f}'
    render=json.loads((out/'render_manifest.json').read_text(encoding='utf8'))
    assert render['status']=='passed' and render['slides']==42 and not render['overflows']
    for field,name in [('html_sha256','PBL6_BAO_CAO_DE_HIEU.html'),('pptx_sha256','PBL6_BAO_CAO_DE_HIEU.pptx'),('scene_sha256','slide_scene.json'),('preview_sha256','XEM_TRUOC.png')]:assert render[field]==sha(out/name)
    for path in out.glob('*.md'):
        for target in re.findall(r'\]\(([^)]+)\)',path.read_text(encoding='utf8')):
            if not target.startswith(('http://','https://','#')):assert (out/target).exists(),target
    for p in [out/'PBL6_BAO_CAO_DE_HIEU.pptx',out/'GIAI_THICH_THEO_CAU_HOI.docx']:
        with zipfile.ZipFile(p) as z:assert z.testzip() is None
    result=dict(status='passed',slides=42,main_slides=30,full_metrics_rows=8,feature_groups=4,
                original_numeric_values_checked=checked,informer_numeric_values_checked=new_checked,
                informer_full_replayed_runs=18,old_publication_artifacts_unchanged=frozen,
                html_text_overflows=0,native_office_rendered=False,test2025_read=False)
    (out/'validation.json').write_text(json.dumps(result,indent=2),encoding='utf8',newline='\n')
    print(json.dumps(result,indent=2))


if __name__=='__main__':main()

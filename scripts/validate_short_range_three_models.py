"""Verify every comparison cell and its three-model scope against frozen results."""
from pathlib import Path
import hashlib
import json
import re
import zipfile

import pandas as pd
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports/short_range_three_models_2026_10_21'
STEM = 'PBL6_SO_SANH_3_MO_HINH'
MODELS = ['lgbm', 'tft', 'informer']
LABELS = ['LightGBM', 'TFT', 'Informer']
METRICS = ['RMSE', 'MAE', 'bias', 'BS', 'BSS', 'CSI', 'POD', 'FAR']


def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    manifest = json.loads((OUT / 'build_manifest.json').read_text(encoding='utf8'))
    assert manifest['models'] == MODELS and manifest['new_training_runs'] == 0 and not manifest['test2025_read']
    for name, digest in manifest['sources'].items(): assert sha(ROOT / name) == digest, name
    for name, meta in manifest['files'].items(): assert sha(OUT / name) == meta['sha256'], name
    assert sha(ROOT / 'scripts/build_short_range_three_models.py') == manifest['code_sha256']
    frozen = 0
    for name in ['plan', 'pilot', 'ablation', 'brief', 'explained', 'presentation', 'revised']:
        directory = ROOT / f'reports/short_range_{name}_2026_10_21'
        prior = json.loads((directory / 'build_manifest.json').read_text(encoding='utf8'))
        for filename, meta in prior['files'].items(): assert sha(directory / filename) == meta['sha256']; frozen += 1
    pilot = ROOT / 'reports/informer_regional_pilot_2026_10_21'
    prior = json.loads((pilot / 'publication_manifest.json').read_text(encoding='utf8'))
    for filename, meta in prior['files'].items(): assert sha(pilot / filename) == meta['sha256']; frozen += 1
    source = pd.read_csv(ROOT / 'reports/short_range_ablation_2026_10_21/scores_ensemble_pooled.csv')
    new = pd.read_csv(pilot / 'scores_pooled.csv')
    source = pd.concat([source[source.model.isin(['tft', 'lgbm'])], new[new.model == 'informer']])
    source = source[(source.season == 'all') & (source.threshold_mm == 16)]
    records = pd.read_csv(OUT / 'KET_QUA_3_MO_HINH.csv')
    assert len(records) == 72 and len(records[records.status == 'available']) == 54
    missing = records[records.status == 'not_run']
    assert len(missing) == 18 and set(missing.model) == {'informer'}
    assert set(missing.experiment) == {f'E{i}' for i in range(1, 7)}
    assert missing[METRICS].isna().all().all()
    csv_values = 0
    for r in records[records.status == 'available'].itertuples():
        q = source[(source.model == r.model) & (source.experiment == r.experiment) & (source.lead_day == r.lead_day)]
        assert len(q) == 1 and r.n == 1075 and r.events == 114
        for key in METRICS: assert abs(getattr(r, key) - q.iloc[0][key]) < 1e-12; csv_values += 1
    scene = json.loads((OUT / 'slide_scene.json').read_text(encoding='utf8'))
    prs = Presentation(OUT / (STEM + '.pptx'))
    assert len(scene) == len(prs.slides) == manifest['slides']
    tables, table_specs = {}, {}
    for slide, spec in zip(prs.slides, scene):
        assert len(slide.shapes) == len(spec['shapes'])
        assert slide.notes_slide.notes_text_frame.text == spec['note']
        for shape, item in zip(slide.shapes, spec['shapes']):
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 100 and shape.top + shape.height <= prs.slide_height + 100
            if item['kind'] == 'text': assert shape.text == item['text'].replace('**', '') and '\ufffd' not in shape.text
            if item['kind'] != 'table': continue
            for ri, row in enumerate([item['headers']] + item['rows']):
                for ci, value in enumerate(row): assert shape.table.cell(ri, ci).text == str(value)
            # Every table comparing model labels has the complete model set.
            labels = {v for row in [item['headers']] + item['rows'] for v in row if v in LABELS}
            if len(labels) >= 2: assert labels == set(LABELS), (spec['number'], labels)
            if item.get('tag'):
                tables[item['tag']] = shape.table
                table_specs[item['tag']] = item
    assert len(tables['metrics_complete'].rows) == 9 and len(tables['groups_complete'].rows) == 5
    assert [tables['metrics_complete'].cell(i, 0).text for i in range(1, 9)] == ['RMSE', 'MAE', 'Bias', 'BS', 'BSS', 'CSI', 'POD', 'FAR']
    assert [tables['groups_complete'].cell(i, 1).text for i in range(1, 5)] == ['11', '12', '20', '1']
    coverage = tables['coverage_three_models']
    assert [coverage.cell(i, 5).text for i in range(1, 9)] == ['9', *['Chưa chạy'] * 6, '9']
    checked = blank = 0
    for day in [1, 2, 3]:
        main = tables[f'e7_3models_{day}']
        assert [main.cell(i, 0).text for i in range(1, 4)] == LABELS
        for i, model in enumerate(MODELS, 1):
            r = source[(source.model == model) & (source.experiment == 'E7') & (source.lead_day == day)].iloc[0]
            for col, key in enumerate(METRICS, 1): assert main.cell(i, col).text == f'{r[key]:.3f}'; checked += 1
        bss = tables[f'bss_3models_{day}']
        assert [bss.cell(0, i).text for i in range(1, 4)] == LABELS
        for ei in range(8):
            for col, model in enumerate(MODELS, 1):
                r = source[(source.model == model) & (source.experiment == f'E{ei}') & (source.lead_day == day)]
                expected = f'{r.iloc[0].BSS:.3f}' if len(r) else 'Chưa chạy'
                assert bss.cell(ei + 1, col).text == expected
        for block in [1, 2, 3]:
            tag = f'full_3models_{day}_{block}'
            spec = table_specs[tag]
            for i, row in enumerate(spec['rows'], 1):
                model = MODELS[LABELS.index(row[1])]
                r = source[(source.model == model) & (source.experiment == row[0]) & (source.lead_day == day)]
                for col, key in enumerate(METRICS, 2):
                    assert row[col] == (f'{r.iloc[0][key]:.3f}' if len(r) else '—')
                    if len(r): checked += 1
                    else: blank += 1
                    # No two-model winner is highlighted while Informer is missing.
                    if row[0] not in ['E0', 'E7']: assert not spec['styles'].get(f'{i},{col}', {}).get('bold', False)
    counts = tables['counts_three_models']
    assert [counts.cell(i, 0).text for i in range(1, 4)] == LABELS
    assert [[counts.cell(i, j).text for j in range(1, 4)] for i in range(1, 4)] == [['60', '54', '32'], ['73', '41', '66'], ['70', '44', '45']]
    render = json.loads((OUT / 'render_manifest.json').read_text(encoding='utf8'))
    assert render['status'] == 'passed' and not render['overflows'] and render['slides'] == len(scene)
    for name, digest in render['files'].items(): assert sha(OUT / name) == digest
    for file in OUT.glob('*.md'):
        for target in re.findall(r'\]\(([^)]+)\)', file.read_text(encoding='utf8')):
            if not target.startswith(('http://', 'https://', '#')): assert (OUT / target).exists(), target
    for name in [STEM + '.pptx', 'GIAI_THICH_SO_SANH_3_MO_HINH.docx']:
        with zipfile.ZipFile(OUT / name) as z: assert z.testzip() is None
    result = dict(status='passed', slides=len(scene), main_slides=manifest['main_slides'],
                  complete_model_set_in_comparison_tables=True, full_metrics_rows=8, feature_groups=4,
                  csv_values_checked=csv_values, slide_metric_values_checked=checked,
                  missing_informer_cells_verified=blank, old_artifacts_unchanged=frozen,
                  html_text_overflows=0, native_office_rendered=False, new_training_runs=0, test2025_read=False)
    (OUT / 'validation.json').write_text(json.dumps(result, indent=2), encoding='utf8', newline='\n')
    print(json.dumps(result, indent=2))


if __name__ == '__main__': main()
